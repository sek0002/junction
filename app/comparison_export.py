"""Printable comparison results; no sequences or database writes are required."""
import io
from datetime import datetime,timezone
from openpyxl import Workbook
from openpyxl.styles import PatternFill,Font

from .exports import _format_sheet,_safe_cell
from .plate_labels import label_workbook_plates


def exact_match(match):
    return (match.get('exact') is True and match.get('identity')==1 and match.get('coverage',1)==1
            and match.get('reference_coverage',1)==1 and match.get('query_coverage',1)==1)


def span_text(spans):
    return '; '.join(f'{int(a)+1}–{int(b)}' if a!=b else f'gap after {int(a)}' for a,b in spans)


def comparison_workbook(records,scope):
    chosen=[]
    for record in records:
        matches=record.get('matches') or []
        perfect=next((m for m in matches if exact_match(m)),None)
        if scope=='exact' and (record.get('error') or perfect is None):continue
        chosen.append((record,perfect or (matches[0] if matches else {})))
    if not chosen:raise ValueError('No results match the export selection.')
    if sum(len(r.get('fragment_matches',[]))+len(m.get('changes',[])) for r,m in chosen)>100000:
        raise ValueError('Select fewer files; this export exceeds 100,000 detail rows.')
    wb=Workbook();results=wb.active;results.title='Results'
    results.append(['File','Record','Query bp','Topology','Closest match','Source','Reference bp','Identity %','100% match',
                    'Strand','Query rotation','Substituted bp','Inserted bp','Deleted bp','Fragments','Exact fragments','Enzymes','Aliases','Status',
                    'Reference coverage %','Query coverage %','Aligned bp'])
    fragments=wb.create_sheet('Fragments');fragments.append(['File','Record','Reference','Fragment','Plate','Well','Representation','Query positions (1-based)','Strand','Identity %','Coverage %','100% match','Status'])
    changes=wb.create_sheet('Differences');changes.append(['File','Record','Reference','Type','Reference positions (1-based)','Original query positions (1-based)','Reference bases','Query bases'])
    def append(sheet,row):sheet.append([_safe_cell(v) for v in row])
    for record,match in chosen:
        filename=record.get('filename','');name=record.get('query_name','');hits=record.get('fragment_matches') or []
        diffs=match.get('changes') or [];identity=match.get('identity')
        counts={kind:sum(len(c.get('query_bases' if kind=='insertion' else 'reference_bases','')) for c in diffs if c.get('type')==kind) for kind in ('substitution','insertion','deletion')}
        detection=record.get('enzyme_detection') or {}
        append(results,[filename,name,record.get('query_length'), 'Circular' if record.get('query_circular') else 'Linear',match.get('name',''),match.get('source',''),match.get('reference_length'),
                        identity*100 if identity is not None else None,'Yes' if exact_match(match) else 'No',match.get('orientation',''),match.get('rotation'),
                        counts['substitution'],counts['insertion'],counts['deletion'],len(hits),sum(exact_match(h) for h in hits),
                        '; '.join(detection.get('compatible_assembly_modes') or detection.get('tested') or []),
                        '; '.join(a.get('name','') for a in match.get('aliases',[])),record.get('error') or ('Exact' if exact_match(match) else 'Near match' if match else 'No assembly match'),
                        match['reference_coverage']*100 if match.get('reference_coverage') is not None else None,
                        match['query_coverage']*100 if match.get('query_coverage') is not None else None,match.get('aligned_bp')])
        for hit in hits:
            append(fragments,[filename,name,match.get('name',''),hit.get('name',''),hit.get('plate_id',''),hit.get('well',''),hit.get('representation',''),span_text(hit.get('query_spans',[])),
                             hit.get('strand',''),hit.get('identity',0)*100,hit.get('coverage',0)*100,'Yes' if exact_match(hit) else 'No',hit.get('status','')])
        for change in diffs:
            append(changes,[filename,name,match.get('name',''),change.get('type',''),span_text([[change.get('reference_start',0),change.get('reference_end',0)]]),span_text(change.get('original_query_spans',[])),change.get('reference_bases',''),change.get('query_bases','')])
    metadata=wb.create_sheet('Export');metadata.append(['Setting','Value'])
    metadata.append(['Selection','100% matches' if scope=='exact' else 'Selected files'])
    metadata.append(['Records',len(chosen)])
    metadata.append(['Exported (UTC)',datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')])
    metadata.append(['Coordinates','1-based, inclusive; gap positions are between bases. Query coordinates refer to the original uploaded sequence.'])
    metadata.append(['Identity','Identical alignment columns / all alignment columns, including gaps. Exact requires the exact-match flag and identity 1.'])
    metadata.append(['Coverage','Paired reference/query bases, including substitutions but excluding gaps, divided by reference/query length respectively. Full coverage alone is not an exact match.'])
    metadata.append(['Scope','One closest reference per record; fragments can overlap. Not detected is not proof of absence.'])
    for sheet,widths in [(results,[28,28,12,14,35,24,14,14,14,10,14,16,14,14,12,15,25,45,35,19,19,14]),(fragments,[28,28,32,30,25,9,25,30,10,14,14,14,20]),(changes,[28,28,32,16,26,30,45,45]),(metadata,[24,95])]:
        _format_sheet(sheet,widths,paper='A3' if sheet==results else 'A4');sheet.oddFooter.right.text='Junction comparison'
    for sheet,column in ((results,9),(fragments,12)):
        for row in range(2,sheet.max_row+1):
            if sheet.cell(row,column).value=='Yes':
                cell=sheet.cell(row,column);cell.fill=PatternFill('solid',fgColor='E9F3EE');cell.font=Font(name='Aptos',size=10,bold=True,color='176C58')
    for sheet,columns in ((results,(8,20,21)),(fragments,(10,11))):
        for col in columns:
            for row in range(2,sheet.max_row+1):sheet.cell(row,col).number_format='0.0000'
    label_workbook_plates(wb)
    output=io.BytesIO();wb.save(output);return output.getvalue()
