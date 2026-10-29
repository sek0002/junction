"""One row per reaction, with stock/location metadata separated from transfers."""
import csv
import colorsys
import hashlib
import io
import math
import re
from xml.sax.saxutils import escape

from .exports import _safe_cell
from .bench_exports import plate_key

TEXT_PALETTE = ['225EA8', 'A4470A', '7343A1', '23713D', 'AA3364', '876500', '007D85', '5353A8', '865A36', '5C7116', '97408F', '476581']
META_HEADERS = ['Fragment', 'Type', 'Plate', 'Well', 'Physical bp', 'Stock ng/µL', 'Resuspend µL', 'Yield ng', 'Required µL', 'Required ng', 'Stock basis', 'Assemblies']



def _plate_text_colors(key):
    """Readable, deterministic text colors, unique within each exported plate set."""
    used=set()
    for name,entry in key.items():
        seed=int(hashlib.sha256(name.encode()).hexdigest()[:8],16)
        color=next((TEXT_PALETTE[(seed+i)%len(TEXT_PALETTE)] for i in range(len(TEXT_PALETTE)) if TEXT_PALETTE[(seed+i)%len(TEXT_PALETTE)] not in used),None)
        while color is None or color in used:
            color=''.join(f'{round(v*255):02X}' for v in colorsys.hls_to_rgb((seed%360)/360,.26,.5+((seed//360)%30)/100))
            seed+=1
        used.add(color);entry['text_color']=color


def _plate_runs(value,key):
    """Color just the plate alias in each (Plate-Well-Volume) tuple."""
    aliases={entry['short']:entry['text_color'] for entry in key.values()}
    text=str(value)
    if not aliases:return [(text,None)]
    pattern=r'\(('+'|'.join(re.escape(alias) for alias in sorted(aliases,key=len,reverse=True))+r')-'
    runs=[];position=0
    for match in re.finditer(pattern,text):
        runs.extend([(text[position:match.start(1)],None),(match[1],aliases[match[1]])])
        position=match.end(1)
    runs.append((text[position:],None))
    return runs


def _pack_inputs(value,fits):
    """Keep transfer order while fitting whole tuples side by side."""
    lines=[];current=''
    for item in str(value).splitlines():
        candidate=current+'  '+item if current else item
        if current and not fits(candidate):
            lines.append(current);current=item
        else:current=candidate
    if current:lines.append(current)
    return '\n'.join(lines)

def matrix_data(sections):
    metadata=[];lookup={};key=plate_key(sections);_plate_text_colors(key);rows=[]
    headers=['Assembly','Enzyme','Inputs (Plate-Well-µL)','DNA µL','Other reagents µL','Water µL','Total µL','Checks']
    for section in sections:
        inputs=[];reagents={};issues=list(section['issues']);dna=0;known=True;count=0
        for row in section['rows']:
            if row['kind'] in ('reagent','total'):
                reagents[row['name']]=row['use'];continue
            count+=1;identity=tuple(row.get(k) for k in ('kind','id','name','plate','well','physical_length','stock','resuspend','yield_ng','basis'))
            owner=section['name']+' / '+section['stage']
            if identity not in lookup:
                lookup[identity]=len(metadata)
                metadata.append([row['name'],row['kind'],key.get(row['plate'],{}).get('short',''),row['well'],row.get('physical_length'),row['stock'],row['resuspend'],row['yield_ng'],0.0,0.0,row['basis'],owner])
            elif owner not in metadata[lookup[identity]][11].split('\n'):
                metadata[lookup[identity]][11]+='\n'+owner
            material=metadata[lookup[identity]]
            for column,value in ((8,row['use']),(9,row.get('required_ng'))):
                material[column]=material[column]+value if material[column] is not None and value is not None else None
            volume=row['use'];plate=key.get(row['plate'],{}).get('short') or row['name'];well=row['well'] or '—'
            inputs.append(f"({plate}-{well}-{'?' if volume is None else format(volume,'.2f')})")
            if volume is None:known=False
            else:dna+=volume
            issues.extend(row['name']+': '+issue for issue in row['checks'])
        checks='; '.join(dict.fromkeys(issues)) or ('Pass' if section['valid'] else 'Check')
        name=section['name']+((' / '+section['stage']) if section['stage'] not in ('Final assembly','Unavailable','Stock preparation') else '')
        rows.append([name,section['enzyme'],'\n'.join(inputs),
                    dna if known and count else None,reagents.get('Master mix + reagents'),reagents.get('Water'),reagents.get('Total'),checks])
    return headers,rows,metadata,key


def _csv(sections):
    headers,rows,metadata,key=matrix_data(sections);out=io.StringIO();writer=csv.writer(out)
    def write(values):writer.writerow([_safe_cell(f'{v:.2f}' if isinstance(v,float) else v) for v in values])
    write(['Assemblies']);write(headers)
    for row in rows:write(row)
    write([]);write(['Fragment metadata']);write(META_HEADERS)
    for row in metadata:write(row)
    write([]);write(['Plate key']);write(['Plate','Full plate identifier','Color'])
    for entry in key.values():write([entry['short'],entry['name'],'#'+entry['text_color']])
    return out.getvalue().encode('utf-8-sig')


def _excel(sections,legacy_factory):
    from openpyxl import Workbook,load_workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.cell.rich_text import CellRichText,TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.utils import get_column_letter
    headers,rows,metadata,key=matrix_data(sections)
    wb=load_workbook(io.BytesIO(legacy_factory())) if legacy_factory else Workbook()
    if not legacy_factory:wb.remove(wb.active)
    tables=[('Assemblies',headers,rows,[18,9,43,7,10,7,7,12]),
            ('Fragment metadata',META_HEADERS,metadata,[32,9,9,7,9,12,12,11,12,12,18,17]),
            ('Plate key',['Plate','Full plate identifier','Color'],[[e['short'],e['name'],'#'+e['text_color']] for e in key.values()],[10,56,12])]
    for index,(name,columns,values,widths) in enumerate(tables):
        ws=wb.create_sheet(name,index);ws.append([name]);ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(columns));ws['A1'].font=Font(name='Arial',size=14,bold=True,color='1D443B');ws.row_dimensions[1].height=24
        # Keep the title/header row coordinates stable for spreadsheet consumers.
        ws.append([]);ws.row_dimensions[2].hidden=True
        ws.append(columns)
        for row in values:ws.append([_safe_cell(v) for v in row])
        for c in ws[3]:c.fill=PatternFill('solid',fgColor='1D443B');c.font=Font(name='Arial',size=9,bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[3].height=36
        if index==0:
            for r in range(4,ws.max_row+1):
                ws.cell(r,3).value=_pack_inputs(ws.cell(r,3).value,lambda text:len(text)<=widths[2]-2)
        for r in range(4,ws.max_row+1):
            lines=1
            for c,width in enumerate(widths,1):
                cell=ws.cell(r,c);cell.fill=PatternFill('solid',fgColor='F1F6F3' if r%2==0 else 'FFFFFF');cell.alignment=Alignment(wrap_text=True,vertical='top',horizontal='center' if index==0 and 4<=c<=7 else None);cell.font=Font(name='Arial',size=11)
                if isinstance(cell.value,float):cell.number_format='0.00'
                if index!=1 or c<=10:
                    display=f'{cell.value:.2f}' if isinstance(cell.value,float) else str(cell.value or '')
                    lines=max(lines,sum(max(1,math.ceil(len(line)/max(8,width-2))) for line in display.split('\n')))
            ws.row_dimensions[r].height=max(22,lines*15+4)
            if index==0:
                cell=ws.cell(r,3)
                cell.value=CellRichText([TextBlock(InlineFont(rFont='Arial',sz=11,color=color,b=True),text) if color else text for text,color in _plate_runs(cell.value,key)])
            elif index==1:
                plate=next((e for e in key.values() if e['short']==ws.cell(r,3).value),None)
                if plate:ws.cell(r,3).font=Font(name='Arial',size=11,bold=True,color=plate['text_color'])
            elif index==2:
                for col in (1,2):ws.cell(r,col).font=Font(name='Arial',size=11,bold=True,color=values[r-4][2].lstrip('#'))
        for c,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(c)].width=width
        ws.freeze_panes='C4';ws.auto_filter.ref=f'A3:{get_column_letter(len(columns))}{ws.max_row}';ws.sheet_view.showGridLines=False;ws.print_title_rows='1:3'
        ws.oddFooter.center.text='Junction | Page &P of &N'
        if index==1:
            # Ownership is already shown by the assembly input tuples. Keep the
            # audit fields available to unhide, without inflating printed rows.
            ws.column_dimensions['K'].hidden=True;ws.column_dimensions['L'].hidden=True
        ws.print_area=f'A1:{get_column_letter(10 if index==1 else len(columns))}{ws.max_row}'
    for ws in wb:
        ws.page_setup.orientation='portrait';ws.page_setup.paperSize=ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.page_margins.left=ws.page_margins.right=.25;ws.page_margins.top=ws.page_margins.bottom=.35
    wb.active=0;out=io.BytesIO();wb.save(out);return out.getvalue()


def _pdf(sections):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak,LongTable,TableStyle
    headers,rows,metadata,key=matrix_data(sections)
    page_width,_=A4;available=page_width-52;out=io.BytesIO()
    widths=[82,40,205,33,49,33,33,available-475]
    styles={k:ParagraphStyle(k,fontSize=s,leading=s+2,fontName='Helvetica-Bold' if k in ('title','head') else 'Helvetica',textColor=colors.white if k=='head' else colors.HexColor('#1D443B')) for k,s in [('title',14),('head',7),('cell',8),('meta',7),('note',8)]}
    def packed_inputs(value):
        def fits(text):
            return sum(stringWidth(run,'Helvetica-Bold' if color else 'Helvetica',8) for run,color in _plate_runs(text,key))<=widths[2]-6
        return _pack_inputs(value,fits)
    def p(v,kind='cell',plate_label=False,input_tuples=False):
        if isinstance(v,float):v=f'{v:.2f}'
        text='—' if v is None else str(v)
        if input_tuples:runs=_plate_runs(text,key)
        elif plate_label:
            plate=next((e for e in key.values() if text in (e['short'],e['name'])),None)
            runs=[(text,plate['text_color'] if plate else None)]
        else:runs=[(text,None)]
        markup=''.join(('<font color="#'+color+'"><b>'+escape(text)+'</b></font>') if color else escape(text) for text,color in runs)
        return Paragraph(markup.replace('µ','&#181;').replace('\n','<br/>'),styles[kind])
    def table(columns,values,widths,kind='cell'):
        data=[[p(c,'head') for c in columns]]+[[p(v,kind,plate_label=columns[i] in ('Plate','Full plate identifier'),input_tuples=columns[i]=='Inputs (Plate-Well-µL)') for i,v in enumerate(row)] for row in values]
        t=LongTable(data,colWidths=widths,repeatRows=1,splitInRow=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#1D443B')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.HexColor('#F1F6F3'),colors.white]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),3),('RIGHTPADDING',(0,0),(-1,-1),3),('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3)]));return t
    display_rows=[row[:2]+[packed_inputs(row[2])]+row[3:-1]+['Pass' if row[-1]=='Pass' else 'Check'] for row in rows]
    flow=[p('Assemblies','title'),Spacer(1,7),table(headers,display_rows,widths)]
    for row in rows:
        if row[-1]!='Pass':flow.extend([Spacer(1,7),p('CHECK: '+row[0]+': '+row[-1],'note')])
    flow.extend([PageBreak(),p('Fragment metadata','title'),Spacer(1,10)])
    meta=table(META_HEADERS[:10],[row[:10] for row in metadata],[144,42,44,25,42,50,50,45,50,available-492])
    flow.extend([meta,Spacer(1,16),p('Plate key','title'),Spacer(1,10)])
    plates=table(['Plate','Full plate identifier'],[[e['short'],e['name']] for e in key.values()],[70,available-70])
    flow.append(plates)
    def footer(canvas,doc):
        canvas.saveState();canvas.setFont('Helvetica',8);canvas.drawString(20,14,'Junction');canvas.drawRightString(page_width-20,14,f'Page {doc.page}');canvas.restoreState()
    SimpleDocTemplate(out,pagesize=A4,leftMargin=20,rightMargin=20,topMargin=20,bottomMargin=26,title='Junction assemblies').build(flow,onFirstPage=footer,onLaterPages=footer)
    return out.getvalue()


def matrix_bytes(format,sections,legacy_factory=None):
    if format=='xlsx':return _excel(sections,legacy_factory)
    if format=='csv':return _csv(sections)
    if format=='pdf':return _pdf(sections)
    raise ValueError('Unsupported export format')
