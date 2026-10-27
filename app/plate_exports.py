"""Read-only plate snapshots: common data for printable and tabular exports."""
import csv
import io
import math
from xml.sax.saxutils import escape
from zipfile import ZipFile, ZIP_DEFLATED

from .exports import _safe_cell
from .plate_labels import plate_label

PMOL = .05
CAPTION = 'Backbone 0.05 pmol | Insert:backbone 1:1 | Each fragment 0.05 pmol'
NOTE = 'Use = physical bp × 660 × 0.05 / 1000 / stock ng/µL. Stocks and resuspension are current saved/planned values; no records are changed. Blank = unknown.'
HEADERS = ['Plate', 'Well', 'Fragment', 'Sequence bp', 'Physical bp', 'Resuspend µL', 'Stock ng/µL', 'Use µL', 'Enzymes', "5' overhang", "3' overhang", 'Stock basis', 'Status', 'Notes', 'Target pmol', 'Insert:backbone']
MIME = {'pdf':'application/pdf', 'xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'csv':'application/zip'}


def plate_title(data):
    label=data.get('display_id') or plate_label(data['plate_id'])
    return label if label==data['plate_id'] else label+' · '+data['plate_id']


def palindromic(value):
    return bool(value) and value == value.translate(str.maketrans('ACGT', 'TGCA'))[::-1]


def well_key(row):
    well=row.get('well','')
    return (well[:1], int(well[1:]) if well[1:].isdigit() else 0, row.get('name',''))


def snapshot(plate, layout=None):
    """Keep actual stock values even when they violate the requested aliquot limits."""
    result={**plate, 'display_id':plate.get('display_id') or plate_label(plate['plate_id']), 'format':layout or plate['format'], 'rows':[]}
    for part in sorted(plate['fragments'],key=well_key):
        inv=part['inventory'];stock=inv['stock_ng_ul'];length=part['physical_length']
        target=length*660*PMOL/1000 if length else None
        use=target/stock if target is not None and stock and math.isfinite(stock) and stock>0 else None
        warnings=[]
        if use is None:warnings.append('Missing length or stock')
        elif not 1-1e-9<=use<=5+1e-9:warnings.append('Use outside 1–5 µL')
        if target and inv['yield_ng'] is not None and inv['yield_ng']<target:warnings.append('Insufficient yield for one reaction')
        if inv['resuspension_ul'] is not None and inv['resuspension_ul']>=100:warnings.append('Resuspension ≥100 µL')
        ends={end['label']:end['overhang'] for end in part['fusion_ends']}
        left,right=ends.get('5′',''),ends.get('3′','')
        if any(palindromic(s.get('overhang','')) for s in part['sites']):warnings.append('Palindromic overhang')
        archived=bool(plate.get('archived_at') or part.get('archived_at'))
        basis=('Saved' if inv['saved'] else 'Planned (unmeasured)')
        if inv['saved'] and 'unmeasured' in (inv.get('concentration_basis') or '').lower():basis='Saved plan (unmeasured)'
        values=[result['display_id'],part['well'],part['name'],part['length'],length,inv['resuspension_ul'],stock,use,
                ', '.join(part['detected_enzymes']),left,right,basis,'Archived' if archived else 'Active',
                '; '.join(warnings),PMOL,'1:1']
        result['rows'].append({'values':values,'part':part,'archived':archived,'warnings':warnings})
    return result


def plate_png(data):
    """Render a fixed-coordinate plate with names; the well list retains full labels."""
    from PIL import Image, ImageDraw, ImageFont
    columns,rows=(24,16) if data['format']==384 else (12,8)
    cell_w,cell_h=132,64
    width,height=columns*cell_w+80,rows*cell_h+150
    image=Image.new('RGB',(width,height),'white');draw=ImageDraw.Draw(image)
    font=ImageFont.load_default(size=16);small=ImageFont.load_default(size=13);title=ImageFont.load_default(size=25)
    draw.text((24,16),plate_title(data),font=title,fill='#153d38')
    draw.text((24,52),f"{data['format']} wells · {len(data['rows'])} occupied · yellow: palindromic · grey: archived",font=small,fill='#465b57')
    for c in range(columns):draw.text((80+c*cell_w+cell_w/2,86),str(c+1),font=font,fill='#465b57',anchor='mm')
    occupied={row['part']['well']:row for row in data['rows']}
    for r in range(rows):
        draw.text((36,110+r*cell_h+cell_h/2),chr(65+r),font=font,fill='#465b57',anchor='mm')
        for c in range(columns):
            well=chr(65+r)+str(c+1);row=occupied.get(well);x,y=66+c*cell_w,108+r*cell_h
            fill='#f3f6f5' if row is None else '#e1e5e4' if row['archived'] else '#d8ecdf'
            if row and any(palindromic(s.get('overhang','')) for s in row['part']['sites']):fill='#fff0b5'
            draw.rounded_rectangle((x,y,x+cell_w-6,y+cell_h-6),radius=9,fill=fill,outline='#aabdb7',width=1)
            draw.text((x+7,y+4),well,font=small,fill='#46665d')
            if row:
                name=row['part']['name'];lines=['']
                for char in name:
                    if draw.textlength(lines[-1]+char,font=small)>cell_w-20:
                        if len(lines)==2:
                            while draw.textlength(lines[-1]+'…',font=small)>cell_w-20:lines[-1]=lines[-1][:-1]
                            lines[-1]+='…';break
                        lines.append(char.lstrip())
                    else:lines[-1]+=char
                for i,line in enumerate(lines):draw.text((x+7,y+21+i*16),line,font=small,fill='#153d38')
    output=io.BytesIO();image.save(output,format='PNG');return output.getvalue()


def csv_bytes(headers, rows):
    output=io.StringIO();writer=csv.writer(output);writer.writerow(headers)
    for row in rows:writer.writerow([_safe_cell(value) if isinstance(value,str) else round(value,2) if isinstance(value,float) and value!=PMOL else value for value in row])
    return output.getvalue().encode('utf-8-sig')


def excel_bytes(data, png):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb=Workbook();plate=wb.active;plate.title='Plate map'
    plate.append([_safe_cell(plate_title(data))]);plate.append([CAPTION]);plate.append([NOTE])
    plate.merge_cells('A1:P1');plate.merge_cells('A2:P2');plate.merge_cells('A3:P4');plate['A3'].alignment=Alignment(wrap_text=True,vertical='top')
    picture=Image(io.BytesIO(png));picture.height=picture.height*1050/picture.width;picture.width=1050;plate.add_image(picture,'A6')
    for c in range(1,17):plate.column_dimensions[get_column_letter(c)].width=10
    plate.print_area=f'A1:P{6+math.ceil(picture.height/20)}'
    for name,headers,rows,widths in [
        ('Wells',HEADERS,[row['values'] for row in data['rows']],[25,8,35,13,13,16,17,13,24,15,15,27,12,48,13,18])]:
        ws=wb.create_sheet(name);ws.append([_safe_cell(plate_title(data))]);ws.append([CAPTION]);ws.append([NOTE]);ws.merge_cells(start_row=3,start_column=1,end_row=3,end_column=len(headers));ws.row_dimensions[3].height=30;ws['A3'].alignment=Alignment(wrap_text=True,vertical='top');ws.append([]);ws.append(headers)
        for row in rows:ws.append([_safe_cell(v) if isinstance(v,str) else v for v in row])
        for cell in ws[5]:cell.fill=PatternFill('solid',fgColor='235C50');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True)
        ws.row_dimensions[5].height=30
        for r in range(6,ws.max_row+1):
            for cell in ws[r]:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(cell.value,float):cell.number_format='0.00'
                if r%2==0:cell.fill=PatternFill('solid',fgColor='F1F6F3')
            if name=='Wells':
                for c in (10,11):
                    if palindromic(ws.cell(r,c).value):ws.cell(r,c).fill=PatternFill('solid',fgColor='FFF0B5')
                ws.row_dimensions[r].height=max(32,14*max(math.ceil(len(str(v or ''))/max(8,width-2)) for v,width in zip(rows[r-6],widths)))
        for c,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(c)].width=width
        ws.freeze_panes='D6';ws.auto_filter.ref=f'A5:{get_column_letter(len(headers))}{ws.max_row}';ws.print_title_rows='1:5'
    for ws in wb:
        ws.sheet_view.showGridLines=False;ws['A1'].font=Font(size=16,bold=True,color='235C50');ws['A2'].font=Font(size=11,color='46665D')
        ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.oddFooter.center.text='Page &P of &N'
    output=io.BytesIO();wb.save(output);return output.getvalue()


def pdf_bytes(data,png):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak, LongTable, TableStyle
    output=io.BytesIO();styles=getSampleStyleSheet();small=ParagraphStyle('PlateSmall',fontName='Helvetica',fontSize=8,leading=10)
    paragraph=lambda text:Paragraph(escape(str(text)).replace('µ','&#181;').replace('′',"'").replace('→',' to ').replace('≥','>=').replace('–','-').replace('\n','<br/>'),small)
    story=[Paragraph(escape(plate_title(data)),styles['Heading1']),paragraph(CAPTION),Spacer(1,10)]
    image=Image(io.BytesIO(png));image.drawHeight=image.imageHeight*780/image.imageWidth;image.drawWidth=780;story.extend([image,Spacer(1,12),paragraph(NOTE),Spacer(1,8),paragraph('Full labels, stocks and transfers follow. Planned stocks are unmeasured. Yellow marks palindromic overhangs. Archived wells remain labelled.'),PageBreak(),Paragraph('Well list',styles['Heading2']),paragraph(CAPTION+' | Planned stocks are unmeasured'),Spacer(1,8)])
    columns=[1,2,3,4,5,6,7,8,9,10,11]
    headers=['Well','Fragment','Seq.\nbp','Physical\nbp','Resuspend\nµL','Stock\nng/µL','Use\nµL','Enzymes',"5'","3'",'Basis']
    table_rows=[[paragraph(v) for v in headers]]
    for row in data['rows']:
        vals=row['values'];values=[vals[c] for c in columns];values[1]+= ' [Archived]' if row['archived'] else '';values[10]='Saved plan' if 'Saved plan' in values[10] else 'Planned' if 'Planned' in values[10] else values[10]
        table_rows.append([paragraph(f'{v:.2f}' if isinstance(v,float) else '—' if v is None else v) for v in values])
    def table(rows,widths):
        result=LongTable(rows,colWidths=widths,repeatRows=1,hAlign='LEFT')
        result.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#d8ecdf')),('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f3f7f5')]),('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor('#235c50')),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]));return result
    wells=table(table_rows,[32,167,43,43,57,57,48,84,45,45,59])
    for r,row in enumerate(data['rows'],1):
        for c in (8,9):
            if palindromic(row['values'][c+1]):wells.setStyle(TableStyle([('BACKGROUND',(c,r),(c,r),colors.HexColor('#fff0b5'))]))
    story.append(wells)
    for row in data['rows']:
        warnings=[warning for warning in row['warnings'] if warning!='Palindromic overhang']
        if warnings:story.extend([Spacer(1,5),paragraph(row['part']['well']+' · '+row['part']['name']+': '+'; '.join(warnings))])
    def footer(canvas,doc):
        canvas.saveState();canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#46665d'));canvas.drawString(26,16,data['plate_id']);canvas.drawRightString(816,16,str(doc.page));canvas.restoreState()
    SimpleDocTemplate(output,pagesize=landscape(A4),leftMargin=26,rightMargin=26,topMargin=22,bottomMargin=28,title=data['plate_id']+' well list').build(story,onFirstPage=footer,onLaterPages=footer)
    return output.getvalue()


def export_bytes(plate,format,layout=None):
    data=snapshot(plate,layout);png=plate_png(data)
    if format=='xlsx':return excel_bytes(data,png)
    if format=='pdf':return pdf_bytes(data,png)
    output=io.BytesIO()
    with ZipFile(output,'w',ZIP_DEFLATED) as archive:
        archive.writestr('wells.csv',csv_bytes(HEADERS,[row['values'] for row in data['rows']]))
        archive.writestr('plate-map.png',png)
        archive.writestr('Readme.txt',plate_title(data)+'\n'+CAPTION+'\n'+NOTE+'\n'+f"Layout: {data['format']} wells\n")
    return output.getvalue()
