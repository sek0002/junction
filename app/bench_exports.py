"""Compact, shared laboratory guides for Excel, PDF and CSV exports.

All formats consume the same calculated snapshot. They never recalculate a
different molar ratio or replace missing stock concentrations with assumptions.
"""
from __future__ import annotations

import csv
import hashlib
import io
import math
from xml.sax.saxutils import escape

from .exports import _safe_cell, calculate_volumes, reaction_summary, resuspension_plan
from .plate_labels import plate_aliases

FORMATS = {
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'pdf': 'application/pdf',
    'csv': 'text/csv; charset=utf-8',
}
PALETTE = ['D9EAF7', 'FCE2CC', 'E6DDF3', 'D8ECDD', 'F7DDE4', 'FFF0BE', 'D5EEEE', 'E5E6FA', 'E8DFC9', 'DDE6C8', 'F4DAED', 'D8E2ED']


def plate_key(sections):
    plates = sorted({str(row['plate']) for section in sections for row in section['rows'] if row.get('plate')})
    aliases = plate_aliases(plates)
    key = {}
    used_colors = set()
    for plate in plates:
        short = aliases[plate]
        seed = int(hashlib.sha256(plate.encode()).hexdigest()[:8], 16)
        color = next((PALETTE[(seed + i) % len(PALETTE)] for i in range(len(PALETTE)) if PALETTE[(seed + i) % len(PALETTE)] not in used_colors), None)
        if color is None:
            color = ''.join(f'{190 + (seed >> shift) % 55:02X}' for shift in (0, 8, 16))
            while color in used_colors:
                seed += 1
                color = ''.join(f'{190 + (seed >> shift) % 55:02X}' for shift in (0, 8, 16))
        used_colors.add(color)
        key[plate] = {'short': short, 'color': color, 'name': plate}
    return key


def _checks(values):
    return list(dict.fromkeys(str(value) for value in values if value))


def assembly_sections(assembly, settings, enzyme=''):
    """Keep physically separate nested reactions before their final reaction."""
    sections = []
    for stage in assembly.get('nested_reactions', []):
        sections.extend(assembly_sections({**stage, 'name': assembly.get('name', 'Assembly'), 'stage': stage['name']}, stage['settings'], stage['enzyme']))
    reaction = assembly.get('reaction')
    if reaction is None:
        rows = calculate_volumes(assembly.get('parts', []), settings)
        summary = reaction_summary(rows, settings)
        reaction = dict(rows=rows, water_ul=summary['water_volume_ul'], master_mix_ul=summary['master_mix_volume_ul'], reaction_volume_ul=summary['reaction_volume_ul'], warnings=summary['warnings'])
    prep = {str(row.get('id')): row for row in assembly.get('resuspension', {}).get('rows', [])}
    rows = []
    for index, row in enumerate(reaction['rows'], 1):
        source = prep.get(str(row.get('id')), {})
        resuspend = row.get('recorded_resuspension_ul')
        if resuspend is None:
            resuspend = source.get('resuspension_ul')
        if (row.get('role') == 'backbone' and row.get('recorded_resuspension_ul') is None) or str(row.get('id', '')).startswith('prepared:'):
            resuspend = None
        rows.append(dict(order=index, id=row.get('id'), physical_length=row.get('physical_length') or row.get('length_bp'), kind=row.get('role', 'insert'), name=row['name'], plate=row.get('plate_id') or '', well=row.get('well') or '',
                         stock=row.get('concentration_ng_ul'), resuspend=resuspend, use=row.get('volume_ul'),
                         yield_ng=source.get('yield_ng', row.get('available_ng')), fmol=row.get('target_fmol'), required_ng=row.get('target_ng'), basis=row.get('concentration_basis', ''),
                         checks=_checks(row.get('warnings', []))))
    for name, field in [('Master mix + reagents', 'master_mix_ul'), ('Water', 'water_ul'), ('Total', 'reaction_volume_ul')]:
        rows.append(dict(order='', kind='reagent' if name != 'Total' else 'total', name=name, plate='', well='', stock=None, resuspend=None, use=reaction.get(field), yield_ng=None, fmol=None, basis='', checks=[]))
    issues = _checks(assembly.get('errors', []) + reaction.get('warnings', []))
    sections.append(dict(name=assembly.get('name') or 'Assembly', stage=assembly.get('stage', 'Final assembly'), enzyme=enzyme or assembly.get('enzyme', ''),
                         ratio=settings.get('insert_ratio', 1), pmol=settings.get('backbone_pmol', .05),
                         mode=settings.get('planner_stock_mode') or 'Recorded / specified stocks', rows=rows, issues=issues,
                         valid=not issues and all(row.get('use') is not None for row in rows)))
    return sections


def build_sections(items):
    sections = []
    for item in items:
        request = item['request']
        if item.get('plan'):
            sections.extend(assembly_sections(item['plan'], request['settings'], request['enzyme']))
        else:
            sections.append(dict(name=request['name'], stage='Unavailable', enzyme=request['enzyme'], ratio=request['settings'].get('insert_ratio', 1),
                                 pmol=request['settings'].get('backbone_pmol', .05), mode='', rows=[], issues=[item.get('error') or 'Assembly unavailable'], valid=False))
    return sections


def inventory_sections(fragments, settings, name='Resuspension'):
    rows = [dict(order=i, kind='fragment', name=row['name'], plate=row.get('plate_id') or '', well=row.get('well') or '',
                 stock=row.get('target_ng_ul'), resuspend=row.get('resuspension_ul'), use=row.get('input_volume_ul'),
                 yield_ng=row.get('yield_ng'), fmol=row.get('target_fmol'), basis=row.get('stock_basis') or 'Planned stock (unmeasured)', checks=row.get('warnings', []))
            for i, row in enumerate(resuspension_plan(fragments, settings), 1)]
    return [dict(name=name, stage='Stock preparation', enzyme='', ratio=settings.get('insert_ratio', 1), pmol=settings.get('backbone_pmol', .05),
                 mode=settings.get('planner_stock_mode') or 'Planned stocks', rows=rows, issues=[], valid=True)]


HEADERS = ['#', 'Fragment / reagent', 'Plate', 'Well', 'Stock ng/µL', 'Resuspend µL', 'Use µL']


def _values(row, key):
    return [row['order'], row['name'], key.get(row['plate'], {}).get('short', ''), row['well'], row['stock'], row['resuspend'], row['use']]


def _number(value):
    return f'{value:.2f}' if isinstance(value, (int, float)) and math.isfinite(value) else '—' if value is None else str(value)


def _caption(section):
    return f"{section['enzyme']} · {section['pmol']:g} pmol · {section['ratio']:g}:1 · {section['mode']}".strip(' ·')


def excel_guide(sections, legacy=None):
    """Put a printable snapshot first; retain existing detailed/editable sheets."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.utils import get_column_letter
    wb = load_workbook(io.BytesIO(legacy)) if legacy else Workbook()
    if not legacy:
        wb.remove(wb.active)
    ws = wb.create_sheet('Bench guide', 0)
    wb.active = 0
    ws.sheet_view.showGridLines = False
    widths = [5, 39, 12, 9, 15, 15, 13]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    key = plate_key(sections)
    dark = '1D443B'
    def banner(text, color=dark, size=13):
        ws.append([_safe_cell(text)])
        n = ws.max_row
        ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=7)
        cell = ws.cell(n, 1)
        cell.font = Font(name='Aptos', size=size, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=color)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[n].height = max(27, math.ceil(len(text)/105)*19)
    for index, section in enumerate(sections):
        if index:
            ws.row_breaks.append(Break(id=ws.max_row))
        banner(section['name'] + ' / ' + section['stage'])
        for text in [_caption(section), 'Resuspend = total stock volume. Use = transfer per reaction. Blank / — = unavailable or not applicable.']:
            ws.append([_safe_cell(text)])
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=7)
            ws.row_dimensions[ws.max_row].height = 23
        ws.append(HEADERS)
        header = ws.max_row
        for cell in ws[header]:
            cell.fill = PatternFill('solid', fgColor=dark)
            cell.font = Font(name='Aptos', size=10, bold=True, color='FFFFFF')
        ws.row_dimensions[header].height = 30
        for i, row in enumerate(section['rows']):
            ws.append([_safe_cell(value) for value in _values(row, key)])
            n = ws.max_row
            for j, cell in enumerate(ws[n], 1):
                cell.font = Font(name='Aptos', size=11, bold=row['kind'] == 'total', color='173D34')
                cell.alignment = Alignment(vertical='center', wrap_text=True, horizontal='right' if j in (5, 6, 7) else 'left')
                cell.fill = PatternFill('solid', fgColor='E4EFE8' if row['kind'] == 'total' else 'F3F6F5' if i % 2 == 0 else 'FFFFFF')
                if j in (5, 6, 7):
                    cell.number_format = '0.00'
                if row['kind'] == 'total':
                    cell.border = Border(top=Side(style='thin', color=dark))
            if row['plate']:
                ws.cell(n, 3).fill = PatternFill('solid', fgColor=key[row['plate']]['color'])
            ws.row_dimensions[n].height = max(27, math.ceil(len(row['name'])/38)*17)
            if row['checks']:
                ws.append([_safe_cell('Check ' + row['name'] + ': ' + '; '.join(row['checks']))])
                ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=7)
                ws.cell(ws.max_row, 1).font = Font(name='Aptos', size=10, color='9C3A20')
                ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical='center')
                ws.row_dimensions[ws.max_row].height = max(24, math.ceil((len(row['name'])+len('; '.join(row['checks']))+8)/120)*16)
        for issue in section['issues']:
            ws.append([_safe_cell('CHECK: ' + issue)])
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=7)
            ws.cell(ws.max_row, 1).font = Font(name='Aptos', size=10, color='9C3A20')
            ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True)
            ws.row_dimensions[ws.max_row].height = max(24, math.ceil(len(issue)/120)*16)
        ws.append([])
    ws.freeze_panes = 'C5'
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = .25
    ws.page_margins.top = ws.page_margins.bottom = .35
    ws.oddFooter.center.text = 'Junction | Page &P of &N'
    ws.print_area = f'A1:G{ws.max_row}'
    legend = wb.create_sheet('Plate key', 1)
    legend.append(['Plate', 'Full plate identifier', 'Color'])
    for entry in key.values():
        legend.append([_safe_cell(entry['short']), _safe_cell(entry['name']), '#'+entry['color']])
        legend.cell(legend.max_row, 1).fill = PatternFill('solid', fgColor=entry['color'])
    for col, width in zip('ABC', [16, 65, 15]):
        legend.column_dimensions[col].width = width
    legend.sheet_view.showGridLines = False
    for cell in legend[1]:
        cell.fill = PatternFill('solid', fgColor=dark)
        cell.font = Font(bold=True, color='FFFFFF')
    for row in legend:
        legend.row_dimensions[row[0].row].height = 25
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='center')
    legend.freeze_panes = 'A2'
    legend.page_setup.orientation = 'landscape'
    legend.page_setup.paperSize = legend.PAPERSIZE_A4
    legend.sheet_properties.pageSetUpPr.fitToPage = True
    legend.page_setup.fitToWidth = 1
    legend.page_setup.fitToHeight = 0
    legend.print_title_rows = '1:1'
    legend.print_area = f'A1:C{legend.max_row}'
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def csv_guide(sections):
    """One rectangular table, with explicit full IDs/color codes for plain CSV."""
    key = plate_key(sections)
    output = io.StringIO(newline='')
    writer = csv.writer(output)
    writer.writerow(['Assembly', 'Stage', 'Enzyme', 'Ratio', 'DNA pmol', 'Order', 'Type', 'Fragment / reagent', 'Plate', 'Well', 'Stock ng/µL', 'Resuspend µL', 'Use µL', 'Yield ng', 'Required fmol', 'Full plate ID', 'Plate color', 'Stock basis', 'Checks'])
    for section in sections:
        for row in section['rows'] or [dict(order='', kind='unavailable', name='', plate='', well='', stock=None, resuspend=None, use=None, yield_ng=None, fmol=None, basis='', checks=[])]:
            plate = key.get(row['plate'], {})
            values = [section['name'], section['stage'], section['enzyme'], f"{section['ratio']:g}:1", section['pmol'], row['order'], row['kind'], row['name'], plate.get('short', ''), row['well'], row['stock'], row['resuspend'], row['use'], row['yield_ng'], row['fmol'], row['plate'], '#'+plate['color'] if plate else '', row['basis'], '; '.join(_checks(row['checks']+section['issues']))]
            writer.writerow([_safe_cell(f'{value:.2f}' if isinstance(value, float) else value) for value in values])
    return ('\ufeff'+output.getvalue()).encode('utf-8')


def pdf_guide(sections):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
    key = plate_key(sections)
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=28, leftMargin=28, topMargin=27, bottomMargin=28, title='Junction assembly guide', author='Junction')
    styles = {
        'title': ParagraphStyle('title', fontName='Helvetica-Bold', fontSize=16, leading=20, textColor=colors.HexColor('#1D443B')),
        'small': ParagraphStyle('small', fontSize=9, leading=12, textColor=colors.HexColor('#40574E')),
        'cell': ParagraphStyle('cell', fontSize=10, leading=13),
        'num': ParagraphStyle('num', fontSize=10, leading=13, alignment=TA_RIGHT),
        'header': ParagraphStyle('header', fontName='Helvetica-Bold', fontSize=9, leading=12, textColor=colors.white),
        'check': ParagraphStyle('check', fontSize=9, leading=12, textColor=colors.HexColor('#9C3A20')),
    }
    def p(value, style='cell'):
        return Paragraph(escape(str(value)).replace('\n', '<br/>'), styles[style])
    flow = []
    for index, section in enumerate(sections):
        if index:
            flow.append(PageBreak())
        flow.append(KeepTogether([p(section['name']+' / '+section['stage'], 'title'), Spacer(1, 7), p(_caption(section), 'small'), Spacer(1, 4),
                                 p('Resuspend = total stock volume. Use = transfer per reaction. — = unavailable or not applicable.', 'small'), Spacer(1, 12)]))
        data = [[p(h, 'header') for h in HEADERS]]
        table_styles = [('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1D443B')), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 7), ('RIGHTPADDING', (0, 0), (-1, -1), 7), ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7)]
        for i, row in enumerate(section['rows'], 1):
            values = _values(row, key)
            data.append([p(_number(v) if j >= 4 else str(v), 'num' if j >= 4 else 'cell') for j, v in enumerate(values)])
            table_styles.append(('BACKGROUND', (0, i), (-1, i), colors.HexColor('#E4EFE8' if row['kind'] == 'total' else '#F3F6F5' if i % 2 else '#FFFFFF')))
            if row['plate']:
                table_styles.append(('BACKGROUND', (2, i), (2, i), colors.HexColor('#'+key[row['plate']]['color'])))
            if row['kind'] == 'total':
                table_styles.append(('LINEABOVE', (0, i), (-1, i), .7, colors.HexColor('#1D443B')))
        table = Table(data, colWidths=[27, 263, 77, 53, 94, 98, 73], repeatRows=1, hAlign='LEFT')
        table.setStyle(TableStyle(table_styles))
        flow.append(table)
        issues = section['issues'] + [row['name']+': '+'; '.join(row['checks']) for row in section['rows'] if row['checks']]
        for issue in _checks(issues):
            flow.extend([Spacer(1, 5), p('CHECK: '+issue, 'check')])
    if key:
        flow.extend([PageBreak(), p('Plate key', 'title'), Spacer(1, 12)])
        data = [[p('Plate', 'header'), p('Full plate identifier', 'header')]]
        style = [('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1D443B')), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('TOPPADDING', (0, 0), (-1, -1), 8), ('BOTTOMPADDING', (0, 0), (-1, -1), 8)]
        for i, entry in enumerate(key.values(), 1):
            data.append([p(entry['short']), p(entry['name'])])
            style.append(('BACKGROUND', (0, i), (0, i), colors.HexColor('#'+entry['color'])))
        table = Table(data, colWidths=[100, 685], repeatRows=1)
        table.setStyle(TableStyle(style))
        flow.append(table)
    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.HexColor('#40574E'))
        canvas.drawString(28, 14, 'Junction | Export snapshot')
        canvas.drawRightString(landscape(A4)[0]-28, 14, f'Page {document.page}')
        canvas.restoreState()
    doc.build(flow, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def guide_bytes(format, sections, legacy_factory=None):
    if any(section['stage']!='Stock preparation' for section in sections):
        from .assembly_matrix import matrix_bytes
        return matrix_bytes(format,sections,legacy_factory)
    if format == 'xlsx':
        return excel_guide(sections, legacy_factory() if legacy_factory else None)
    if format == 'pdf':
        return pdf_guide(sections)
    if format == 'csv':
        return csv_guide(sections)
    raise ValueError('Unsupported export format')
