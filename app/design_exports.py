"""Synthesis records keep target DNA separate from complete adapter-flanked DNA."""
import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def fragment_rows(result):
    rows = [dict(zip([
        'Name', 'Assembly row', 'Region', 'Target sequence', 'Full sequence',
        'Target length nt', 'Full length nt', 'GC %', '5′ overhang', '3′ overhang',
        '5′ RE', '3′ RE', '5′ recognition site', '3′ recognition site',
        '5′ complete end', '3′ complete end', 'RE sites included',
        'Reference start', 'Reference end', 'Output orientation'], [
        row['name'], row['row_name'], row['region_mode'], row['target_sequence'], row['full_sequence'],
        row['target_length'], row['length'], round(row['gc'], 2), row['left_overhang'], row['right_overhang'],
        row['left_adapter_enzyme'] if row['include_re_sites'] else '',
        row['right_adapter_enzyme'] if row['include_re_sites'] else '',
        row['left_re_site'], row['right_re_site'], row['five_prime_sequence'], row['three_prime_sequence'],
        row['include_re_sites'], row['source_start'] + 1, row['source_end'], row.get('output_orientation','Forward')])) for row in result['generated']]
    if result.get('make_twist_compatible'):
        for exported, row in zip(rows, result['generated']):
            exported.update({'Unpadded full sequence': row['unpadded_full_sequence'],
                             'Unpadded full length nt': row['unpadded_length'],
                             row.get('padding_end','3′')+' padding sequence': row['synthesis_padding'],
                             row.get('padding_end','3′')+' padding length nt': row['padding_length']})
    return rows


def assembly_rows(result):
    return [{'Name': row['name'], 'Length bp': row['length'], 'Circular': row['circular'],
             'GC %': round(row['gc'], 2), 'Target sequence': row['sequence'],
             'Output orientation': row.get('output_orientation','Forward'),
             'Fragment order': ' → '.join(feature['label'] for feature in row['features']
                                          if feature.get('source') in ('fragment','retained_backbone'))}
            for row in result.get('assemblies', [])]


def literal(value):
    return "'" + value if isinstance(value,str) and value.startswith(('=','+','-','@')) else value


def csv_bytes(rows):
    output = io.StringIO()
    if rows:
        writer = csv.writer(output)
        writer.writerow(rows[0].keys())
        writer.writerows([literal(value) for value in row.values()] for row in rows)
    return output.getvalue().encode('utf-8-sig')


def workbook_bytes(result):
    workbook = Workbook()
    workbook.remove(workbook.active)
    overhangs = [{'Region': row['name'], '5′ overhang': row['left_overhang'],
                 '3′ overhang': row['right_overhang'], 'Constant across rows': row['valid'],
                 'Output orientation': row.get('output_orientation','Forward')}
                for row in result.get('constant_overhangs', [])]
    for title, rows in [('Fragments',fragment_rows(result)),('Assemblies',assembly_rows(result)),('Overhangs',overhangs)]:
        sheet = workbook.create_sheet(title)
        if not rows:
            continue
        sheet.append(list(rows[0]))
        for row in rows:
            if any(isinstance(value,str) and len(value)>32767 for value in row.values()):
                raise ValueError('A sequence exceeds Excel’s cell limit. Export CSV or FASTA for complete sequences.')
            sheet.append([literal(value) for value in row.values()])
        for cell in sheet[1]:
            cell.font = Font(bold=True,color='FFFFFF')
            cell.fill = PatternFill('solid',fgColor='176C58')
            cell.alignment = Alignment(wrap_text=True,vertical='center')
        for column in sheet.columns:
            header = str(column[0].value)
            sheet.column_dimensions[column[0].column_letter].width = 60 if 'sequence' in header.lower() else 24 if header in ('Name','Assembly row','Fragment order') else 18
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True,vertical='top')
                if isinstance(cell.value,float):
                    cell.number_format = '0.00'
            sheet.row_dimensions[row[0].row].height = 48
        sheet.row_dimensions[1].height = 32
        sheet.freeze_panes = 'D2' if title == 'Fragments' else 'B2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
