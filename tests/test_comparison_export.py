import io
import unittest
from openpyxl import load_workbook
from app.comparison_export import comparison_workbook


def record(filename,exact=True,identity=1):
    return {'filename':filename,'query_name':'Query','query_length':100,'query_circular':True,
            'matches':[{'name':'Reference','source':'saved plasmid','reference_length':100,'identity':identity,'exact':exact,'orientation':'+','rotation':0,
                        'changes':[] if exact else [{'type':'substitution','reference_start':9,'reference_end':10,'reference_bases':'A','query_bases':'G','original_query_spans':[[9,10]]}]}],
            'fragment_matches':[{'name':'Fragment','plate_id':'Plate','well':'A1','representation':'Payload','query_spans':[[0,50]],'strand':'+','identity':1,'coverage':1,'exact':True,'status':'Exact'}]}


class ComparisonExportTests(unittest.TestCase):
    def test_exact_export_excludes_rounded_near_matches_and_error_records(self):
        rows=[record('exact.gb'),record('near.gb',False,.9999999),record('flag-but-not-exact.gb',True,.9999999),{'filename':'error.gb','error':'Bad record'}]
        wb=load_workbook(io.BytesIO(comparison_workbook(rows,'exact')))
        sheet=wb['Results'];self.assertEqual(sheet.max_row,2);self.assertEqual(sheet['A2'].value,'exact.gb');self.assertEqual(sheet['I2'].value,'Yes')
        self.assertEqual(wb['Fragments'].max_row,2)
        self.assertEqual(sheet['I2'].fill.fgColor.rgb,'00E9F3EE')

    def test_selected_export_preserves_nonexact_results_and_difference_coordinates(self):
        rows=[record('exact.gb'),record('mutated.gb',False,.99)]
        wb=load_workbook(io.BytesIO(comparison_workbook(rows,'selected')))
        self.assertEqual(wb['Results'].max_row,3);self.assertEqual(wb['Results']['I3'].value,'No')
        self.assertEqual(wb['Results']['H3'].value,99)
        self.assertEqual(wb['Results']['L3'].value,1)
        self.assertEqual(wb['Differences']['E2'].value,'10–10')
        self.assertEqual(wb['Differences']['F2'].value,'10–10')
        self.assertEqual(wb['Fragments']['E2'].value,'Plate')
        self.assertEqual(wb['Fragments']['F2'].value,'A1')

    def test_names_are_excel_literals_and_no_hits_export_with_clear_status(self):
        data=record('=danger.gb');data['query_name']='@query'
        wb=load_workbook(io.BytesIO(comparison_workbook([data,{'filename':'no-hit.fa','query_name':'No hit','matches':[]}],'selected')))
        self.assertEqual(wb['Results']['A2'].data_type,'s');self.assertEqual(wb['Results']['A2'].value,"'=danger.gb")
        self.assertEqual(wb['Results']['S3'].value,'No assembly match')
        with self.assertRaises(ValueError):comparison_workbook([record('near',False,.99)],'exact')
