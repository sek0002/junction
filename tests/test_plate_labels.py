"""The same picking labels appear in the app and every tabular export."""
import io
import unittest
from openpyxl import load_workbook
from app.plate_labels import plate_label, plate_aliases
from app.exports import make_workbook
from app.comparison_export import comparison_workbook
from app.plate_exports import snapshot,export_bytes
from test_plate_exports import example_plate
from test_comparison_export import record


class PlateLabelTests(unittest.TestCase):
    def test_prefix_alphanumeric_and_non_twist_ids(self):
        self.assertEqual(plate_label('pSHPs0405B231201IQ'),'0405')
        self.assertEqual(plate_label('pSHPs0622B228801IQ'),'0622')
        self.assertEqual(plate_label('PSHPS04-05B231201IQ'),'0405')
        self.assertEqual(plate_label('Custom Plate'),'Custom Plate')
        self.assertEqual(plate_label(''),'')

    def test_collisions_are_unique_and_order_independent(self):
        ids=['pSHPs0921B1498027B','pSHPs0405B231201IQ','pSHPs0921B1498017B','0921-1']
        aliases=plate_aliases(ids)
        self.assertEqual(aliases,plate_aliases(reversed(ids)))
        self.assertEqual(aliases['pSHPs0405B231201IQ'],'0405')
        self.assertEqual(len(set(aliases.values())),4)
        self.assertTrue(aliases['pSHPs0921B1498017B'].startswith('0921-'))

    def test_legacy_calculation_sheets_use_aliases_and_keep_full_id_comments(self):
        identifier='pSHPs0405B231201IQ'
        part=dict(id=1,name='Fragment',plate_id=identifier,well='A1',sequence='ACGT'*100,yield_ng=100)
        wb=load_workbook(io.BytesIO(make_workbook([part])))
        for sheet,cell in [('Resuspension inventory','C2'),('Resuspension plan','B3'),('Fragment selection','B4')]:
            self.assertEqual(wb[sheet][cell].value,'0405')
            self.assertEqual(wb[sheet][cell].comment.text,identifier)
        self.assertEqual(wb['Resuspension inventory']['H2'].data_type,'f')

    def test_comparison_export_uses_same_alias_with_source_id_comment(self):
        data=record('exact.gb');data['fragment_matches'][0]['plate_id']='pSHPs0622B228801IQ'
        wb=load_workbook(io.BytesIO(comparison_workbook([data],'exact')))
        self.assertEqual(wb['Fragments']['E2'].value,'0622')
        self.assertEqual(wb['Fragments']['E2'].comment.text,'pSHPs0622B228801IQ')

    def test_plate_well_exports_share_label_and_keep_full_title(self):
        plate=example_plate();plate['plate_id']='pSHPs0405B231201IQ'
        self.assertEqual(snapshot(plate)['rows'][0]['values'][0],'0405')
        wb=load_workbook(io.BytesIO(export_bytes(plate,'xlsx')))
        self.assertEqual(wb['Wells']['A6'].value,'0405')
        self.assertEqual(wb['Plate map']['A1'].value,'0405 · pSHPs0405B231201IQ')
