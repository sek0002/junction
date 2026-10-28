"""RC is a coherent output strand, including end metadata and GenBank features."""
from copy import deepcopy
import csv
import io
import json
import os
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from auth_client import AuthenticatedClient as TestClient
from app.main import app
from app.design_engine import design
from app.design_table import sequence_metadata
from app.design_output import output_view, reverse_features
from app.sequence_engine import digest_fragment, reverse_complement as rc
from app.exports import parse_sequence_file
from test_design import clean_dna, placeholder


class OutputTests(unittest.TestCase):
    def test_mixed_end_metadata_and_exact_full_molecules_reverse_together(self):
        result=sequence_metadata(design([placeholder(clean_dna())],'BsmBI + BsaI',left_overhang='TTTG',right_overhang='ACCT',left_adapter_enzyme='BsaI',right_adapter_enzyme='BsmBI'))
        self.assertTrue(result['valid'],result['errors']);before=deepcopy(result)
        out=output_view(result,True)
        for original,row in zip(result['generated'],out['generated']):
            self.assertEqual(row['name'],original['export_name'])
            self.assertEqual(row['full_sequence'],rc(original['full_sequence']))
            self.assertEqual(row['target_sequence'],rc(original['target_sequence']))
            self.assertEqual(row['left_overhang'],rc(original['right_overhang']))
            self.assertEqual(row['right_overhang'],rc(original['left_overhang']))
            self.assertEqual(row['left_adapter_enzyme'],'BsmBI');self.assertEqual(row['right_adapter_enzyme'],'BsaI')
            self.assertEqual(row['five_prime_sequence'],rc(original['three_prime_sequence']))
            self.assertEqual(row['five_prime_sequence']+row['target_sequence']+row['three_prime_sequence'],row['full_sequence'])
            parsed=digest_fragment(row['full_sequence'],'BsmBI + BsaI')
            self.assertTrue(parsed['valid'],parsed['errors']);self.assertEqual(parsed['target'],row['target_sequence'])
            self.assertEqual(row['gc'],original['gc']);self.assertEqual(row['source_start'],original['source_start'])
        self.assertEqual(out['assembly']['sequence'],rc(result['assembly']['sequence']))
        self.assertEqual(out['assemblies'][0]['sequence'],rc(result['assemblies'][0]['sequence']))
        self.assertEqual(result,before)

    def test_feature_coordinates_strands_and_compound_segments_remap_without_losing_biological_order(self):
        features=[dict(label='Gene',start=10,end=30,strand=1,type='CDS'),dict(label='Joined',start=2,end=8,strand=1,segments=[dict(start=2,end=8,strand=1),dict(start=50,end=60,strand=1)]),dict(label='ACCT fusion: A → B',source='fusion',start=40,end=44,strand=1)]
        mapped=reverse_features(features,100);by_name={f['label']:f for f in mapped}
        self.assertEqual((by_name['Gene']['start'],by_name['Gene']['end'],by_name['Gene']['strand']),(70,90,-1))
        self.assertEqual(by_name['Joined']['segments'],[dict(start=92,end=98,strand=-1),dict(start=40,end=50,strand=-1)])
        self.assertIn('AGGT fusion: B → A',by_name)

    def test_adapter_free_output_remains_exactly_overhang_target_overhang(self):
        original=sequence_metadata(design([placeholder(clean_dna())],left_overhang='TTTG',right_overhang='ACCT'),False)
        out=output_view(original,True)
        self.assertEqual(output_view(original)['generated'][0]['name'],original['generated'][0]['export_name'])
        row=out['generated'][0];self.assertEqual(row['full_sequence'],row['left_overhang']+row['target_sequence']+row['right_overhang'])
        self.assertEqual(row['left_re_site'],'');self.assertEqual(row['right_re_site'],'')


class OutputAPITests(unittest.TestCase):
    def setUp(self):
        self.folder=tempfile.TemporaryDirectory();self.old=os.environ.get('GOLDEN_GATE_DB')
        path=Path(self.folder.name)/'test.sqlite3';os.environ['GOLDEN_GATE_DB']=str(path)
        path.with_suffix('.sources.json').write_text(json.dumps({'plate_files':[],'genbank_files':[]}))
        self.client=TestClient(app);self.client.__enter__()
        self.payload=dict(name='Output test',enzyme='BsaI',entries=[placeholder(clean_dna(),[dict(name='Cargo',start=100,end=260,mode='dynamic')])],left_overhang='TTTG',right_overhang='ACCT')

    def tearDown(self):
        self.client.__exit__(None,None,None)
        if self.old is None:os.environ.pop('GOLDEN_GATE_DB',None)
        else:os.environ['GOLDEN_GATE_DB']=self.old
        self.folder.cleanup()

    def test_all_export_formats_reverse_without_changing_preview(self):
        payload=self.payload|dict(reverse_output=True)
        preview=self.client.post('/api/design/preview?generate=true',json=payload).json()
        baseline=self.client.post('/api/design/preview?generate=true',json=self.payload).json()
        self.assertEqual(preview['assembly']['sequence'],baseline['assembly']['sequence'])
        for format in ('csv','xlsx','fasta','target_fasta','assemblies_csv','gb','example_gb'):
            response=self.client.post('/api/design/export?format='+format,json=payload)
            self.assertEqual(response.status_code,200,response.text[:100] if format!='xlsx' else '')
            self.assertIn('_RC.',response.headers['Content-Disposition'])
            if format=='csv':
                rows=list(csv.DictReader(io.StringIO(response.content.decode('utf-8-sig'))))
                self.assertEqual(rows[0]['Full sequence'],rc(preview['generated'][0]['full_sequence']))
                self.assertEqual(rows[0]['Output orientation'],'Reverse complement')
                self.assertEqual(rows[0]['5′ overhang'],rc(preview['generated'][0]['right_overhang']))
            elif format=='xlsx':
                wb=load_workbook(io.BytesIO(response.content));self.assertEqual(wb['Fragments'].cell(2,5).value,rc(preview['generated'][0]['full_sequence']))
            elif format in ('fasta','target_fasta'):
                key='full_sequence' if format=='fasta' else 'target_sequence'
                records=response.text.split('>')[1:]
                self.assertTrue(all(record.splitlines()[0].endswith('|RC') for record in records))
                self.assertEqual([''.join(record.splitlines()[1:]) for record in records],[rc(row[key]) for row in preview['generated']])
            elif format=='assemblies_csv':
                rows=list(csv.DictReader(io.StringIO(response.content.decode('utf-8-sig'))))
                self.assertEqual(rows[0]['Target sequence'],rc(preview['assembly']['sequence']))
            else:
                record=parse_sequence_file(response.text)[0];self.assertEqual(record['sequence'],rc(preview['assembly']['sequence']))
                if format=='example_gb':
                    region=next(feature for feature in record['features'] if feature['label']=='Cargo')
                    self.assertEqual(region['strand'],-1)
                    self.assertEqual(rc(record['sequence'][region['start']:region['end']]),clean_dna()[100:260])
        self.assertEqual(self.client.post('/api/design/preview?generate=true',json=payload).json()['assembly']['sequence'],preview['assembly']['sequence'])


if __name__=='__main__':unittest.main()
