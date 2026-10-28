"""HTTP integration checks using isolated SQLite databases and real examples."""
from __future__ import annotations

import csv
import io
import json
import os
from pathlib import Path
import tempfile
import unittest

from auth_client import AuthenticatedClient as TestClient
from openpyxl import load_workbook

from app.main import app
from app.exports import parse_sequence_file

APP_ROOT = Path(__file__).resolve().parents[1]


def reverse_complement(sequence):
    return sequence.translate(str.maketrans('ACGT', 'TGCA'))[::-1]


class WorkbenchIntegrationTests(unittest.TestCase):
    def test_comparison_export_filters_exact_results_and_selected_near_matches(self):
        records=[{'filename':'exact.gb','query_name':'Exact','matches':[{'name':'Reference','identity':1,'exact':True,'reference_coverage':1,'query_coverage':1,'aligned_bp':100}]},
                 {'filename':'near.gb','query_name':'Near','matches':[{'name':'Reference','identity':.9999,'exact':False}]}]
        response=self.client.post('/api/compare/export',json={'scope':'exact','records':records})
        self.assertEqual(response.status_code,200,response.text if response.status_code!=200 else '')
        wb=load_workbook(io.BytesIO(response.content));self.assertEqual(wb['Results'].max_row,2)
        self.assertEqual(wb['Results']['A2'].value,'exact.gb')
        self.assertEqual(wb['Results']['T2'].value,100)
        self.assertEqual(wb['Results']['U2'].value,100)
        self.assertEqual(wb['Results']['V2'].value,100)
        self.assertIn('comparison-exact.xlsx',response.headers['content-disposition'])
        selected=self.client.post('/api/compare/export',json={'scope':'selected','records':records[1:]})
        self.assertEqual(selected.status_code,200)
        self.assertEqual(load_workbook(io.BytesIO(selected.content))['Results']['I2'].value,'No')
        self.assertEqual(self.client.post('/api/compare/export',json={'scope':'exact','records':records[1:]}).status_code,422)

    def test_advanced_compare_is_unavailable(self):
        self.assertEqual(self.client.get('/api/compare/advanced/status').status_code,404)
        self.assertEqual(self.client.post('/api/compare/advanced',json={}).status_code,404)
        self.assertFalse(any('/compare/advanced' in path for path in self.client.get('/openapi.json').json()['paths']))

    def test_backbone_detection_flags_are_independent_of_selected_enzyme(self):
        for enzyme in ('BsaI','BsmBI','BsmBI + BsaI'):
            response=self.client.get('/api/plasmids',params={'enzyme':enzyme})
            self.assertEqual(response.status_code,200,response.text)
            example=next(p for p in response.json() if p['name']=='pIVT_V2')
            self.assertEqual(example['backbone_options']['site_counts'],{'BsmBI':2,'BsaI':0})
        from app.main import backbone_template_info
        self.assertEqual(backbone_template_info('ACGTACGT',True)['site_counts'],{'BsmBI':0,'BsaI':0})
        both=backbone_template_info('GGTCTCAAACGTCTCAAAA',True)
        self.assertEqual(both['site_counts'],{'BsmBI':1,'BsaI':1})
        self.assertFalse(both['eligible'])

    def test_dual_enzyme_across_search_build_queue_export_and_compare(self):
        enzyme='BsmBI + BsaI'
        first='GCGTCGTCTCATTTGAACCGGATGGCGAGACCACGCA'
        second='GCGTGGTCTCAATGGCCAAAATTTGCGAGACGACGCA'
        content='Name,Plate ID,Well,Sequence,Yield ng\nDual left,dual,A1,'+first+',300\nDual right,dual,A2,'+second+',300\n'
        uploaded=self.client.post('/api/uploads/plate?no_backbone=true',files={'file':('dual.csv',content,'text/csv')})
        self.assertEqual(uploaded.status_code,200,uploaded.text)
        parts=self.client.get('/api/fragments',params={'enzyme':enzyme,'plate':'dual'}).json()
        self.assertTrue(all(p['digest']['valid'] for p in parts))
        self.assertEqual(len(parts),2)
        ids=[p['id'] for p in parts]
        self.assertTrue(self.client.get('/api/fragments/'+str(ids[0]),params={'enzyme':enzyme}).json()['digest']['valid'])
        self.assertEqual(self.client.get('/api/plasmids',params={'enzyme':enzyme}).status_code,200)
        payload={'name':'Dual assembly','enzyme':enzyme,'fragment_ids':ids}
        result=self.client.post('/api/assemblies/preview',json=payload)
        self.assertEqual(result.status_code,200,result.text)
        self.assertTrue(result.json()['valid'],result.json()['errors'])
        queue=self.client.post('/api/build-list',json=payload)
        self.assertEqual(queue.status_code,200,queue.text)
        self.assertEqual(queue.json()['request']['enzyme'],enzyme)
        genbank=self.client.post('/api/assemblies/genbank',json=payload)
        self.assertEqual(genbank.status_code,200,genbank.text)
        self.assertIn(enzyme,genbank.text)
        worksheet=self.client.post('/api/assemblies/worksheet',json=payload)
        self.assertEqual(worksheet.status_code,200)
        book=load_workbook(io.BytesIO(worksheet.content))
        chosen=[row for row in book['Fragment selection'].iter_rows(min_row=4,values_only=True) if row[0]=='Dual left']
        self.assertEqual(chosen[0][3],enzyme)
        compared=self.client.post('/api/compare',params={'enzyme':enzyme},files={'file':('dual.fa','>dual\n'+result.json()['sequence'],'text/plain')})
        self.assertEqual(compared.status_code,200,compared.text)
        self.assertGreater(compared.json()['inferred_count'],0)
        self.assertTrue(compared.json()['matches'][0]['exact'])

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.previous_database = os.environ.get('GOLDEN_GATE_DB')
        self.database = str(Path(self.directory.name) / 'isolated.sqlite3')
        os.environ['GOLDEN_GATE_DB'] = self.database
        # Pin the original two-plate regression fixture independently of shipped defaults.
        sources=Path(self.directory.name)/'fixture_sources'
        sources.mkdir()
        plates=[APP_ROOT/'examples'/name for name in ('platemap_pSHPs0921B1498017B.csv','platemap_pSHPs0921B1498027B.csv')]
        genbanks=sorted((APP_ROOT/'examples').glob('*.gb'))
        for path in plates+genbanks:(sources/path.name).write_bytes(path.read_bytes())
        Path(self.database).with_suffix('.sources.json').write_text(json.dumps({
            'plate_files':['fixture_sources/'+path.name for path in plates],
            'genbank_files':['fixture_sources/'+path.name for path in genbanks]}))
        self.client = TestClient(app)
        self.client.__enter__()

    def tearDown(self):
        self.client.__exit__(None, None, None)
        if self.previous_database is None:
            os.environ.pop('GOLDEN_GATE_DB', None)
        else:
            os.environ['GOLDEN_GATE_DB'] = self.previous_database
        self.directory.cleanup()

    def recipes(self):
        response = self.client.get('/api/recipes')
        self.assertEqual(response.status_code, 200)
        return response.json()

    def payload(self, recipe):
        return {key: value for key, value in recipe.items() if key not in {'available', 'missing'}}

    def preview(self, name='pHR_V2 · U6'):
        recipe = next(recipe for recipe in self.recipes() if recipe['name'] == name)
        response = self.client.post('/api/assemblies/preview', json=self.payload(recipe))
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_seed_examples_download_and_all_available_recipes(self):
        summary = self.client.get('/api/summary').json()
        self.assertEqual((summary['fragments'], summary['plasmids'], summary['plates']), (102, 7, 2))
        backbones = [part for part in self.client.get('/api/plasmids').json() if part['recipe'].get('example_backbone')]
        self.assertEqual({part['name'] for part in backbones}, {'pIVT_V2', 'pULT_V2', 'pHR_V2'})
        self.assertEqual(len(backbones), 3)
        self.assertTrue(all(part['circular'] for part in backbones))
        legacy = [part for part in self.client.get('/api/plasmids').json() if part['name'].startswith(('PHR_VE', 'PHR_U6'))]
        self.assertEqual(len(legacy), 4)
        self.assertTrue(all(part['recipe'].get('legacy_example') and not part['recipe'].get('example_backbone') for part in legacy))
        examples = self.client.get('/api/examples').json()
        self.assertEqual(sum(item['kind'] == 'plate' for item in examples), 2)
        self.assertEqual(sum(item['kind'] == 'GenBank' for item in examples), 4)
        for item in examples:
            downloaded = self.client.get('/api/examples/' + item['filename'])
            self.assertEqual(downloaded.status_code, 200)
            self.assertEqual(downloaded.content, (APP_ROOT / 'examples' / item['filename']).read_bytes())
        lengths = []
        for recipe in self.recipes():
            if not recipe['available']:
                self.assertTrue(recipe['missing'])
                continue
            result = self.client.post('/api/assemblies/preview', json=self.payload(recipe))
            self.assertEqual(result.status_code, 200, result.text)
            assembly = result.json()
            self.assertTrue(assembly['valid'], assembly['errors'])
            self.assertTrue(all(j['compatible'] for j in assembly['junctions']))
            self.assertTrue(assembly['junctions'][-1]['circular_closure'])
            self.assertFalse([site for site in assembly['restriction_sites'] if site['enzyme'] == 'BsaI'])
            lengths.append(assembly['length'])
        self.assertEqual(sorted(lengths), [2827, 7457, 7760, 8201])

    def test_fragment_search_matches_names_without_matching_plate_or_well(self):
        from app import database as store
        plate='location-token-927'
        with store.connection() as db:
            for name,well in [('F6 donor','A1'),('Promoter','F6'),('F6 variant','B1')]:
                db.execute('INSERT INTO fragments(name,plate_id,well,sequence) VALUES(?,?,?,?)',
                    (name,plate,well,'ACGT'*75))
        results=self.client.get('/api/fragments',params={'q':' f6 ','plate':plate}).json()
        self.assertEqual({row['name'] for row in results},{'F6 donor','F6 variant'})
        for query in ('A1',plate):
            self.assertEqual(self.client.get('/api/fragments',params={'q':query,'plate':plate}).json(),[])
        typo=self.client.get('/api/fragments',params={'q':'promter','plate':plate}).json()
        self.assertEqual([row['name'] for row in typo],['Promoter'])
        self.assertEqual(len(self.client.get('/api/fragments',params={'plate':plate}).json()),3)

    def test_upload_rules_preview_import_and_library_agree_without_changing_defaults(self):
        defaults=self.client.get('/api/inventory/rules').json()['rules']
        rules=defaults | {'backbone_pmol':.1,'insert_ratio':1,'preferred_resuspension_ul':60,'max_resuspension_ul':90,'min_input_volume_ul':2,'max_input_volume_ul':4}
        params=rules | {'apply_import_rules':True,'no_backbone':True}
        content='Name,Plate ID,Well,Sequence,Yield ng,Stock ng ul,Resuspension ul\nImport plan,import-flags,A1,'+'ATAC'*125+',600,,\nProvided stock,import-flags,A2,'+'ATAC'*150+',600,25,24\n'
        def upload(preview,values=params):
            return self.client.post('/api/uploads/plate',params=values | {'preview':preview},files={'file':('rules.csv',content,'text/csv')})
        count=self.client.get('/api/summary').json()['fragments']
        preview=upload(True)
        self.assertEqual(preview.status_code,200,preview.text)
        self.assertEqual(self.client.get('/api/summary').json()['fragments'],count)
        planned=preview.json()['resuspension']['rows'][0]
        self.assertEqual(planned['resuspension_ul'],60)
        self.assertAlmostEqual(planned['input_volume_ul'],3.3)
        committed=upload(False)
        self.assertEqual(committed.status_code,200,committed.text)
        self.assertEqual(committed.json()['count'],2)
        self.assertEqual(self.client.get('/api/inventory/rules').json()['rules'],defaults)
        fragments=self.client.get('/api/fragments',params={'plate':'import-flags'}).json()
        by_name={part['name']:part for part in fragments}
        records={row['id']:row for row in self.client.get('/api/inventory/records').json()['rows']}
        imported=records[by_name['Import plan']['id']]
        self.assertEqual(imported['applied_rules'],rules)
        self.assertFalse(imported['saved'])
        self.assertIn('unmeasured',imported['concentration_basis'])
        for key in ('yield_ng','resuspension_ul','input_volume_ul'):
            self.assertAlmostEqual(imported[key],planned[key])
        self.assertAlmostEqual(imported['stock_ng_ul'],planned['target_ng_ul'])
        provided=records[by_name['Provided stock']['id']]
        self.assertEqual((provided['stock_ng_ul'],provided['resuspension_ul']),(25,24))
        repeated=upload(False,params | {'preferred_resuspension_ul':80})
        self.assertEqual(repeated.status_code,200,repeated.text)
        self.assertEqual(repeated.json()['count'],0)
        with TestClient(app) as reopened:
            record=reopened.get('/api/fragments/'+str(by_name['Import plan']['id'])).json()
            self.assertEqual(record['stock_rules'],rules)

    def test_upload_rule_limits_reject_before_import(self):
        content='Name,Plate ID,Well,Sequence,Yield ng\nRejected,invalid-flags,A1,'+'ATAC'*125+',600\n'
        count=self.client.get('/api/summary').json()['fragments']
        for bad in ({'max_resuspension_ul':0},{'preferred_resuspension_ul':-1},{'min_input_volume_ul':0},{'max_input_volume_ul':-6},{'min_input_volume_ul':4,'max_input_volume_ul':2}):
            response=self.client.post('/api/uploads/plate',params={'apply_import_rules':True,'no_backbone':True} | bad,files={'file':('bad.csv',content,'text/csv')})
            self.assertEqual(response.status_code,422,response.text)
            self.assertEqual(self.client.get('/api/summary').json()['fragments'],count)

    def test_unrestricted_rules_save_apply_import_export_and_restart(self):
        rules={'backbone_pmol':.05,'insert_ratio':1,'preferred_resuspension_ul':15000,
            'max_resuspension_ul':20000,'min_input_volume_ul':.25,'max_input_volume_ul':20001}
        saved=self.client.put('/api/inventory/rules',json=rules)
        self.assertEqual(saved.status_code,200,saved.text)
        self.assertEqual(saved.json()['rules'],rules)
        self.assertEqual(self.client.get('/api/inventory/rules').json()['rules'],rules)
        for key in ('backbone_pmol','insert_ratio'):
            response=self.client.put('/api/inventory/rules',json=rules | {key:1000})
            self.assertEqual(response.status_code,200,response.text)
        self.client.put('/api/inventory/rules',json=rules)
        applied=self.client.post('/api/inventory/apply-rules',json={'fragment_ids':[1],'rules':rules})
        self.assertEqual(applied.status_code,200,applied.text)
        content='Name,Plate ID,Well,Sequence,Yield ng\nUnrestricted,custom-wide,A1,'+'ATAC'*125+',600\n'
        uploaded=self.client.post('/api/uploads/plate',params=rules | {'apply_import_rules':True,'no_backbone':True},files={'file':('wide.csv',content,'text/csv')})
        self.assertEqual(uploaded.status_code,200,uploaded.text)
        part=next(p for p in self.client.get('/api/fragments',params={'plate':'custom-wide'}).json() if p['name']=='Unrestricted')
        with TestClient(app) as reopened:
            self.assertEqual(reopened.get('/api/inventory/rules').json()['rules'],rules)
            row=reopened.get('/api/fragments/'+str(part['id'])).json()
            self.assertEqual(row['stock_rules'],rules)
            applied_row=reopened.get('/api/fragments/1').json()
            self.assertEqual(applied_row['stock_rules'],rules)
        for path in ('/api/inventory/resuspension','/api/inventory/worksheet'):
            result=self.client.get(path,params=rules | {'no_backbone':True,'plate':'custom-wide'})
            self.assertEqual(result.status_code,200,result.text if result.status_code!=200 else '')
        queued=self.client.post('/api/build-list',json={'name':'Large reaction','fragment_ids':[1],
            'settings':rules | {'reaction_volume_ul':50000,'master_mix_volume_ul':12000}})
        self.assertEqual(queued.status_code,200,queued.text)
        patched=self.client.patch('/api/build-list/'+str(queued.json()['id'])+'/settings',json={'reaction_volume_ul':60000,'insert_ratio':1000})
        self.assertEqual(patched.status_code,200,patched.text)
        self.assertEqual(patched.json()['request']['settings']['reaction_volume_ul'],60000)
        self.assertEqual(patched.json()['request']['settings']['insert_ratio'],1000)

    def test_selected_custom_rules_persist_without_changing_global_rules_or_yields(self):
        initial=self.client.get('/api/fragments/1').json()
        untouched=self.client.get('/api/fragments/2').json()
        self.client.patch('/api/fragments/1',json={'stock_ng_ul':20,'resuspension_ul':30,'derive':'yield_ng','yield_locked':False})
        defaults=self.client.get('/api/inventory/rules').json()['rules']
        custom=defaults | {'backbone_pmol':.1,'insert_ratio':2,'preferred_resuspension_ul':30,
            'max_resuspension_ul':90,'min_input_volume_ul':2,'max_input_volume_ul':4}
        applied=self.client.post('/api/inventory/apply-rules',json={'fragment_ids':[1,1],'rules':custom})
        self.assertEqual(applied.status_code,200,applied.text)
        self.assertEqual(applied.json()['count'],1)
        self.assertEqual(self.client.get('/api/inventory/rules').json()['rules'],defaults)
        row=next(r for r in self.client.get('/api/inventory/records').json()['rows'] if r['id']==1)
        self.assertEqual(row['yield_ng'],600)
        self.assertFalse(row['yield_locked'])
        self.assertEqual(row['applied_rules'],custom)
        self.assertAlmostEqual(row['stock_ng_ul']*row['resuspension_ul'],600)
        self.assertAlmostEqual(row['input_volume_ul']*row['stock_ng_ul'],.1*2*initial['physical_length']*660/1000)
        self.assertGreaterEqual(row['input_volume_ul'],2)
        self.assertLessEqual(row['input_volume_ul'],4)
        self.assertLessEqual(row['resuspension_ul'],90)
        self.assertIn('unmeasured',row['concentration_basis'])
        self.assertEqual(self.client.get('/api/fragments/2').json(),untouched)
        planned=self.client.post('/api/assemblies/preview',json={'fragment_ids':[1],'settings':{'planner_stock_mode':'saved'}})
        self.assertEqual(planned.status_code,200,planned.text)
        reaction_row=planned.json()['reaction']['rows'][0]
        self.assertAlmostEqual(reaction_row['concentration_ng_ul'],row['stock_ng_ul'])
        self.assertIn('unmeasured',reaction_row['concentration_basis'])
        with TestClient(app) as reopened:
            persisted=reopened.get('/api/fragments/1').json()
            self.assertEqual(persisted['stock_rules'],custom)
            self.assertAlmostEqual(persisted['stock_ng_ul'],row['stock_ng_ul'])
        self.client.put('/api/inventory/rules',json=defaults | {'backbone_pmol':.03})
        still=next(r for r in self.client.get('/api/inventory/records').json()['rows'] if r['id']==1)
        self.assertEqual(still['rules'],custom)
        self.assertEqual(still['input_volume_ul'],row['input_volume_ul'])
        self.client.post('/api/inventory/reset',json={'fragment_ids':[1]})
        reset=self.client.get('/api/fragments/1').json()
        self.assertIsNone(reset['stock_rules'])
        self.assertEqual(reset['yield_ng'],initial['default_yield_ng'])
        self.assertFalse(reset['stock_saved'])

    def test_selected_rules_fail_atomically_for_infeasible_or_unknown_rows(self):
        rules=self.client.get('/api/inventory/rules').json()['rules']
        self.client.patch('/api/fragments/2',json={'yield_ng':.001})
        before=self.client.get('/api/fragments/1').json()
        failed=self.client.post('/api/inventory/apply-rules',json={'fragment_ids':[1,2],'rules':rules})
        self.assertEqual(failed.status_code,422,failed.text)
        self.assertIn('No rows changed',failed.json()['detail']['message'])
        self.assertEqual(self.client.get('/api/fragments/1').json(),before)
        missing=self.client.post('/api/inventory/apply-rules',json={'fragment_ids':[1,999999],'rules':rules})
        self.assertEqual(missing.status_code,404)
        self.assertEqual(self.client.get('/api/fragments/1').json(),before)
        self.assertEqual(self.client.post('/api/inventory/apply-rules',json={'fragment_ids':[],'rules':rules}).status_code,422)
        self.assertEqual(self.client.post('/api/inventory/apply-rules',json={'fragment_ids':[1],'rules':rules | {'min_input_volume_ul':5,'max_input_volume_ul':2}}).status_code,422)

    def test_fragment_detection_lists_all_present_enzymes_independent_of_digest(self):
        sequence='GGTCTCAAAACGTCTCAAAAGAAGACAAAAAAGCTT'
        content='Name,Plate ID,Well,Sequence,Yield ng\nAll enzymes,enzyme-flags,A1,'+sequence+',300\nNo enzymes,enzyme-flags,A2,'+'ATAC'*75+',300\n'
        uploaded=self.client.post('/api/uploads/plate?no_backbone=true',files={'file':('flags.csv',content,'text/csv')})
        self.assertEqual(uploaded.status_code,200,uploaded.text)
        for enzyme in ('BsaI','BsmBI','BsmBI + BsaI'):
            rows=self.client.get('/api/fragments',params={'plate':'enzyme-flags','enzyme':enzyme}).json()
            detected={row['name']:row['sites'] for row in rows}
            self.assertEqual({site['enzyme'] for site in detected['All enzymes']},{'BsaI','BsmBI','BbsI','HindIII'})
            self.assertEqual(detected['No enzymes'],[])

    def test_inventory_defaults_linked_edits_persistence_and_yield_reset(self):
        initial=self.client.get('/api/fragments/1').json()
        def record(fid=1):
            return next(r for r in self.client.get('/api/inventory/records').json()['rows'] if r['id']==fid)
        default=record()
        self.assertFalse(default['saved'])
        self.assertAlmostEqual(default['stock_ng_ul']*default['resuspension_ul'],initial['yield_ng'])
        self.assertEqual(default['default_yield_ng'],initial['yield_ng'])
        stock=self.client.patch('/api/fragments/1',json={'stock_ng_ul':10}).json()
        self.assertEqual(stock['yield_ng'],initial['yield_ng'])
        self.assertAlmostEqual(stock['resuspension_ul'],initial['yield_ng']/10)
        volume=self.client.patch('/api/fragments/1',json={'resuspension_ul':50}).json()
        self.assertAlmostEqual(volume['stock_ng_ul'],initial['yield_ng']/50)
        changed=self.client.patch('/api/fragments/1',json={'stock_ng_ul':25,'resuspension_ul':20,'derive':'yield_ng'})
        self.assertEqual(changed.status_code,200,changed.text)
        self.assertEqual(changed.json()['yield_ng'],500)
        with TestClient(app) as reopened:
            saved=reopened.get('/api/fragments/1').json()
            self.assertEqual((saved['yield_ng'],saved['stock_ng_ul'],saved['resuspension_ul']),(500,25,20))
            self.assertEqual(saved['default_yield_ng'],initial['yield_ng'])
        self.assertTrue(record()['saved'])
        reset=self.client.post('/api/inventory/reset',json={'fragment_ids':[1]})
        self.assertEqual(reset.json()['count'],1)
        after=record()
        self.assertFalse(after['saved'])
        self.assertEqual(after['yield_ng'],initial['yield_ng'])
        self.assertEqual(after['stock_ng_ul'],default['stock_ng_ul'])
        self.assertEqual(after['resuspension_ul'],default['resuspension_ul'])

    def test_yield_lock_persists_without_committing_planned_stocks_and_reset_is_selected(self):
        initial=self.client.get('/api/fragments/2').json()
        toggled=self.client.patch('/api/fragments/2',json={'yield_locked':False})
        self.assertEqual(toggled.status_code,200,toggled.text)
        self.assertFalse(toggled.json()['yield_locked'])
        self.assertFalse(toggled.json()['inventory']['saved'])
        self.assertIsNone(toggled.json()['stock_ng_ul'])
        with TestClient(app) as reopened:
            self.assertFalse(reopened.get('/api/fragments/2').json()['yield_locked'])
        self.client.patch('/api/fragments/1',json={'stock_ng_ul':17})
        before=self.client.get('/api/fragments/1').json()
        changed=self.client.patch('/api/fragments/2',json={
            'stock_ng_ul':20,'resuspension_ul':30,'derive':'yield_ng','yield_locked':False})
        self.assertEqual(changed.status_code,200,changed.text)
        self.assertEqual(changed.json()['yield_ng'],600)
        reset=self.client.post('/api/inventory/reset',json={'fragment_ids':[2]})
        self.assertEqual(reset.json()['count'],1)
        after=self.client.get('/api/fragments/2').json()
        self.assertTrue(after['yield_locked'])
        self.assertEqual(after['yield_ng'],initial['yield_ng'])
        self.assertFalse(after['stock_saved'])
        self.assertEqual(self.client.get('/api/fragments/1').json()['stock_ng_ul'],before['stock_ng_ul'])
        self.assertEqual(self.client.post('/api/inventory/reset',json={'fragment_ids':[]}).json()['count'],0)

    def test_inventory_rules_persist_and_do_not_replace_saved_stocks(self):
        self.client.patch('/api/fragments/2',json={'stock_ng_ul':20})
        saved=self.client.get('/api/fragments/2').json()
        rules=self.client.get('/api/inventory/rules').json()['rules']
        custom=rules | {'backbone_pmol':.1,'insert_ratio':2,'preferred_resuspension_ul':60,
            'max_resuspension_ul':90,'min_input_volume_ul':2,'max_input_volume_ul':4}
        result=self.client.put('/api/inventory/rules',json=custom)
        self.assertEqual(result.status_code,200,result.text)
        with TestClient(app) as reopened:
            self.assertEqual(reopened.get('/api/inventory/rules').json()['rules'],custom)
            records={row['id']:row for row in reopened.get('/api/inventory/records').json()['rows']}
            self.assertEqual(records[2]['stock_ng_ul'],20)
            self.assertEqual(records[2]['resuspension_ul'],saved['resuspension_ul'])
            self.assertFalse(records[1]['saved'])
            target=.1*1000*2*1808*660/1_000_000
            self.assertAlmostEqual(records[1]['stock_ng_ul']*records[1]['input_volume_ul'],target)
            recipe=next(r for r in reopened.get('/api/recipes').json() if r['name']=='pIVT_V2')
            preview=reopened.post('/api/assemblies/preview',json={'fragment_ids':recipe['no_backbone_fragment_ids']}).json()
            self.assertTrue(all(row['target_fmol']==200 for row in preview['reaction']['rows']))
            self.assertEqual(preview['effective_settings']['max_resuspension_ul'],90)
        invalid=self.client.put('/api/inventory/rules',json=custom | {'min_input_volume_ul':5,'max_input_volume_ul':1})
        self.assertEqual(invalid.status_code,422)
        self.assertEqual(self.client.get('/api/inventory/rules').json()['rules'],custom)
        self.assertEqual(self.client.post('/api/inventory/rules/reset').json()['rules'],rules)
        self.assertEqual(self.client.get('/api/fragments/2').json()['stock_ng_ul'],20)

    def test_missing_yield_can_be_derived_and_reset_to_unknown(self):
        content='Name,Plate ID,Well,Sequence,Yield ng\nUnknown mass,unknown_yield,A1,'+'ACGT'*75+',\n'
        response=self.client.post('/api/uploads/plate?no_backbone=true',files={'file':('unknown.csv',content,'text/csv')})
        self.assertEqual(response.status_code,200,response.text)
        part=self.client.get('/api/fragments',params={'plate':'unknown_yield'}).json()[0]
        self.assertIsNone(part['yield_ng'])
        filled=self.client.patch('/api/fragments/'+str(part['id']),json={'yield_ng':300})
        self.assertEqual(filled.status_code,200,filled.text)
        self.assertAlmostEqual(filled.json()['stock_ng_ul']*filled.json()['resuspension_ul'],300)
        response=self.client.patch('/api/fragments/'+str(part['id']),json={'stock_ng_ul':10,'resuspension_ul':30,'derive':'yield_ng'})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['yield_ng'],300)
        self.client.post('/api/inventory/reset',json={'fragment_ids':[part['id']]})
        reset=self.client.get('/api/fragments/'+str(part['id'])).json()
        self.assertIsNone(reset['yield_ng'])
        self.assertIsNone(reset['stock_ng_ul'])

    def test_planner_saved_and_custom_stocks_are_separate_from_inventory(self):
        self.client.patch('/api/fragments/1',json={'stock_ng_ul':25})
        before=self.client.get('/api/fragments/1').json()
        payload={'fragment_ids':[1,2], 'settings':{'planner_stock_mode':'saved','backbone_pmol':.1,
            'insert_ratio':2,'reaction_volume_ul':40,'master_mix_volume_ul':4,
            'preferred_resuspension_ul':10}}
        saved=self.client.post('/api/assemblies/preview',json=payload).json()
        self.assertEqual(saved['reaction']['rows'][0]['concentration_ng_ul'],25)
        self.assertEqual(saved['resuspension']['rows'][0]['resuspension_ul'],before['resuspension_ul'])
        records={r['id']:r for r in self.client.get('/api/inventory/records').json()['rows']}
        self.assertEqual(saved['reaction']['rows'][1]['concentration_ng_ul'],records[2]['stock_ng_ul'])
        self.assertEqual(saved['reaction']['rows'][0]['target_fmol'],200)
        payload['settings'].update(planner_stock_mode='custom',preferred_resuspension_ul=40)
        custom=self.client.post('/api/assemblies/preview',json=payload).json()
        plan=custom['resuspension']['rows'][0]
        self.assertAlmostEqual(custom['reaction']['rows'][0]['concentration_ng_ul'],plan['target_ng_ul'])
        self.assertNotEqual(plan['target_ng_ul'],25)
        self.assertTrue(1<=plan['input_volume_ul']<=5)
        self.assertLess(plan['resuspension_ul'],100)
        override={'yield_ng':600,'stock_ng_ul':12,'resuspension_ul':50,'yield_locked':False}
        payload['settings']['fragment_stock_overrides']={'1':override}
        result=self.client.post('/api/assemblies/preview',json=payload).json()
        self.assertEqual(result['reaction']['rows'][0]['concentration_ng_ul'],12)
        self.assertEqual(result['resuspension']['rows'][0]['yield_ng'],600)
        self.assertEqual(result['reaction']['rows'][0]['concentration_basis'],'Instance override')
        self.assertEqual(self.client.get('/api/fragments/1').json(),before)
        inventory=self.client.post('/api/inventory/resuspension',json=payload['settings']|{'no_backbone':True,'fragment_ids':[1]})
        self.assertEqual(inventory.status_code,200,inventory.text)
        self.assertEqual(len(inventory.json()['rows']),1)
        self.assertEqual(inventory.json()['rows'][0]['target_ng_ul'],12)
        for endpoint,body in [('/api/assemblies/worksheet',payload),('/api/inventory/worksheet',payload['settings']|{'no_backbone':True,'fragment_ids':[1]})]:
            response=self.client.post(endpoint,json=body)
            self.assertEqual(response.status_code,200,response.text if response.status_code!=200 else '')
            book=load_workbook(io.BytesIO(response.content),data_only=False)
            sheet=book['Resuspension plan']
            row=next(r for r in sheet.iter_rows(min_row=3,values_only=True) if r[0]==before['name'])
            self.assertEqual(row[4:8],(600,12,50,result['resuspension']['rows'][0]['input_volume_ul']))
            self.assertEqual(book['Resuspension inventory']['G2'].value,'=IF(J2="","",J2)')
        self.assertEqual(self.client.get('/api/fragments/1').json(),before)

    def test_planner_backbone_stock_remains_manual_in_both_modes(self):
        pid=next(p['id'] for p in self.client.get('/api/plasmids').json() if p['name']=='pIVT_V2')
        for mode in ('saved','custom'):
            payload={'backbone_id':pid,'fragment_ids':[1], 'settings':{'planner_stock_mode':mode,
                'backbone_stock_ng_ul':100,'backbone_pmol':.05,'insert_ratio':2}}
            response=self.client.post('/api/assemblies/preview',json=payload)
            self.assertEqual(response.status_code,200,response.text)
            row=response.json()['reaction']['rows'][0]
            self.assertEqual(row['concentration_ng_ul'],100)
            self.assertAlmostEqual(row['volume_ul'],.05*2827*660/1000/100)
            self.assertEqual(row['feasibility'],'manual_stock')

    def test_plate_well_list_exports_and_validation_are_read_only(self):
        plates=self.client.get('/api/plates').json();plate=plates[0]['plate_id']
        before=self.client.get('/api/plates/detail',params={'plate_id':plate}).json()
        for format,mime in [('pdf','application/pdf'),('xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),('csv','application/zip')]:
            response=self.client.get('/api/plates/well-list',params={'plate_id':plate,'format':format})
            self.assertEqual(response.status_code,200,response.text[:200] if response.status_code!=200 else '')
            self.assertEqual(response.headers['content-type'],mime)
            self.assertIn('-well-list.',response.headers['content-disposition'])
        self.assertEqual(before,self.client.get('/api/plates/detail',params={'plate_id':plate}).json())
        self.assertEqual(self.client.get('/api/plates/well-list',params={'plate_id':'unknown'}).status_code,404)
        self.assertEqual(self.client.get('/api/plates/well-list',params={'plate_id':plate,'format':'exe'}).status_code,422)
        self.assertEqual(self.client.get('/api/plates/well-list',params={'plate_id':plate,'layout':'12'}).status_code,422)

    def test_stored_plate_maps_list_populated_wells_and_current_stock_details(self):
        plates=self.client.get('/api/plates').json()
        self.assertEqual(len(plates),2)
        plate=next(p for p in plates if p['occupied_wells']==96)
        self.assertEqual((plate['format'],plate['row_count'],plate['column_count']),(96,8,12))
        self.assertEqual(plate['active_wells'],96)
        self.assertTrue(plate['filenames'][0].startswith('platemap_'))
        response=self.client.get('/api/plates/detail',params={'plate_id':plate['plate_id'],'enzyme':'BsaI'})
        self.assertEqual(response.status_code,200,response.text)
        rows=response.json()['fragments']
        self.assertEqual(len(rows),96)
        first=next(p for p in rows if p['name']=='F1')
        inventory=next(p for p in self.client.get('/api/inventory/records').json()['rows'] if p['id']==first['id'])
        self.assertEqual(first['inventory'],inventory)
        self.assertEqual(first['digest']['left_overhang'],'TTTC')
        self.assertNotIn('sequence',first)
        self.assertNotIn('sequence',first['digest'])
        self.assertTrue(first['sites'])
        self.assertEqual(self.client.get('/api/plates/detail',params={'plate_id':'unknown'}).status_code,404)

    def test_plate_detection_is_per_well_and_independent_of_assembly_enzyme(self):
        bsai='GGTCTCAACCTAACCGGTATGCGAGACC'
        bsmbi='CGTCTCAACCTAACCGGTATGCGAGACG'
        content='Name,Plate ID,Well,Sequence,Construct sequence,Yield ng\n'
        content+=f'BsaI donor,mixed,A1,{bsai},AAAA{bsai}TTTT,50\n'
        content+=f'BsmBI donor,mixed,A2,{bsmbi},,50\n'
        content+='No sites,mixed,A3,ACGTACGT,,50\n'
        response=self.client.post('/api/uploads/plate?no_backbone=true',files={'file':('mixed.csv',content,'text/csv')})
        self.assertEqual(response.status_code,200,response.text)
        detected=self.client.get('/api/plates/detail',params={'plate_id':'mixed'}).json()
        for legacy_enzyme in ('BsaI','BsmBI','BsmBI + BsaI'):
            self.assertEqual(detected,self.client.get('/api/plates/detail',params={'plate_id':'mixed','enzyme':legacy_enzyme}).json())
        a,b,c=sorted(detected['fragments'],key=lambda p:p['well'])
        self.assertEqual(a['detected_enzymes'],['BsaI'])
        self.assertEqual(b['detected_enzymes'],['BsmBI'])
        self.assertEqual(c['detected_enzymes'],[])
        self.assertEqual([e['overhang'] for e in b['fusion_ends']],['ACCT','TATG'])
        self.assertEqual(a['fusion_ends'][0]['start'],7)
        self.assertEqual(a['fusion_ends'][0]['construct_start'],11)
        self.assertNotIn('enzyme',detected)
        self.assertNotIn('errors',b['digest'])
        self.assertNotIn('input_sequence',b['digest'])

    def test_plate_layout_384_detection_and_multi_plate_upload(self):
        text='Name,Plate ID,Well,Sequence,Yield ng\nHigh row,384 row,I1,ACGTACGT,50\nHigh column,384 column,H13,ACGTACGT,50\nLast well,384 row,P24,ACGTACGT,50\nSparse,96 sparse,A01,ACGTACGT,50\n'
        response=self.client.post('/api/uploads/plate?no_backbone=true',files={'file':('new-maps.csv',text,'text/csv')})
        self.assertEqual(response.status_code,200,response.text)
        plates={p['plate_id']:p for p in self.client.get('/api/plates').json()}
        for name in ('384 row','384 column'):
            self.assertEqual((plates[name]['format'],plates[name]['row_count'],plates[name]['column_count']),(384,16,24))
            self.assertEqual(plates[name]['filenames'],['new-maps.csv'])
        self.assertEqual(plates['96 sparse']['format'],96)
        details=self.client.get('/api/plates/detail',params={'plate_id':'96 sparse'}).json()
        self.assertEqual(details['fragments'][0]['well'],'A1')
        self.assertIn('sparse',details['layout_basis'])

    def test_plate_archive_is_atomic_persistent_and_does_not_reset_hidden_stocks(self):
        from app import database as store
        plate=self.client.get('/api/plates').json()[0]
        parts=[p for p in store.fragments() if p['plate_id']==plate['plate_id']]
        keep=parts[0];self.client.patch('/api/fragments/'+str(keep['id']),json={'stock_ng_ul':19})
        saved=store.get_fragment(keep['id'])
        invalid=self.client.post('/api/bin/archive',json={'fragment_ids':[keep['id']],'plate_ids':[plate['plate_id'],'missing']})
        self.assertEqual(invalid.status_code,404)
        self.assertIsNotNone(store.get_fragment(keep['id']))
        result=self.client.post('/api/bin/archive',json={'plate_ids':[plate['plate_id']]})
        self.assertEqual(result.json()['count'],1)
        self.assertEqual(self.client.get('/api/bin').json()['rows'][0]['kind'],'plate')
        self.assertNotIn(plate['plate_id'],self.client.get('/api/summary').json()['plate_ids'])
        self.assertEqual(self.client.get('/api/fragments/'+str(keep['id'])).status_code,404)
        self.assertEqual(self.client.post('/api/inventory/reset',json={'fragment_ids':[keep['id']]}).status_code,404)
        self.assertEqual(self.client.post('/api/inventory/apply-rules',json={'fragment_ids':[keep['id']],'rules':{}}).status_code,404)
        self.client.post('/api/inventory/reset',json={})
        self.assertEqual(store.plate_fragments(plate['plate_id'])[0]['stock_ng_ul'],saved['stock_ng_ul'])
        with TestClient(app) as reopened:
            self.assertNotIn(plate['plate_id'],[p['plate_id'] for p in reopened.get('/api/plates').json()])
            details=reopened.get('/api/plates/detail',params={'plate_id':plate['plate_id']}).json()
            self.assertTrue(details['archived_at'])
            self.assertEqual(details['active_wells'],0)
        self.assertEqual(self.client.post('/api/bin/restore',json={'plate_ids':[plate['plate_id']]}).json()['count'],1)
        self.assertEqual(store.get_fragment(keep['id'])['stock_ng_ul'],19)
        self.assertEqual(self.client.get('/api/bin').json()['rows'],[])

    def test_plate_restore_preserves_individual_archives_and_blocks_new_import_to_archived_plate(self):
        part=self.client.get('/api/fragments/1').json();plate=part['plate_id']
        self.client.post('/api/bin/archive',json={'fragment_ids':[1]})
        self.client.post('/api/bin/archive',json={'plate_ids':[plate]})
        archived=self.client.get('/api/bin').json()['rows']
        self.assertEqual([r['kind'] for r in archived],['plate'])
        text='Name,Plate ID,Well,Sequence\nNew,'+plate+',P24,ACGTACGT\n'
        response=self.client.post('/api/uploads/plate?no_backbone=true',files={'file':('extra.csv',text,'text/csv')})
        self.assertEqual(response.status_code,422)
        self.assertIn('archived',response.json()['detail'])
        self.client.post('/api/bin/restore',json={'plate_ids':[plate]})
        self.assertEqual(self.client.get('/api/fragments/1').status_code,404)
        self.assertEqual(self.client.get('/api/fragments/2').status_code,200)
        self.assertEqual([r['kind'] for r in self.client.get('/api/bin').json()['rows']],['fragment'])
        self.client.post('/api/bin/restore',json={'fragment_ids':[1]})
        self.assertEqual(self.client.get('/api/fragments/1').status_code,200)

    def test_plate_export_and_selection_support_non_url_safe_plate_identifiers(self):
        identifier='Plate / "quoted" & one'
        content=io.StringIO();writer=csv.writer(content);writer.writerow(['Name','Plate ID','Well','Sequence']);writer.writerow(['test',identifier,'P24','ACGTACGT'])
        response=self.client.post('/api/uploads/plate?no_backbone=true',files={'file':('map.csv',content.getvalue(),'text/csv')})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.client.get('/api/plates/detail',params={'plate_id':identifier}).json()['format'],384)
        exported=self.client.get('/api/plates/export',params={'plate_id':identifier})
        self.assertEqual(exported.status_code,200)
        rows=list(csv.DictReader(io.StringIO(exported.content.decode('utf-8-sig'))))
        self.assertEqual(rows[0]['Plate ID'],identifier)
        self.assertEqual(rows[0]['Well'],'P24')
        self.assertEqual(rows[0]['Sequence'],'ACGTACGT')
        self.assertEqual(self.client.post('/api/bin/archive',json={'plate_ids':[identifier]}).json()['count'],1)
        exported=self.client.get('/api/plates/export',params={'plate_id':identifier})
        self.assertEqual(list(csv.DictReader(io.StringIO(exported.content.decode('utf-8-sig'))))[0]['Archived'],'True')

    def test_plate_migration_backfills_legacy_inventory_without_unarchiving(self):
        from app import database as store
        with store.connection() as db:db.execute('DROP TABLE plates')
        store.initialize()
        plates=self.client.get('/api/plates').json()
        self.assertEqual(len(plates),2)
        self.client.post('/api/bin/archive',json={'plate_ids':[plates[0]['plate_id']]})
        store.initialize()
        self.assertEqual(len(self.client.get('/api/plates').json()),1)
        self.assertEqual(len(self.client.get('/api/bin').json()['rows']),1)

    def test_bin_preserves_and_restores_fragments_plasmids_and_recipes(self):
        self.client.patch('/api/fragments/1',json={'stock_ng_ul':15})
        before=self.client.get('/api/fragments/1').json()
        pid=next(p['id'] for p in self.client.get('/api/plasmids').json() if p['name']=='pIVT_V2')
        plasmid=self.client.get(f'/api/plasmids/{pid}').json()
        selected={'fragment_ids':[1],'plasmid_ids':[pid]}
        invalid=self.client.post('/api/bin/archive',json={'fragment_ids':[1],'plasmid_ids':[999999]})
        self.assertEqual(invalid.status_code,404)
        self.assertEqual(self.client.get('/api/fragments/1').status_code,200)
        response=self.client.post('/api/bin/archive',json=selected)
        self.assertEqual(response.json()['count'],2)
        self.assertEqual(self.client.get('/api/fragments/1').status_code,404)
        self.assertEqual(self.client.get(f'/api/plasmids/{pid}').status_code,404)
        self.assertEqual({row['kind'] for row in self.client.get('/api/bin').json()['rows']},{'fragment','plasmid'})
        self.assertFalse(next(r for r in self.recipes() if r['name']=='pIVT_V2')['available'])
        with TestClient(app) as reopened:
            self.assertEqual(len(reopened.get('/api/bin').json()['rows']),2)
            self.assertEqual(reopened.get('/api/summary').json()['fragments'],101)
            self.assertEqual(reopened.get('/api/summary').json()['plasmids'],6)
        self.client.post('/api/inventory/reset',json={})
        restored=self.client.post('/api/bin/restore',json=selected)
        self.assertEqual(restored.json()['count'],2)
        after=self.client.get('/api/fragments/1').json()
        for field in ('sequence','yield_ng','default_yield_ng','stock_ng_ul','resuspension_ul','notes'):
            self.assertEqual(after[field],before[field])
        actual=self.client.get(f'/api/plasmids/{pid}').json()
        for field in ('sequence','features','recipe','circular'):
            self.assertEqual(actual[field],plasmid[field])
        self.assertEqual(self.client.get('/api/bin').json()['rows'],[])
        self.assertTrue(next(r for r in self.recipes() if r['name']=='pIVT_V2')['available'])

    def test_first_round_incompatible_f6_cannot_be_saved_or_exported(self):
        recipe = next(recipe for recipe in self.recipes() if recipe['name'] == 'pHR_V2 · U6')
        originals = self.client.get('/api/fragments', params={'q': 'F6'}).json()
        original = next(part for part in originals if part['name'] == "F6 (3'LTR_BsaI original)")
        u6 = next(part for part in originals if part['name'] == "F6 (3'LTR_U6_BsmBI)")
        payload = self.payload(recipe)
        payload['fragment_ids'] = [original['id'] if value == u6['id'] else value for value in payload['fragment_ids']]
        response = self.client.post('/api/assemblies/preview', json=payload)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json()['valid'])
        self.assertTrue(response.json()['errors'])
        for endpoint in ['/api/assemblies', '/api/assemblies/genbank']:
            rejected = self.client.post(endpoint, json=payload)
            self.assertEqual(rejected.status_code, 422)
        self.assertEqual(self.client.get('/api/summary').json()['plasmids'], 7)

    def test_stock_persists_and_reaction_uses_physical_dna_length(self):
        fragment = self.client.get('/api/fragments/1').json()
        response = self.client.patch('/api/fragments/1', json={
            'stock_ng_ul': 25, 'resuspension_ul': 50, 'notes': 'Measured stock for integration check'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['stock_ng_ul'], 25)
        # A second app lifespan uses the same persistent file and idempotent seeding.
        with TestClient(app) as second:
            stored = second.get('/api/fragments/1').json()
            self.assertEqual(stored['stock_ng_ul'], 25)
            self.assertEqual(stored['resuspension_ul'], 50)
            self.assertEqual(second.get('/api/summary').json()['fragments'], 102)
        recipe = next(recipe for recipe in self.recipes() if recipe['name'] == 'pIVT_V2')
        payload = self.payload(recipe)
        payload['settings'] = {'backbone_ng': 50, 'backbone_stock_ng_ul': 25,
                               'insert_ratio': 2, 'target_stock_ng_ul': 20, 'resuspension_mode': 'one_ul',
                               'assume_target_stock': True}
        result = self.client.post('/api/assemblies/preview', json=payload).json()
        rows = result['reaction']['rows']
        physical = fragment['physical_length']
        self.assertEqual(rows[0]['length_bp'], physical)
        self.assertEqual(rows[0]['volume_ul'], 2)
        self.assertAlmostEqual(rows[0]['target_fmol'], 50e6 / (660 * physical))
        self.assertAlmostEqual(rows[1]['target_fmol'], 2 * rows[0]['target_fmol'])
        self.assertTrue(any(word in rows[1]['concentration_basis'].lower() for word in ('assumed', 'planned')))
        self.assertTrue(result['reaction']['volume_complete'])

    def test_constrained_resuspension_defaults_to_fifty_fmol_and_valid_pipetting(self):
        response = self.client.get('/api/inventory/resuspension')
        self.assertEqual(response.status_code, 200, response.text)
        plan = response.json()
        self.assertEqual(plan['settings']['resuspension_mode'], 'constrained')
        self.assertEqual(plan['reference']['backbone_pmol'], 0.05)
        self.assertAlmostEqual(plan['reference']['backbone_ng'], 0.05 * 1808 * 660 / 1000)
        self.assertEqual(len(plan['rows']), 102)
        for row in plan['rows']:
            if row.get('manual_stock'):
                self.assertIsNone(row['resuspension_ul'])
                self.assertEqual(row['input_volume_ul'],1)
                self.assertIsNone(row['target_ng_ul'])
                continue
            self.assertTrue(row['feasible'], (row['name'], row['warnings']))
            self.assertAlmostEqual(row['target_fmol'], 50)
            self.assertGreater(row['resuspension_ul'], 0)
            self.assertLess(row['resuspension_ul'], 100)
            self.assertGreaterEqual(row['input_volume_ul'] + 1e-10, 1)
            self.assertLessEqual(row['input_volume_ul'], 5 + 1e-10)
            # A stock prepared with the proposed solvent volume delivers 50 fmol.
            amount = row['yield_ng'] / row['resuspension_ul'] * row['input_volume_ul']
            self.assertAlmostEqual(amount * 1e6 / (row['physical_length'] * 660), 50)
        # Editable planner limits remain positive and ordered.
        for params in ({'max_resuspension_ul': 0}, {'min_input_volume_ul': 0},
                       {'max_input_volume_ul': -1},
                       {'min_input_volume_ul': 5, 'max_input_volume_ul': 1}):
            rejected = self.client.get('/api/inventory/resuspension', params=params)
            self.assertEqual(rejected.status_code, 422, rejected.text)
        self.assertIsNone(self.client.get('/api/fragments/2').json()['stock_ng_ul'])

    def test_viewer_rename_persists_and_preserves_plasmid_sequence_and_template(self):
        recipe=next(r for r in self.recipes() if r['name']=='pIVT_V2')
        saved=self.client.post('/api/assemblies',json={'name':'Viewer rename test',
            'fragment_ids':recipe['no_backbone_fragment_ids'],'save_as_backbone':True}).json()
        route='/api/plasmids/'+str(saved['id'])
        before=self.client.get(route).json()
        renamed=self.client.patch(route,json={'name':'  Renamed_viewer_plasmid  '})
        self.assertEqual(renamed.status_code,200,renamed.text)
        after=renamed.json()
        self.assertEqual(after['name'],'Renamed_viewer_plasmid')
        for key in ('id','sequence','features','circular','notes'):
            self.assertEqual(before[key],after[key])
        self.assertEqual(after['recipe']['name'],'Renamed_viewer_plasmid')
        self.assertTrue(after['recipe']['backbone_template'])
        self.assertIn('Renamed_viewer_plasmid',self.client.get(route+'/genbank').text)
        with TestClient(app) as reopened:
            self.assertEqual(reopened.get(route).json()['name'],'Renamed_viewer_plasmid')
            self.assertTrue(reopened.get(route).json()['recipe']['backbone_template'])

    def test_renamed_example_is_not_recreated_and_invalid_names_are_rejected(self):
        examples=[p for p in self.client.get('/api/plasmids').json() if p['recipe'].get('example_backbone')]
        example=examples[0]
        route='/api/plasmids/'+str(example['id'])
        count=self.client.get('/api/summary').json()['plasmids']
        result=self.client.patch(route,json={'name':'My example backbone'})
        self.assertEqual(result.status_code,200,result.text)
        with TestClient(app) as reopened:
            self.assertEqual(reopened.get('/api/summary').json()['plasmids'],count)
            self.assertEqual(reopened.get(route).json()['name'],'My example backbone')
            self.assertTrue(reopened.get(route).json()['recipe']['example_backbone'])
        for name in ('','   ','A'*101):
            self.assertEqual(self.client.patch(route,json={'name':name}).status_code,422)
        from app import database as store
        other=store.save_plasmid('Conflicting name',self.client.get(route).json()['sequence'])
        self.assertNotEqual(other,example['id'])
        self.assertEqual(self.client.patch(route,json={'name':'Conflicting name'}).status_code,409)
        self.assertEqual(self.client.get(route).json()['name'],'My example backbone')

    def test_build_list_persists_separate_designs_and_settings_and_exports_all(self):
        recipes=[r for r in self.recipes() if r['name'] in ('pIVT_V2','pHR_V2')]
        items=[]
        for index,recipe in enumerate(recipes,1):
            payload={'name':recipe['name'],'enzyme':recipe['enzyme'],'fragment_ids':recipe['no_backbone_fragment_ids'],
                'settings':{'backbone_pmol':.05*index,'reaction_volume_ul':40,'assume_target_stock':True}}
            added=self.client.post('/api/build-list',json=payload)
            self.assertEqual(added.status_code,200,added.text)
            item=added.json()
            self.assertIsNotNone(item['plan'])
            self.assertEqual(item['request']['fragment_ids'],recipe['no_backbone_fragment_ids'])
            self.assertEqual([part['id'] for part in item['plan']['fragment_checks']],recipe['no_backbone_fragment_ids'])
            self.assertTrue(all(len(part['left_overhang'])==len(part['right_overhang'])==4
                                for part in item['plan']['fragment_checks']))
            self.assertTrue(all(row['target_fmol']==50*index for row in item['plan']['reaction']['rows']))
            items.append(item)
        with TestClient(app) as reopened:
            current=reopened.get('/api/build-list').json()['items']
            self.assertEqual([item['id'] for item in current],[item['id'] for item in items])
            self.assertEqual([item['request']['settings']['backbone_pmol'] for item in current],[.05,.1])
        changed=items[0]['request']
        changed['settings']['insert_ratio']=2
        updated=self.client.put('/api/build-list/'+str(items[0]['id']),json=changed)
        self.assertEqual(updated.status_code,200,updated.text)
        current=self.client.get('/api/build-list').json()['items']
        self.assertEqual(current[0]['request']['settings']['insert_ratio'],2)
        self.assertEqual(current[1]['request']['settings']['insert_ratio'],1)
        exported=self.client.get('/api/build-list/worksheet')
        self.assertEqual(exported.status_code,200,exported.text)
        workbook=load_workbook(io.BytesIO(exported.content))
        self.assertEqual(workbook['Build list'].max_row,3)
        self.assertEqual(workbook['Reaction settings']['C2'].value,2)
        self.assertEqual(workbook['Reaction settings']['C3'].value,1)
        names={row[0] for row in workbook['Pipetting'].iter_rows(min_row=2,values_only=True)}
        self.assertEqual(names,{item['request']['name'] for item in items})
        workbook.close()

    def test_build_plan_fields_save_independently_and_backbone_groups_do_not_cross(self):
        ivt=next(r for r in self.recipes() if r['name']=='pIVT_V2')
        hr=next(r for r in self.recipes() if r['name']=='pHR_V2')
        builds=[]
        for index,recipe in enumerate((ivt,ivt,hr)):
            payload=self.payload(recipe)|{'name':'Build '+str(index),'settings':{'backbone_stock_ng_ul':100,'assume_target_stock':True}}
            if index==2:
                payload['starter_fragment_id']=2;payload['fragment_ids']=[i for i in payload['fragment_ids'] if i!=2]
            response=self.client.post('/api/build-list',json=payload)
            self.assertEqual(response.status_code,200,response.text);builds.append(response.json())
        before_inventory=self.client.get('/api/inventory/records').json()
        route='/api/build-list/'+str(builds[0]['id'])+'/settings'
        response=self.client.patch(route,json={'backbone_pmol':.08,'insert_ratio':2,'master_mix_volume_ul':0,'reaction_volume_ul':40,'backbone_stock_ng_ul':80})
        self.assertEqual(response.status_code,200,response.text)
        item=response.json();self.assertEqual(item['request']['fragment_ids'],builds[0]['request']['fragment_ids'])
        self.assertEqual(item['request']['settings']['master_mix_volume_ul'],0)
        self.assertEqual(item['plan']['reaction']['master_mix_ul'],0)
        self.assertEqual(item['plan']['reaction']['reaction_volume_ul'],40)
        backbone=next(row for row in item['plan']['reaction']['rows'] if row['role']=='backbone')
        self.assertAlmostEqual(backbone['volume_ul'],.08*1808*.66/80)
        self.assertTrue(all(row['target_fmol']==160 for row in item['plan']['reaction']['rows'] if row['role']!='backbone'))
        self.client.patch(route,json={'insert_ratio':3})
        current=self.client.get('/api/build-list').json()['items']
        self.assertEqual(current[0]['request']['settings']['backbone_pmol'],.08)
        self.assertEqual(current[1]['request']['settings']['insert_ratio'],1)
        applied=self.client.put('/api/build-list/backbone-stocks/apply',json={'kind':'fragment','id':builds[0]['request']['starter_fragment_id'],'stock_ng_ul':50})
        self.assertEqual(applied.status_code,200,applied.text);self.assertEqual(len(applied.json()['items']),2)
        current=self.client.get('/api/build-list').json()['items']
        self.assertEqual([b['request']['settings']['backbone_stock_ng_ul'] for b in current],[50,50,100])
        self.assertEqual(current[0]['request']['settings']['insert_ratio'],3)
        self.assertEqual(current[0]['request']['settings']['reaction_volume_ul'],40)
        with TestClient(app) as reopened:
            self.assertEqual([b['request']['settings']['backbone_stock_ng_ul'] for b in reopened.get('/api/build-list').json()['items']],[50,50,100])
        for patch in ({'backbone_pmol':0},{'insert_ratio':-1},{'reaction_volume_ul':None},{'backbone_stock_ng_ul':0}):
            self.assertEqual(self.client.patch(route,json=patch).status_code,422)
        self.assertEqual(self.client.patch(route,json={'backbone_stock_ng_ul':None}).status_code,200)
        self.assertEqual(self.client.get('/api/inventory/records').json(),before_inventory)
        self.assertEqual(self.client.patch('/api/build-list/9999/settings',json={'insert_ratio':2}).status_code,404)

    def test_build_list_copies_have_distinct_names_and_keep_manual_backbone_stocks(self):
        recipe=next(r for r in self.recipes() if r['name']=='pIVT_V2')
        payload=self.payload(recipe)|{'name':'IVT test','settings':{'backbone_stock_ng_ul':2000,'assume_target_stock':True}}
        first=self.client.post('/api/build-list',json=payload).json()
        payload['settings']['backbone_stock_ng_ul']=100
        second=self.client.post('/api/build-list',json=payload).json()
        self.assertEqual(first['request']['name'],'IVT test')
        self.assertEqual(second['request']['name'],'IVT test (2)')
        for item,stock in [(first,2000),(second,100)]:
            backbone=next(row for row in item['plan']['reaction']['rows'] if row['role']=='backbone')
            self.assertEqual(backbone['concentration_ng_ul'],stock)
            self.assertAlmostEqual(backbone['volume_ul'],.05*1808*660/1000/stock)
            self.assertTrue(backbone['feasible'])
            self.assertIsNone(backbone['resuspension_ul'])

    def test_build_list_removal_is_reversible_and_missing_parts_are_reported(self):
        recipe=next(r for r in self.recipes() if r['name']=='pIVT_V2')
        payload={'name':'Queued IVT','fragment_ids':recipe['no_backbone_fragment_ids']}
        item=self.client.post('/api/build-list',json=payload).json()
        route='/api/build-list/'+str(item['id'])
        self.assertEqual(self.client.delete(route).status_code,200)
        removed=self.client.get('/api/build-list').json()
        self.assertEqual(removed['items'],[])
        self.assertEqual(removed['last_removed_id'],item['id'])
        self.assertEqual(self.client.get('/api/build-list/worksheet').status_code,422)
        self.assertEqual(self.client.post(route+'/restore').status_code,200)
        part=payload['fragment_ids'][0]
        self.client.post('/api/bin/archive',json={'fragment_ids':[part]})
        current=self.client.get('/api/build-list').json()['items'][0]
        self.assertEqual(current['status'],'Unavailable')
        self.assertIsNone(current['plan'])
        self.assertIn('not found',current['error'])
        exported=self.client.get('/api/build-list/worksheet')
        self.assertEqual(exported.status_code,200)
        workbook=load_workbook(io.BytesIO(exported.content))
        self.assertEqual(workbook['Build list']['C2'].value,'Unavailable')
        workbook.close()
        self.client.post('/api/bin/restore',json={'fragment_ids':[part]})
        self.assertIsNotNone(self.client.get('/api/build-list').json()['items'][0]['plan'])
        invalid=self.client.put(route,json=payload|{'fragment_ids':[999999]})
        self.assertEqual(invalid.status_code,404)
        self.assertEqual(self.client.get('/api/build-list').json()['items'][0]['request']['fragment_ids'],payload['fragment_ids'])

    def test_planner_overrides_recalculate_and_export_without_changing_inventory_defaults(self):
        defaults=self.client.get('/api/inventory/rules').json()['rules']
        recipe=next(r for r in self.recipes() if r['name']=='pIVT_V2')
        payload={'fragment_ids':recipe['no_backbone_fragment_ids'],'settings':{
            'backbone_pmol':.1,'insert_ratio':3,'min_input_volume_ul':.5,
            'max_input_volume_ul':8,'max_resuspension_ul':150,
            'preferred_resuspension_ul':120,'assume_target_stock':True}}
        response=self.client.post('/api/assemblies/preview',json=payload)
        self.assertEqual(response.status_code,200,response.text)
        self.assertTrue(all(row['target_fmol']==300 for row in response.json()['reaction']['rows']))
        exported=self.client.post('/api/assemblies/worksheet',json=payload)
        self.assertEqual(exported.status_code,200,exported.text)
        wb=load_workbook(io.BytesIO(exported.content))
        self.assertEqual(wb['Settings']['B5'].value,3)
        self.assertEqual(wb['Settings']['B15'].value,150)
        self.assertEqual(wb['Settings']['B16'].value,.5)
        self.assertEqual(wb['Settings']['B17'].value,8)
        self.assertEqual(wb['Resuspension plan']['I3'].value,300)
        wb.close()
        self.assertEqual(self.client.get('/api/inventory/rules').json()['rules'],defaults)

    def test_manual_backbone_is_exempt_from_fragment_planning_limits(self):
        recipe=next(r for r in self.recipes() if r['name']=='pIVT_V2')
        payload=self.payload(recipe)
        payload['settings']={'backbone_stock_ng_ul':1000,'assume_target_stock':True,
            'min_input_volume_ul':2,'max_input_volume_ul':3,'max_resuspension_ul':1}
        response=self.client.post('/api/assemblies/preview',json=payload)
        self.assertEqual(response.status_code,200,response.text)
        row=next(row for row in response.json()['reaction']['rows'] if row['role']=='backbone')
        self.assertAlmostEqual(row['volume_ul'],.05*1808*660/1000/1000)
        self.assertTrue(row['feasible'])
        self.assertIsNone(row['resuspension_ul'])
        self.assertEqual(row['warnings'],[])
        plan=self.client.get('/api/inventory/resuspension',params={
            'starter_fragment_id':1,'backbone_stock_ng_ul':1000,
            'min_input_volume_ul':2,'max_input_volume_ul':3}).json()
        manual=next(row for row in plan['rows'] if row.get('manual_stock'))
        self.assertAlmostEqual(manual['input_volume_ul'],.059664)
        self.assertIsNone(manual['resuspension_ul'])
        payload['settings'].pop('backbone_stock_ng_ul')
        row=next(row for row in self.client.post('/api/assemblies/preview',json=payload).json()['reaction']['rows'] if row['role']=='backbone')
        self.assertEqual(row['volume_ul'],1)
        self.assertIsNone(row['concentration_ng_ul'])
        self.assertEqual(row['feasibility'],'default_volume')
        self.assertEqual(row['warnings'],[])

    def test_constrained_upload_reports_infeasible_rows_without_invented_volumes(self):
        content = io.StringIO()
        fields = ['Name', 'Plate ID', 'Well', 'Sequence', 'Yield ng']
        writer = csv.DictWriter(content, fieldnames=fields)
        writer.writeheader()
        for name, well, mass in [('Requires 100 uL', 'A1', 990),
                                 ('Too little mass', 'A2', 9.8), ('Unknown mass', 'A3', '')]:
            writer.writerow({'Name': name, 'Plate ID': 'boundary_plate', 'Well': well,
                'Sequence': 'ACGT' * 75, 'Yield ng': mass})
        upload = {'file': ('boundaries.csv', content.getvalue().encode(), 'text/csv')}
        preview = self.client.post('/api/uploads/plate?preview=true&starter_fragment_id=1', files=upload)
        self.assertEqual(preview.status_code, 200, preview.text)
        rows = preview.json()['resuspension']['rows']
        self.assertEqual([row['feasibility'] for row in rows],
            ['resuspension_limit', 'insufficient_mass', 'missing_data'])
        for row in rows:
            self.assertFalse(row['feasible'])
            self.assertIsNone(row['resuspension_ul'])
            self.assertIsNone(row['target_ng_ul'])
            self.assertTrue(row['warnings'])
        committed = self.client.post('/api/uploads/plate?starter_fragment_id=1', files=upload)
        self.assertEqual(committed.status_code, 200, committed.text)
        self.assertEqual(committed.json()['count'], 3)
        saved = self.client.get('/api/fragments', params={'plate': 'boundary_plate'}).json()
        self.assertEqual(sorted(row['yield_ng'] for row in saved if row['yield_ng'] is not None), [9.8, 990])
        self.assertTrue(all(row['stock_ng_ul'] is None for row in saved))

    def test_resuspension_plan_uses_real_reference_and_equal_moles_in_one_ul(self):
        response = self.client.get('/api/inventory/resuspension', params={'resuspension_mode': 'one_ul'})
        self.assertEqual(response.status_code, 200, response.text)
        plan = response.json()
        self.assertEqual((plan['reference']['name'], plan['reference']['length_bp']), ('F1', 1808))
        self.assertTrue(plan['reference']['default_reference'])
        f1 = next(row for row in plan['rows'] if row['name'] == 'F1')
        f2 = next(row for row in plan['rows'] if row['name'] == 'F2')
        self.assertEqual(f2['physical_length'], 344)
        self.assertAlmostEqual(f2['target_ng_ul'], 50 * 344 / 1808)
        self.assertAlmostEqual(f2['resuspension_ul'], 824 / (50 * 344 / 1808))
        self.assertEqual(f2['input_volume_ul'], 1)
        self.assertAlmostEqual(f2['target_fmol'], f1['target_fmol'])
        for row in (f1, f2):
            moles = row['target_ng_ul'] * row['input_volume_ul'] * 1e6 / (660 * row['physical_length'])
            self.assertAlmostEqual(moles, 50e6 / (660 * 1808))
        self.assertIsNone(self.client.get('/api/fragments/2').json()['stock_ng_ul'])
        # Physical reference selection overrides a client-supplied or obsolete length.
        custom = self.client.post('/api/inventory/resuspension', json={
            'starter_fragment_id': f2['id'], 'backbone_length_bp': 1234, 'resuspension_mode': 'one_ul',
            'backbone_ng': 25, 'insert_ratio': 1, 'input_volume_ul': 1,
            'plate': f2['plate_id']}).json()
        self.assertEqual(custom['reference']['length_bp'], 344)
        self.assertEqual(len(custom['rows']), 96)
        same = next(row for row in custom['rows'] if row['id'] == f2['id'])
        self.assertAlmostEqual(same['target_ng_ul'], 25)
        fixed = self.client.get('/api/inventory/resuspension', params={
            'resuspension_mode': 'fixed_stock', 'target_stock_ng_ul': 20}).json()
        self.assertTrue(all(row['target_ng_ul'] == 20 for row in fixed['rows']))
        for invalid in ({'input_volume_ul': 0}, {'insert_ratio': 0}, {'resuspension_mode': 'other'}):
            self.assertEqual(self.client.get('/api/inventory/resuspension', params=invalid).status_code, 422)

    def test_reaction_plan_is_one_ul_only_for_planned_stocks(self):
        recipe = next(recipe for recipe in self.recipes() if recipe['name'] == 'pIVT_V2')
        payload = self.payload(recipe)
        payload['settings'] = {'backbone_ng': 50, 'assume_target_stock': True,
            'backbone_length_bp': 1234, 'resuspension_mode': 'one_ul'}
        response = self.client.post('/api/assemblies/preview', json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()
        rows = result['reaction']['rows']
        self.assertTrue(all(abs(row['volume_ul'] - 1) < 1e-9 for row in rows))
        self.assertTrue(all(abs(row['target_fmol'] - rows[0]['target_fmol']) < 1e-9 for row in rows))
        self.assertEqual(result['effective_settings']['backbone_length_bp'], 1808)
        self.assertEqual(result['recipe']['settings']['backbone_length_bp'], 1234)
        f2 = next(part for part in result['parts'] if part['name'] == 'F2')
        self.client.patch('/api/fragments/' + str(f2['id']), json={'stock_ng_ul': 20})
        measured = self.client.post('/api/assemblies/preview', json=payload).json()
        measured_f2 = next(row for row in measured['reaction']['rows'] if row['id'] == f2['id'])
        self.assertAlmostEqual(measured_f2['volume_ul'], (50 * 344 / 1808) / 20)
        self.assertAlmostEqual(measured_f2['target_ng_ul'], 50 * 344 / 1808)
        self.assertTrue(any('instead of the planned' in warning for warning in measured_f2['warnings']))
        self.assertEqual(self.client.get('/api/fragments/' + str(f2['id'])).json()['stock_ng_ul'], 20)

    def test_inventory_worksheet_keeps_reference_at_one_mole_when_insert_ratio_changes(self):
        params = {'starter_fragment_id': 1, 'insert_ratio': 3, 'resuspension_mode': 'one_ul',
            'plate': 'pSHPs0921B1498017B'}
        plan_response = self.client.get('/api/inventory/resuspension', params=params)
        self.assertEqual(plan_response.status_code, 200, plan_response.text)
        api_rows = {row['name']: row for row in plan_response.json()['rows']}
        downloaded = self.client.get('/api/inventory/worksheet', params=params)
        self.assertEqual(downloaded.status_code, 200, downloaded.text)
        wb = load_workbook(io.BytesIO(downloaded.content))
        snapshot = {row[0]: row for row in wb['Resuspension plan'].iter_rows(min_row=3, values_only=True)}
        self.assertAlmostEqual(snapshot['F1'][5], 50)
        self.assertAlmostEqual(snapshot['F2'][5], 3 * 50 * 344 / 1808)
        for name in ('F1', 'F2'):
            self.assertAlmostEqual(snapshot[name][5], api_rows[name]['target_ng_ul'])
            self.assertAlmostEqual(snapshot[name][6], api_rows[name]['resuspension_ul'])
            self.assertAlmostEqual(snapshot[name][8], api_rows[name]['target_fmol'])
        self.assertAlmostEqual(snapshot['F2'][8], 3 * snapshot['F1'][8])
        # Editable Excel targets apply the same backbone-vs-insert dose rule.
        inventory = {row[1]: row for row in wb['Resuspension inventory'].iter_rows(min_row=2, values_only=True)}
        self.assertIn('/Settings!B8*1/Settings!B12', inventory['F1'][6])
        self.assertIn('/Settings!B8*Settings!B5/Settings!B12', inventory['F2'][6])
        wb.close()

    def test_upload_previews_and_commits_only_incoming_resuspension_rows(self):
        source = APP_ROOT / 'examples' / 'platemap_pSHPs0921B1498017B.csv'
        rows = list(csv.DictReader(io.StringIO(source.read_text())))[:2]
        for row in rows:
            row['Plate ID'] = 'new_plate'
        content = io.StringIO()
        writer = csv.DictWriter(content, fieldnames=list(rows[0]))
        writer.writeheader(); writer.writerows(rows)
        upload = {'file': ('new_plate.csv', content.getvalue().encode(), 'text/csv')}
        before = self.client.get('/api/summary').json()
        preview = self.client.post('/api/uploads/plate?preview=true&starter_fragment_id=1', files=upload)
        self.assertEqual(preview.status_code, 200, preview.text)
        planned = preview.json()['resuspension']
        self.assertEqual(len(planned['rows']), 2)
        self.assertTrue(all(row['plate_id'] == 'new_plate' for row in planned['rows']))
        self.assertEqual(planned['reference']['length_bp'], 1808)
        self.assertEqual(self.client.get('/api/summary').json(), before)
        bad = self.client.post('/api/uploads/plate?starter_fragment_id=999999', files=upload)
        self.assertEqual(bad.status_code, 404)
        self.assertEqual(self.client.get('/api/summary').json(), before)
        committed = self.client.post('/api/uploads/plate?starter_fragment_id=1', files=upload)
        self.assertEqual(committed.status_code, 200, committed.text)
        self.assertEqual(committed.json()['count'], 2)
        self.assertEqual(len(committed.json()['resuspension']['rows']), 2)
        self.assertEqual(self.client.get('/api/summary').json()['fragments'], before['fragments'] + 2)
        uploaded = self.client.get('/api/fragments', params={'plate': 'new_plate'}).json()
        self.assertTrue(all(row['stock_ng_ul'] is None for row in uploaded))
        repeated = self.client.post('/api/uploads/plate?starter_fragment_id=1', files=upload).json()
        self.assertTrue(repeated['already_imported'])
        self.assertEqual(repeated['count'], 0)
        self.assertEqual(len(repeated['resuspension']['rows']), 2)

    def test_inventory_reference_requires_unique_default_and_explicit_plasmid_works(self):
        from app import database as store
        with store.connection() as db:
            db.execute("UPDATE fragments SET name='renamed reference' WHERE name='F1'")
        response = self.client.get('/api/inventory/resuspension')
        self.assertEqual(response.status_code, 422)
        self.assertIn('unique F1', response.json()['detail'])
        source = self.client.get('/api/plasmids').json()[0]
        explicit = self.client.get('/api/inventory/resuspension', params={'backbone_id': source['id']})
        self.assertEqual(explicit.status_code, 200)
        self.assertEqual(explicit.json()['reference']['length_bp'], source['length'])
        self.assertFalse(explicit.json()['reference']['default_reference'])
        self.assertEqual(self.client.get('/api/inventory/resuspension', params={
            'backbone_id': source['id'], 'starter_fragment_id': 1}).status_code, 422)
        worksheet = self.client.get('/api/inventory/worksheet', params={
            'starter_fragment_id': 1, 'plate': 'pSHPs0921B1498017B', 'input_volume_ul': 1, 'resuspension_mode': 'one_ul'})
        self.assertEqual(worksheet.status_code, 200, worksheet.text)
        wb = load_workbook(io.BytesIO(worksheet.content))
        self.assertIn('Resuspension plan', wb.sheetnames)
        self.assertEqual(wb['Settings']['B8'].value, 1808)
        self.assertEqual(wb['Settings']['B11'].value, 'one_ul')
        wb.close()

    def test_duplicate_imports_and_conflicting_well_are_atomic(self):
        plate = APP_ROOT / 'examples' / 'platemap_pSHPs0921B1498017B.csv'
        response = self.client.post('/api/uploads/plate', files={'file': (plate.name, plate.read_bytes(), 'text/csv')})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['already_imported'])
        self.assertEqual(response.json()['count'], 0)
        self.assertEqual(self.client.get('/api/summary').json()['fragments'], 102)
        rows = list(csv.DictReader(io.StringIO(plate.read_text())))
        rows[0]['Name'] = 'Different DNA at occupied well'
        content = io.StringIO()
        writer = csv.DictWriter(content, fieldnames=list(rows[0]))
        writer.writeheader(); writer.writerows(rows)
        prior_history = self.client.get('/api/uploads').json()
        conflict = self.client.post('/api/uploads/plate', files={'file': ('conflict.csv', content.getvalue().encode(), 'text/csv')})
        self.assertEqual(conflict.status_code, 422)
        self.assertIn('already occupied', conflict.json()['detail'])
        self.assertEqual(self.client.get('/api/summary').json()['fragments'], 102)
        self.assertEqual(self.client.get('/api/uploads').json(), prior_history)
        original_fragment = self.client.get('/api/fragments/1').json()
        self.assertEqual(original_fragment['name'], 'F1')

    def test_genbank_export_roundtrip_matches_preview_and_saved_plasmid(self):
        recipe = next(recipe for recipe in self.recipes() if recipe['name'] == 'pHR_V2 · U6')
        payload = self.payload(recipe)
        preview = self.client.post('/api/assemblies/preview', json=payload).json()
        response = self.client.post('/api/assemblies/genbank', json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn('attachment;', response.headers['content-disposition'])
        records = parse_sequence_file(response.text, 'assembly.gb')
        self.assertEqual(len(records), 1)
        self.assertTrue(records[0]['circular'])
        self.assertEqual(records[0]['sequence'], preview['sequence'])
        labels = {f.get('label') for f in records[0]['features']}
        self.assertTrue({'F21(mod)', 'F24', 'F25', 'F26'} <= labels)
        saved = self.client.post('/api/assemblies', json=payload)
        self.assertEqual(saved.status_code, 200)
        download = self.client.get(f"/api/plasmids/{saved.json()['id']}/genbank")
        parsed = parse_sequence_file(download.text, 'saved.gb')[0]
        self.assertEqual(parsed['sequence'], preview['sequence'])
        self.assertEqual(self.client.get('/api/plasmids/' + str(saved.json()['id'])).json()['sequence'], preview['sequence'])

    def test_fragment_only_assemblies_detect_and_preserve_actual_topology(self):
        recipe = next(recipe for recipe in self.recipes() if recipe['name'] == 'pIVT_V2')
        self.assertTrue(recipe['no_backbone'])
        payload = {'name': 'Fragment-only circle', 'fragment_ids': recipe['no_backbone_fragment_ids'],
            'settings': {'assume_target_stock': True}, 'circular': None}
        response = self.client.post('/api/assemblies/preview', json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()
        self.assertTrue(result['valid'], result['errors'])
        self.assertTrue(result['circular'])
        self.assertEqual(result['resuspension']['reference']['type'], 'none')
        self.assertIsNone(result['effective_settings']['backbone_length_bp'])
        self.assertTrue(all(row['role'] == 'insert' for row in result['reaction']['rows']))
        self.assertTrue(all(row['target_fmol'] == 50 for row in result['reaction']['rows']))
        self.assertIsNone(result['reaction']['backbone_fmol'])
        self.assertEqual(result['reaction']['input_fmol'], 50)
        for endpoint in ('/api/assemblies/genbank', '/api/assemblies'):
            saved = self.client.post(endpoint, json=payload)
            self.assertEqual(saved.status_code, 200, saved.text)
            if endpoint.endswith('genbank'):
                parsed = parse_sequence_file(saved.text)[0]
                self.assertTrue(parsed['circular'])
                self.assertEqual(parsed['sequence'], result['sequence'])
            else:
                self.assertTrue(saved.json()['circular'])
                self.assertTrue(self.client.get('/api/plasmids/' + str(saved.json()['id'])).json()['circular'])

        # A short first fragment is an insert, and an open chain stays linear.
        payload.update(name='Fragment-only linear', fragment_ids=[recipe['no_backbone_fragment_ids'][1]])
        result = self.client.post('/api/assemblies/preview', json=payload).json()
        self.assertTrue(result['valid'], result['errors'])
        self.assertFalse(result['circular'])
        row = result['reaction']['rows'][0]
        self.assertLess(row['length_bp'], 1500)
        self.assertEqual(row['role'], 'insert')
        self.assertEqual(row['target_fmol'], 50)
        self.assertTrue(row['feasible'])
        export = self.client.post('/api/assemblies/genbank', json=payload)
        self.assertEqual(export.status_code, 200, export.text)
        self.assertFalse(parse_sequence_file(export.text)[0]['circular'])
        saved = self.client.post('/api/assemblies', json=payload)
        self.assertEqual(saved.status_code, 200, saved.text)
        self.assertFalse(self.client.get('/api/plasmids/' + str(saved.json()['id'])).json()['circular'])

        worksheet = self.client.post('/api/assemblies/worksheet', json=payload)
        self.assertEqual(worksheet.status_code, 200, worksheet.text)
        wb = load_workbook(io.BytesIO(worksheet.content))
        self.assertIsNone(wb['Settings']['B8'].value)
        self.assertEqual(wb['Settings']['B18'].value, 1)
        self.assertEqual(wb['Assembly pipetting']['B4'].value, 'Insert')
        self.assertAlmostEqual(wb['Resuspension plan']['I3'].value, 50)
        wb.close()

    def test_multiple_complete_cassettes_can_be_promoted_and_imported_as_templates(self):
        from app.main import backbone_template_info
        from app.sequence_engine import prepare_backbone
        cassette=lambda left,right,reverse='GAGACG',forward='CGTCTC':left+'C'+reverse+'AAAGATATC'+forward+'A'+right
        examples=[('Two BsmBI cassettes',cassette('AAAC','GGTG')+'ACGT'*20+cassette('TTGT','GGGA')+'ACGT'*20,'BsmBI'),
                  ('Two enzyme cassettes',cassette('AAAC','GGTG')+'ACGT'*20+cassette('TTGT','GGGA','GAGACC','GGTCTC')+'ACGT'*20,'BsmBI + BsaI'),
                  ('Mixed boundary cassette',cassette('AAAC','GGTG','GAGACG','GGTCTC')+'ACGT'*40,'BsmBI + BsaI')]
        for name,sequence,enzyme in examples:
            with self.subTest(name=name):
                info=backbone_template_info(sequence,True)
                self.assertTrue(info['eligible'],info)
                expected=1 if name=='Mixed boundary cassette' else 2
                self.assertEqual(info['cassette_counts'][enzyme],expected)
                if expected==2:
                    self.assertTrue(info['requires_multiple_inserts'])
                    # Eligibility must not hide unresolved cuts in a real reaction.
                    digest=prepare_backbone(sequence,enzyme,'slot_1')
                    self.assertFalse(digest['valid']);self.assertTrue(digest['slot']['retained_site_count'])
                uploaded=self.client.post('/api/uploads/plasmid',files={'file':('template.fa','>'+name.replace(' ','_')+'\n'+sequence,'text/plain')})
                self.assertEqual(uploaded.status_code,200,uploaded.text);pid=uploaded.json()['ids'][0]
                promoted=self.client.patch(f'/api/plasmids/{pid}',json={'backbone_template':True})
                self.assertEqual(promoted.status_code,200,promoted.text)
                self.assertTrue(promoted.json()['recipe']['backbone_template'])
                imported=self.client.post('/api/uploads/plasmid?backbone_template=true',files={'file':('template.fa','>'+name.replace(' ','_')+'\n'+sequence,'text/plain')})
                self.assertEqual(imported.status_code,200,imported.text);self.assertEqual(imported.json()['ids'],[pid])
                for mode in ('BsaI','BsmBI','BsmBI + BsaI'):
                    listed=next(p for p in self.client.get('/api/plasmids',params={'enzyme':mode}).json() if p['id']==pid)
                    self.assertTrue(listed['backbone_options']['eligible'])
                    self.assertEqual(listed['backbone_options']['cassette_counts'][enzyme],expected)
                for offset in (0,1,7,len(sequence)-3):
                    rotated=backbone_template_info(sequence[offset:]+sequence[:offset],True)
                    self.assertTrue(rotated['eligible']);self.assertEqual(rotated['cassette_counts'][enzyme],expected)
        # A stray cut is different from a second complete cassette.
        unpaired=examples[0][1]+'AACGTCTCAAAA'
        self.assertFalse(backbone_template_info(unpaired,True)['eligible'])
        self.assertFalse(backbone_template_info(examples[0][1],False)['eligible'])

    def test_circular_assembly_can_be_promoted_to_persistent_backbone_and_reused(self):
        from app import database as store
        recipe=next(recipe for recipe in self.recipes() if recipe['name']=='pIVT_V2')
        payload={'name':'Saved first-round backbone','fragment_ids':recipe['no_backbone_fragment_ids']}
        preview=self.client.post('/api/assemblies/preview',json=payload).json()
        self.assertTrue(preview['backbone_template']['eligible'])
        self.assertEqual(preview['backbone_template']['enzymes'],['BsmBI','BsmBI + BsaI'])
        ordinary=self.client.post('/api/assemblies',json=payload).json()
        self.assertFalse(ordinary['backbone_template'])
        payload['save_as_backbone']=True
        saved=self.client.post('/api/assemblies',json=payload)
        self.assertEqual(saved.status_code,200,saved.text)
        saved=saved.json()
        self.assertEqual(saved['id'],ordinary['id'])
        self.assertTrue(saved['backbone_template'])
        self.assertEqual(saved['backbone_enzymes'],['BsmBI','BsmBI + BsaI'])
        # Reopening the database retains the template and exact annotated product.
        with TestClient(app) as reopened:
            listed=next(p for p in reopened.get('/api/plasmids?enzyme=BsmBI').json() if p['id']==saved['id'])
            self.assertTrue(listed['recipe']['backbone_template'])
            self.assertEqual(listed['recipe']['fragment_ids'],payload['fragment_ids'])
            self.assertEqual(len(listed['slots']),1)
            self.assertFalse(listed['slots'][0]['errors'])
            product=reopened.get('/api/plasmids/'+str(saved['id'])).json()
            self.assertEqual(product['sequence'],preview['sequence'])
            self.assertEqual(product['features'],preview['features'])
        slot=listed['slots'][0]
        donor='CGTCTCA'+slot['left_overhang']+'ACGTTACGATGACCTAGTCA'+slot['right_overhang']+'AGAGACG'
        with store.connection() as db:
            fid=db.execute('INSERT INTO fragments(name,plate_id,well,sequence,yield_ng) VALUES(?,?,?,?,?)',
                ('Next-round insert','backbone_tests','A1',donor,100)).lastrowid
        second={'name':'Second-round product','backbone_id':saved['id'],'enzyme':'BsmBI',
            'slot_id':slot['id'],'fragment_ids':[fid]}
        result=self.client.post('/api/assemblies/preview',json=second)
        self.assertEqual(result.status_code,200,result.text)
        result=result.json()
        self.assertTrue(result['valid'],result['errors'])
        self.assertTrue(result['circular'])
        self.assertFalse([site for site in result['restriction_sites'] if site['enzyme']=='BsmBI'])
        self.assertFalse(result['backbone_template']['eligible'])
        # An ordinary save of the same product must not demote the template.
        payload['save_as_backbone']=False
        again=self.client.post('/api/assemblies',json=payload).json()
        self.assertEqual(again['id'],saved['id'])
        self.assertTrue(again['backbone_template'])

    def test_backbone_template_requires_circular_product_and_usable_cassette(self):
        from app import database as store
        cases=[
            ('No sites','ACGT'*20,'TTTC',True),
            ('One site','ACGT'*10+'CGTCTCA'+'ACGT'*10,'TTTC',True),
            ('Additional cuts','GAGACGAAAGATATCCGTCTCA'+'ACGT'*10+'CGTCTCA'+'ACGT'*10,'TTTC',True),
            ('Linear with cassette','GAGACGAAAGATATCCGTCTCA'+'ACGT'*20,'ACTG',False),
        ]
        for index,(name,target,right,circular) in enumerate(cases):
            with self.subTest(name=name):
                sequence='GGTCTCA'+'TTTC'+target+right+'AGAGACC'
                with store.connection() as db:
                    fid=db.execute('INSERT INTO fragments(name,plate_id,well,sequence,yield_ng) VALUES(?,?,?,?,?)',
                        (name,'backbone_tests','B'+str(index+1),sequence,100)).lastrowid
                payload={'name':name,'fragment_ids':[fid],'save_as_backbone':True}
                preview=self.client.post('/api/assemblies/preview',json=payload).json()
                self.assertTrue(preview['valid'],preview['errors'])
                self.assertEqual(preview['circular'],circular)
                self.assertFalse(preview['backbone_template']['eligible'])
                count=len(store.plasmids())
                response=self.client.post('/api/assemblies',json=payload)
                self.assertEqual(response.status_code,422,response.text)
                self.assertEqual(len(store.plasmids()),count)

    def test_uploaded_linear_plasmid_can_be_assigned_as_backbone_after_import(self):
        from app import database as store
        source=next(p for p in store.plasmids() if p['name']=='pIVT_V2')
        content='>Uploaded FASTA backbone\n'+source['sequence']
        upload=self.client.post('/api/uploads/plasmid',files={'file':('vector.fasta',content,'text/plain')})
        self.assertEqual(upload.status_code,200,upload.text)
        pid=upload.json()['ids'][0]
        initial=self.client.get(f'/api/plasmids/{pid}').json()
        self.assertFalse(initial['circular'])
        self.assertTrue(initial['backbone_options']['eligible'])
        promoted=self.client.patch(f'/api/plasmids/{pid}',json={'backbone_template':True})
        self.assertEqual(promoted.status_code,200,promoted.text)
        self.assertTrue(promoted.json()['circular'])
        self.assertEqual(promoted.json()['recipe']['backbone_enzymes'],['BsmBI','BsmBI + BsaI'])
        # A duplicate default upload preserves settings assigned after the first upload.
        duplicate=self.client.post('/api/uploads/plasmid',files={'file':('vector.fasta',content,'text/plain')})
        self.assertEqual(duplicate.json()['ids'],[pid])
        with TestClient(app) as reopened:
            listed=next(p for p in reopened.get('/api/plasmids?enzyme=BsmBI').json() if p['id']==pid)
            self.assertTrue(listed['circular'])
            self.assertTrue(listed['recipe']['backbone_template'])
            self.assertEqual(len(listed['slots']),1)
            exported=parse_sequence_file(reopened.get(f'/api/plasmids/{pid}/genbank').text)[0]
            self.assertTrue(exported['circular'])
            self.assertEqual(exported['sequence'],source['sequence'])
        demoted=self.client.patch(f'/api/plasmids/{pid}',json={'circular':False})
        self.assertEqual(demoted.status_code,200,demoted.text)
        self.assertFalse(demoted.json()['circular'])
        self.assertFalse(demoted.json()['recipe'].get('backbone_template'))
        self.assertFalse(parse_sequence_file(self.client.get(f'/api/plasmids/{pid}/genbank').text)[0]['circular'])
        rejected=self.client.patch(f'/api/plasmids/{pid}',json={'circular':False,'backbone_template':True})
        self.assertEqual(rejected.status_code,422,rejected.text)
        self.assertFalse(store.get_plasmid(pid)['circular'])

    def test_backbone_upload_options_validate_all_records_before_import(self):
        from app import database as store
        sequence=next(p['sequence'] for p in store.plasmids() if p['name']=='pIVT_V2')
        files={'file':('vectors.fasta','>Upload one\n'+sequence+'\n>Upload two\n'+reverse_complement(sequence),'text/plain')}
        response=self.client.post('/api/uploads/plasmid?topology=circular&backbone_template=true',files=files)
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['count'],2)
        for pid in response.json()['ids']:
            saved=store.get_plasmid(pid)
            self.assertTrue(saved['circular'])
            self.assertTrue(saved['recipe']['backbone_template'])
            self.assertEqual(saved['recipe']['backbone_enzymes'],['BsmBI','BsmBI + BsaI'])
        # Selecting a backbone implies Circular even for a FASTA with From file selected.
        implicit=self.client.post('/api/uploads/plasmid?backbone_template=true',
            files={'file':('implicit.fa','>Implicit circle\n'+sequence,'text/plain')})
        self.assertEqual(implicit.status_code,200,implicit.text)
        self.assertTrue(store.get_plasmid(implicit.json()['ids'][0])['circular'])
        count=len(store.plasmids())
        mixed='>Must not be imported\n'+sequence+'\n>No cassette\n'+'ACGT'*30
        invalid=self.client.post('/api/uploads/plasmid?backbone_template=true',
            files={'file':('mixed.fa',mixed,'text/plain')})
        self.assertEqual(invalid.status_code,422,invalid.text)
        conflict=self.client.post('/api/uploads/plasmid?topology=linear&backbone_template=true',files=files)
        self.assertEqual(conflict.status_code,422,conflict.text)
        self.assertEqual(len(store.plasmids()),count)

    def test_genbank_reupload_adds_annotations_without_duplicates_or_losing_stock(self):
        from app import database as store
        from app.exports import genbank_bytes
        sequence='ACGT'*250
        response=self.client.post('/api/uploads/plasmid',files={'file':('reference.fa','>Annotation_test\n'+sequence,'text/plain')})
        self.assertEqual(response.status_code,200,response.text);pid=response.json()['ids'][0]
        with store.connection() as db:
            db.execute('UPDATE plasmids SET stock_ng_ul=25,plate_id=?,well=? WHERE id=?',('Stocks','A2',pid))
        annotations=[dict(type='CDS',label='Wrapped CDS',start=0,end=1000,raw_location='complement(join(901..1000,1..80))',qualifiers={'gene':['marker'],'product':['Test protein'],'note':['Imported note'],'ApEinfo_revcolor':['#4A90D9']}),
                     dict(type='protein_bind',label='Short site',start=25,end=26,qualifiers={'note':['One base annotation']}),
                     dict(type='misc_feature',label='Remote feature',start=0,end=0,raw_location='AB123.1:1..40',qualifiers={'note':['Remote annotation']})]
        files={'file':('annotated.gb',genbank_bytes('Annotation_test',sequence,annotations,True),'text/plain')}
        for _ in range(2):
            response=self.client.post('/api/uploads/plasmid',files=files)
            self.assertEqual(response.status_code,200,response.text);self.assertEqual(response.json()['ids'],[pid])
        record=self.client.get(f'/api/plasmids/{pid}').json()
        self.assertEqual(len(record['features']),4)
        wrapped=next(f for f in record['features'] if f['label']=='Wrapped CDS')
        self.assertEqual(wrapped['strand'],-1);self.assertEqual(len(wrapped['segments']),2)
        self.assertEqual(wrapped['qualifiers']['product'],['Test protein'])
        self.assertEqual((record['stock_ng_ul'],record['plate_id'],record['well']),(25,'Stocks','A2'))
        # Updating a qualifier refreshes that feature rather than duplicating it.
        annotations[0]['qualifiers']['note']=['Updated note']
        response=self.client.post('/api/uploads/plasmid',files={'file':('annotated.gb',genbank_bytes('Annotation_test',sequence,annotations,True),'text/plain')})
        self.assertEqual(response.status_code,200,response.text)
        self.client.post('/api/uploads/plasmid',files={'file':('reference.fa','>Annotation_test\n'+sequence,'text/plain')})
        exported=parse_sequence_file(self.client.get(f'/api/plasmids/{pid}/genbank').text)[0]
        self.assertEqual(len(exported['features']),4)
        self.assertEqual(next(f for f in exported['features'] if f['label']=='Wrapped CDS')['qualifiers']['note'],['Updated note'])

    def test_topology_overrides_preserve_uploaded_annotations_and_inventory_metadata(self):
        from app import database as store
        from app.exports import genbank_bytes
        source=next(p for p in store.plasmids() if p['name']=='pIVT_V2')
        content=genbank_bytes('Annotated upload',source['sequence'],source['features'],True)
        files={'file':('annotated.gb',content,'text/plain')}
        uploaded=self.client.post('/api/uploads/plasmid',files=files)
        self.assertEqual(uploaded.status_code,200,uploaded.text)
        pid=uploaded.json()['ids'][0]
        before=store.get_plasmid(pid)
        self.assertTrue(before['circular'])
        self.client.patch(f'/api/plasmids/{pid}',json={'backbone_template':True})
        with store.connection() as db:
            db.execute('UPDATE plasmids SET stock_ng_ul=25,plate_id=?,well=? WHERE id=?',('stocks','C2',pid))
        override=self.client.post('/api/uploads/plasmid?topology=linear',files=files)
        self.assertEqual(override.status_code,200,override.text)
        self.assertEqual(override.json()['ids'],[pid])
        after=store.get_plasmid(pid)
        self.assertFalse(after['circular'])
        self.assertFalse(after['recipe'].get('backbone_template'))
        self.assertEqual(after['features'],before['features'])
        self.assertEqual(after['sequence'],before['sequence'])
        self.assertEqual((after['stock_ng_ul'],after['plate_id'],after['well']),(25,'stocks','C2'))
        no_sites=store.save_plasmid('No cassette reference','ACGT'*100,circular=False)
        rejected=self.client.patch(f'/api/plasmids/{no_sites}',json={'backbone_template':True})
        self.assertEqual(rejected.status_code,422,rejected.text)
        self.assertFalse(store.get_plasmid(no_sites)['circular'])
        self.assertEqual(self.client.patch('/api/plasmids/999999',json={'circular':True}).status_code,404)

    def test_selected_backbone_defaults_to_required_circular_closure(self):
        from app import database as store
        vector = 'ACGCAA' + 'TTTG' + 'C' + 'GAGACC' + 'AAAGATATC' + 'GGTCTC' + 'A' + 'TTTC' + 'AACCGG' + 'A' * 1800
        donor = 'GCGTGGTCTCA' + 'TTTG' + 'A' * 100 + 'ATGG' + 'CGAGACCACGCA'
        backbone_id = store.save_plasmid('Closure regression backbone', vector)
        with store.connection() as db:
            fragment_id = db.execute('INSERT INTO fragments(name,plate_id,well,sequence,yield_ng) VALUES(?,?,?,?,?)',
                ('Wrong closing donor', 'regression', 'A1', donor, 100)).lastrowid
        payload = {'name': 'Wrong closing end', 'backbone_id': backbone_id,
            'fragment_ids': [fragment_id], 'circular': None}
        result = self.client.post('/api/assemblies/preview', json=payload)
        self.assertEqual(result.status_code, 200, result.text)
        result = result.json()
        self.assertTrue(result['circular'])
        self.assertFalse(result['topology_auto_detected'])
        self.assertFalse(result['valid'])
        self.assertTrue(result['junctions'][0]['compatible'])
        self.assertFalse(result['junctions'][-1]['compatible'])
        self.assertTrue(result['junctions'][-1]['circular_closure'])
        self.assertEqual(result['reaction']['backbone_fmol'], 50)
        self.assertIsNone(result['reaction']['input_fmol'])
        for endpoint in ('/api/assemblies', '/api/assemblies/genbank'):
            rejected = self.client.post(endpoint, json=payload)
            self.assertEqual(rejected.status_code, 422, rejected.text)
        # Legacy starter references retain the same circular validation default.
        recipe = next(recipe for recipe in self.recipes() if recipe['name'] == 'pIVT_V2')
        legacy = {'starter_fragment_id': recipe['starter_fragment_id'],
            'fragment_ids': recipe['fragment_ids'][:1]}
        result = self.client.post('/api/assemblies/preview', json=legacy).json()
        self.assertTrue(result['circular'])
        self.assertFalse(result['valid'])

    def test_no_backbone_inventory_upload_and_worksheet_need_no_f1(self):
        from app import database as store
        with store.connection() as db:
            db.execute("UPDATE fragments SET name='Ordinary fragment' WHERE name='F1'")
        for method in ('get', 'post'):
            response = (self.client.get('/api/inventory/resuspension', params={'no_backbone': True})
                if method == 'get' else self.client.post('/api/inventory/resuspension', json={'no_backbone': True}))
            self.assertEqual(response.status_code, 200, response.text)
            plan = response.json()
            self.assertEqual(plan['reference']['type'], 'none')
            self.assertIsNone(plan['reference']['length_bp'])
            self.assertEqual(len(plan['rows']), 102)
            self.assertTrue(all(row['target_fmol'] == 50 for row in plan['rows']))
        contradictory = self.client.get('/api/inventory/resuspension', params={'no_backbone': True, 'starter_fragment_id': 1})
        self.assertEqual(contradictory.status_code, 422)
        worksheet = self.client.get('/api/inventory/worksheet', params={'no_backbone': True})
        self.assertEqual(worksheet.status_code, 200, worksheet.text)
        wb = load_workbook(io.BytesIO(worksheet.content))
        self.assertEqual(wb['Settings']['B18'].value, 1)
        self.assertIsNone(wb['Settings']['B8'].value)
        wb.close()
        content = 'Name,Plate ID,Well,Sequence,Yield ng\nShort insert,no_backbone_plate,A1,' + 'ACGT' * 75 + ',100\n'
        upload = self.client.post('/api/uploads/plate?preview=true&no_backbone=true',
            files={'file': ('short.csv', content.encode(), 'text/csv')})
        self.assertEqual(upload.status_code, 200, upload.text)
        row = upload.json()['resuspension']['rows'][0]
        self.assertTrue(row['feasible'], row['warnings'])
        self.assertEqual(row['target_fmol'], 50)
        self.assertGreaterEqual(row['input_volume_ul'], 1)
        self.assertLessEqual(row['input_volume_ul'], 5)
        self.assertLess(row['resuspension_ul'], 100)

    def test_compare_includes_unsaved_build_list_assemblies_without_fragment_inference(self):
        recipe=next(r for r in self.recipes() if r['name']=='pHR_V2 · U6')
        payload=self.payload(recipe) | {'name':'Unsaved queued plasmid'}
        queued=self.client.post('/api/build-list',json=payload)
        self.assertEqual(queued.status_code,200,queued.text)
        exported=self.client.post('/api/assemblies/genbank',json=payload)
        before=self.client.get('/api/uploads').json()
        compared=self.client.post('/api/compare?infer=false',files={'file':('queued.gb',exported.content,'text/plain')})
        self.assertEqual(compared.status_code,200,compared.text)
        result=compared.json()
        exact=result['matches'][0]
        self.assertTrue(exact['exact'])
        aliases=exact['aliases']
        self.assertTrue(any(a['source']=='Build list' and a['name']=='Unsaved queued plasmid' for a in aliases),aliases)
        self.assertIn('Build list',result['source_counts'])
        self.assertGreater(len(result['fragment_matches']),0)
        self.assertEqual(result['inferred_count'],0)
        self.assertEqual(self.client.get('/api/uploads').json(),before)

    def test_compare_multiple_records_and_record_level_error_do_not_block_good_records(self):
        from app import database as store
        reference=store.plasmids()[0]
        sequence=reference['sequence']
        response=self.client.post('/api/compare?infer=false',files={'file':('multiple.fa',
            '>exact\n'+sequence+'\n>empty\n\n>oversize\n'+'A'*30001+'\n>reverse\n'+reverse_complement(sequence),'text/plain')})
        self.assertEqual(response.status_code,200,response.text)
        result=response.json()
        self.assertEqual(result['record_count'],4)
        records=result['records']
        self.assertTrue(records[0]['matches'][0]['exact'])
        self.assertIn('empty',records[1]['error'])
        self.assertIn('30,000',records[2]['error'])
        self.assertTrue(records[3]['matches'][0]['exact'])
        self.assertEqual(records[3]['matches'][0]['orientation'],'-')
        limit=self.client.post('/api/compare',files={'file':('too-many.fa',''.join('>r'+str(i)+'\nACGT\n' for i in range(101)),'text/plain')})
        self.assertEqual(limit.status_code,422)
        self.assertIn('100',limit.json()['detail'])
        invalid=self.client.post('/api/compare?enzyme=unknown',files={'file':('q.fa','>r\nACGT','text/plain')})
        self.assertEqual(invalid.status_code,422)

    def test_compare_uploaded_genbank_exact_reverse_and_point_mutation(self):
        recipe = next(recipe for recipe in self.recipes() if recipe['name'] == 'pHR_V2 · U6')
        exported = self.client.post('/api/assemblies/genbank', json=self.payload(recipe))
        exact = self.client.post('/api/compare?infer=true', files={'file': ('query.gb', exported.content, 'text/plain')})
        self.assertEqual(exact.status_code, 200, exact.text)
        result = exact.json()
        self.assertGreater(result['inferred_count'], 0)
        self.assertTrue(result['matches'][0]['exact'])
        self.assertEqual(result['matches'][0]['identity'], 1)
        self.assertIn(result['matches'][0]['source'], {'saved plasmid', 'inferred from plate fragments'})
        search = result['search'].get('metadata', result['search'])
        self.assertIn('complete_search', search)
        self.assertIn('routes_explored', search)
        self.assertIn('bounded', result['note'])
        sequence = parse_sequence_file(exported.text, 'query.gb')[0]['sequence']
        rc = reverse_complement(sequence[999:] + sequence[:999])
        reversed_match = self.client.post('/api/compare?infer=false', files={'file': ('reverse.fasta', ('>reverse\n' + rc).encode(), 'text/plain')}).json()['matches'][0]
        self.assertTrue(reversed_match['exact'])
        self.assertEqual(reversed_match['orientation'], '-')
        alternate = next(base for base in 'ACGT' if base != sequence[3432])
        mutated = sequence[:3432] + alternate + sequence[3433:]
        point_response = self.client.post('/api/compare?infer=false', files={'file': ('point.fasta', ('>point\n' + mutated).encode(), 'text/plain')})
        self.assertEqual(point_response.status_code, 200, point_response.text)
        point = point_response.json()['matches'][0]
        self.assertFalse(point['exact'])
        self.assertEqual(len(point['changes']), 1)
        self.assertEqual((point['changes'][0]['type'], point['changes'][0]['reference_start']), ('substitution', 3432))


class DatabaseSourceManifestTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.previous_database = os.environ.get('GOLDEN_GATE_DB')
        self.database = Path(self.directory.name) / 'replacement.sqlite3'
        os.environ['GOLDEN_GATE_DB'] = str(self.database)

    def tearDown(self):
        if self.previous_database is None:
            os.environ.pop('GOLDEN_GATE_DB', None)
        else:
            os.environ['GOLDEN_GATE_DB'] = self.previous_database
        self.directory.cleanup()

    def write_manifest(self, plate_files=None, genbank_files=None):
        self.database.with_suffix('.sources.json').write_text(json.dumps({
            'plate_files': plate_files or [], 'genbank_files': genbank_files or []}), encoding='utf8')

    def test_fresh_install_seeds_eight_current_twist_platemaps(self):
        manifest=json.loads((APP_ROOT/'data'/'golden_gate.sources.json').read_text())
        expected={Path(name).name:APP_ROOT/'data'/name for name in manifest['plate_files']}
        self.assertEqual(len(expected),7)
        with TestClient(app) as first:
            summary=first.get('/api/summary').json()
            self.assertEqual((summary['fragments'],summary['plates'],summary['plasmids']),(218,8,3))
            offered=first.get('/api/examples').json()
            self.assertEqual({row['filename'] for row in offered if row['kind']=='plate'},set(expected))
            for name,path in expected.items():
                downloaded=first.get('/api/examples/'+name)
                self.assertEqual(downloaded.status_code,200)
                self.assertEqual(downloaded.content,path.read_bytes())
            parts=first.get('/api/fragments',params={'limit':1000}).json()
            self.assertEqual(len({(p['plate_id'],p['well']) for p in parts}),218)
            kept,archived=parts[0],parts[-1]
            changed=first.patch('/api/fragments/'+str(kept['id']),json={'stock_ng_ul':7})
            self.assertEqual(changed.status_code,200,changed.text)
            first.post('/api/bin/archive',json={'fragment_ids':[archived['id']]})
            history=first.get('/api/uploads').json()
        with TestClient(app) as restarted:
            self.assertEqual(restarted.get('/api/summary').json()['fragments'],217)
            self.assertEqual(restarted.get('/api/fragments/'+str(kept['id'])).json()['stock_ng_ul'],7)
            self.assertEqual(restarted.get('/api/fragments/'+str(archived['id'])).status_code,404)
            self.assertEqual(restarted.get('/api/uploads').json(),history)

    def test_withdrawn_plate_archives_once_on_upgrade_and_can_be_restored(self):
        from app import database as store
        from app.imports import import_plate
        self.write_manifest()
        with TestClient(app):
            content=b'Name,Plate ID,Well,Sequence,Yield ng\nWithdrawn,pSHPs0309B323501IQ,A1,ACGTACGT,100\n'
            import_plate(content,'withdrawn.csv')
            self.assertEqual(len(store.fragments()),1)
        self.database.with_suffix('.sources.json').unlink()
        with TestClient(app) as upgraded:
            self.assertEqual(upgraded.get('/api/summary').json()['fragments'],218)
            self.assertNotIn('pSHPs0309B323501IQ',upgraded.get('/api/summary').json()['plate_ids'])
            archived=upgraded.get('/api/bin').json()['rows']
            self.assertTrue(any(row['kind']=='plate' and row['plate_id']=='pSHPs0309B323501IQ' for row in archived))
            store.move_to_bin([],[],restore=True,plate_ids=['pSHPs0309B323501IQ'])
        with TestClient(app) as restarted:
            self.assertIn('pSHPs0309B323501IQ',restarted.get('/api/summary').json()['plate_ids'])
            self.assertEqual(restarted.get('/api/summary').json()['fragments'],219)

    def test_replacement_sources_exclude_old_examples_and_survive_restart(self):
        sources = self.database.parent / 'inventory_sources'
        sources.mkdir()
        filename = 'platemap_pSHPs0921B1498027B.csv'
        (sources / filename).write_bytes((APP_ROOT / 'examples' / filename).read_bytes())
        self.write_manifest(['inventory_sources/' + filename])
        with TestClient(app) as first:
            summary = first.get('/api/summary').json()
            self.assertEqual((summary['fragments'], summary['plates'], summary['plasmids']), (6, 1, 3))
            history = first.get('/api/uploads').json()
            self.assertEqual([row['filename'] for row in history if row['kind'] == 'plate'], [filename])
            self.assertEqual(sum(row['kind'] == 'genbank' for row in history), 3)
            self.assertEqual(first.get('/api/examples').json(), [{'filename': filename, 'kind': 'plate'}])
            download = first.get('/api/examples/' + filename)
            self.assertEqual(download.status_code, 200)
            self.assertEqual(download.content, (sources / filename).read_bytes())
            excluded = next(path.name for path in (APP_ROOT / 'examples').glob('*.gb'))
            self.assertEqual(first.get('/api/examples/' + excluded).status_code, 404)
            fragment = first.get('/api/fragments').json()[0]
            updated = first.patch('/api/fragments/' + str(fragment['id']), json={'stock_ng_ul': 17})
            self.assertEqual(updated.status_code, 200)
        with TestClient(app) as restarted:
            summary = restarted.get('/api/summary').json()
            self.assertEqual((summary['fragments'], summary['plates'], summary['plasmids']), (6, 1, 3))
            self.assertEqual(restarted.get('/api/uploads').json(), history)
            self.assertEqual(restarted.get('/api/fragments/' + str(fragment['id'])).json()['stock_ng_ul'], 17)

    def test_empty_manifest_does_not_fall_back_to_examples(self):
        self.write_manifest()
        with TestClient(app) as client:
            self.assertEqual(client.get('/api/summary').json()['fragments'], 0)
            self.assertEqual(client.get('/api/summary').json()['plasmids'], 3)
            self.assertEqual(len(client.get('/api/uploads').json()), 3)
            self.assertTrue(all(row['kind'] == 'genbank' for row in client.get('/api/uploads').json()))
            self.assertEqual(client.get('/api/examples').json(), [])

    def test_manifest_rejects_sources_outside_database_directory(self):
        self.write_manifest(['../outside.csv'])
        with self.assertRaisesRegex(ValueError, 'within the database directory'):
            with TestClient(app):
                pass


if __name__ == '__main__':
    unittest.main()
