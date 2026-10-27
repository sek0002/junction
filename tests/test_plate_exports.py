"""All plate formats share saved stocks and fixed 50 fmol transfer calculations."""
import csv
import io
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook
from pypdf import PdfReader
from PIL import Image

from app.plate_exports import snapshot, export_bytes
from app.plates import fragment_site_info


def example_plate():
    seq='GGTCTCAACCTAACCGAGACGACGTACGTACGTCTCAAACCGGTATGCGAGACC'
    physical=len(seq)+20
    part={'id':1,'well':'A2','name':'Nested fragment','length':len(seq),'physical_length':physical,
          'inventory':{'stock_ng_ul':physical*.033/2,'resuspension_ul':50,'saved':True,
            'yield_ng':100,'input_volume_ul':4,'concentration_basis':'recorded linked values'},
          'archived_at':None,**fragment_site_info(seq,'A'*10+seq+'T'*10)}
    return {'plate_id':'Plate 1234','format':96,'archived_at':None,'fragments':[part]}


class PlateExportTests(unittest.TestCase):
    def test_fixed_amount_uses_physical_length_saved_stock_not_inventory_target(self):
        plate=example_plate();data=snapshot(plate);values=data['rows'][0]['values']
        self.assertAlmostEqual(values[7],2)
        self.assertNotEqual(values[7],plate['fragments'][0]['inventory']['input_volume_ul'])
        self.assertEqual(values[5:7],[50,plate['fragments'][0]['inventory']['stock_ng_ul']])
        self.assertEqual(values[8],'BsaI, BsmBI')
        self.assertEqual(values[9:11],['ACCT','TATG'])
        self.assertEqual(values[14:],[.05,'1:1'])
        self.assertEqual(plate['fragments'][0]['inventory']['input_volume_ul'],4)

    def test_missing_stocks_and_archived_wells_are_explicit_and_not_omitted(self):
        plate=example_plate();part=plate['fragments'][0];part['inventory']['stock_ng_ul']=None
        part['inventory']['resuspension_ul']=None;plate['archived_at']='2026-10-27'
        row=snapshot(plate)['rows'][0]
        self.assertIsNone(row['values'][7]);self.assertIsNone(row['values'][6])
        self.assertEqual(row['values'][12],'Archived');self.assertIn('Missing length or stock',row['warnings'])
        part['inventory']['stock_ng_ul']=.001
        self.assertIn('Use outside 1–5 µL',snapshot(plate)['rows'][0]['warnings'])

    def test_csv_bundle_contains_map_and_well_list_without_site_location_file(self):
        plate=example_plate()
        with ZipFile(io.BytesIO(export_bytes(plate,'csv'))) as archive:
            self.assertEqual(set(archive.namelist()),{'wells.csv','plate-map.png','Readme.txt'})
            rows=list(csv.DictReader(io.StringIO(archive.read('wells.csv').decode('utf-8-sig'))))
            self.assertEqual(rows[0]['Plate'],'Plate 1234');self.assertEqual(float(rows[0]['Use µL']),2)
            self.assertEqual(rows[0]['Target pmol'],'0.05')
            self.assertEqual(rows[0]['Enzymes'],'BsaI, BsmBI')
            self.assertEqual(rows[0]["5' overhang"],'ACCT')
            self.assertEqual(Image.open(io.BytesIO(archive.read('plate-map.png'))).format,'PNG')
            self.assertIn('0.05 pmol',archive.read('Readme.txt').decode())

    def test_excel_keeps_map_and_wells_without_sites_sheet(self):
        plate=example_plate();plate['fragments'][0]['name']='=SUM(1,2)'
        content=export_bytes(plate,'xlsx',384)
        wb=load_workbook(io.BytesIO(content),data_only=False)
        self.assertEqual(wb.sheetnames,['Plate map','Wells'])
        self.assertEqual(len(wb['Plate map']._images),1)
        self.assertEqual(wb['Wells']['C6'].data_type,'s')
        self.assertAlmostEqual(wb['Wells']['H6'].value,2)
        self.assertEqual(wb['Wells']['H6'].number_format,'0.00')
        self.assertEqual(wb['Wells'].freeze_panes,'D6')
        self.assertEqual(wb['Wells']['I6'].value,'BsaI, BsmBI')
        self.assertEqual(wb['Wells']['J6'].value,'ACCT')
        self.assertEqual(wb['Plate map'].page_setup.orientation,'landscape')

    def test_pdf_keeps_map_and_wells_without_enzyme_location_section(self):
        content=export_bytes(example_plate(),'pdf')
        reader=PdfReader(io.BytesIO(content));text='\n'.join(p.extract_text() for p in reader.pages)
        self.assertIn('Plate 1234',text);self.assertIn('Nested fragment',text)
        self.assertIn('BsmBI',text);self.assertIn('ACCT',text)
        self.assertNotIn('Enzymes and overhang locations',text);self.assertNotIn('Internal',text)
        self.assertGreater(len(reader.pages[0].images),0)
        self.assertEqual(len(reader.pages),2)

    def test_sparse_384_layout_includes_edge_wells_and_sorted_rows(self):
        plate=example_plate();plate['format']=384;part=plate['fragments'][0]
        plate['fragments']=[part|{'well':'P24'},part|{'well':'A10'},part|{'well':'A2'}]
        data=snapshot(plate)
        self.assertEqual([r['values'][1] for r in data['rows']],['A2','A10','P24'])
        content=export_bytes(plate,'pdf');self.assertTrue(content.startswith(b'%PDF'))
