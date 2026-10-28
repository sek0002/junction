"""Named rows pair dynamic inserts and retain one common set of native junctions."""
from copy import deepcopy
import csv
import io
import json
import os
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from auth_client import AuthenticatedClient as TestClient
from app.main import app
from app.design_engine import design
from app.design_table import table_design, sequence_metadata
from app.design_exports import workbook_bytes, fragment_rows
from app.sequence_engine import digest_fragment
from app.exports import parse_sequence_file
from test_design import clean_dna, placeholder


class NamedTableTests(unittest.TestCase):
    def setUp(self):
        self.sequence = clean_dna(2, 600)
        self.sections = [dict(name='Promoter', start=0, end=100, mode='constant'),
                         dict(name='Cargo A', start=100, end=200, mode='dynamic'),
                         dict(name='Cargo B', start=350, end=450, mode='dynamic')]
        self.entries = [placeholder(self.sequence, self.sections)]
        self.rows = [dict(name='Alpha', dynamic_inputs={'insert:100:200':'AACCGGTTAA', 'insert:350:450':'TTGGAACCAA'}),
                     dict(name='Beta', dynamic_inputs={'insert:100:200':clean_dna(13,220), 'insert:350:450':'CCTTAACCGGTT'})]

    def base(self, enzyme='BsaI', **kwargs):
        result = design(self.entries, enzyme, left_overhang='TTTG', right_overhang='ACCT', **kwargs)
        self.assertTrue(result['valid'], result['errors'])
        return result

    def test_paired_rows_preserve_all_constant_bases_and_ends_without_cartesian_products(self):
        base = self.base()
        result = table_design(base, self.rows, 'BsaI')
        self.assertTrue(result['valid'], result['errors'])
        self.assertEqual([row['name'] for row in result['assemblies']], ['Alpha','Beta'])
        for row, product in zip(self.rows,result['assemblies']):
            expected = 'TTTG'+self.sequence[:100]+row['dynamic_inputs']['insert:100:200']+self.sequence[200:350]+row['dynamic_inputs']['insert:350:450']+self.sequence[450:]+'ACCT'
            self.assertEqual(product['sequence'], expected)
        constants = [row for row in result['generated'] if row['row_name']=='Shared']
        self.assertEqual(len(constants), len([row for row in base['generated'] if not row['dynamic_regions']]))
        for original in base['generated']:
            variants = [row for row in result['generated'] if row['id'].split('_row')[0]==original['id']]
            self.assertEqual(len(variants), 2 if original['dynamic_regions'] else 1)
            for variant in variants:
                parsed = digest_fragment(variant['sequence'], 'BsaI')
                self.assertTrue(parsed['valid'], parsed['errors'])
                self.assertEqual((parsed['left_overhang'],parsed['right_overhang']), (original['left_overhang'],original['right_overhang']))
        self.assertEqual(len(result['constant_overhangs']), 2)
        self.assertEqual(base['assemblies'][0]['sequence'],'TTTG'+self.sequence+'ACCT')

    def test_both_mixed_adapter_orders_round_trip_under_dual_cut(self):
        for left,right in [('BsaI','BsmBI'),('BsmBI','BsaI')]:
            result = table_design(self.base('BsmBI + BsaI',left_adapter_enzyme=left,right_adapter_enzyme=right), self.rows, 'BsmBI + BsaI')
            self.assertTrue(result['valid'],result['errors'])
            for row in result['generated']:
                parsed=digest_fragment(row['sequence'],'BsmBI + BsaI')
                self.assertTrue(parsed['valid'],parsed['errors'])
                self.assertEqual(parsed['target'],row['target'])
                self.assertEqual(row['left_adapter_enzyme'],left)
                self.assertEqual(row['right_adapter_enzyme'],right)
        bad=design(self.entries,'BsaI',left_adapter_enzyme='BsmBI')
        self.assertFalse(bad['valid']);self.assertIn('5′ BsmBI',' '.join(bad['errors']))

    def test_replaced_reference_sites_do_not_leave_stale_product_warnings(self):
        self.entries[0]['sequence']=self.sequence[:130]+'GAGACC'+self.sequence[136:]
        base=self.base('BsmBI')
        self.assertTrue(any('Retained BsaI' in message for message in base['warnings']))
        result=table_design(base,self.rows,'BsmBI')
        self.assertTrue(result['valid'],result['errors'])
        self.assertFalse(any('Retained BsaI' in message for message in result['warnings']))
        self.assertFalse(any(site['enzyme']=='BsaI' for site in result['assembly_sites']))

    def test_invalid_rows_and_deleted_region_inputs_block_generation(self):
        for change, message in [({'name':' '},'row name'), ({'name':'ALPHA'},'unique'),
                                ({'dynamic_inputs':{'insert:100:200':'ACGN'}},'Dynamic region'),
                                ({'dynamic_inputs':{'insert:100:200':'GGTCTC','insert:350:450':'ACGT'}},'BsaI')]:
            rows=deepcopy(self.rows);rows[1].update(change)
            result=table_design(self.base(),rows,'BsaI')
            self.assertFalse(result['valid']);self.assertIn(message,' '.join(result['errors']))
            self.assertEqual(result['assemblies'],[])

    def test_sites_created_at_replacement_boundaries_block_export(self):
        self.entries[0]['sequence']=self.sequence[:97]+'GGT'+self.sequence[100:]
        base=self.base()
        rows=deepcopy(self.rows);rows[0]['dynamic_inputs']['insert:100:200']='CTCAACCTT'
        result=table_design(base,rows,'BsaI')
        self.assertFalse(result['valid']);self.assertIn('BsaI',' '.join(result['errors']))

    def test_adapter_free_exports_keep_overhangs_and_metadata_truthful(self):
        result=sequence_metadata(table_design(self.base(),self.rows,'BsaI'),False)
        self.assertTrue(result['valid'],result['errors'])
        for row in result['generated']:
            self.assertEqual(row['full_sequence'],row['left_overhang']+row['target_sequence']+row['right_overhang'])
            self.assertEqual(row['length'],len(row['target_sequence'])+8)
            self.assertEqual(row['left_re_site'],'');self.assertEqual(row['right_re_site'],'')
            self.assertFalse(any(site['enzyme']=='BsaI' for site in row['sites']))
        result=sequence_metadata(table_design(self.base(),self.rows,'BsaI'))
        for row in result['generated']:
            self.assertEqual(row['full_sequence'],row['five_prime_sequence']+row['target_sequence']+row['three_prime_sequence'])
            self.assertEqual(row['length'],row['target_length']+34)

    def test_excel_escapes_formulas_preserves_sequences_and_rejects_truncation(self):
        rows=deepcopy(self.rows);rows[0]['name']='=1+1'
        result=sequence_metadata(table_design(self.base(),rows,'BsaI'))
        wb=load_workbook(io.BytesIO(workbook_bytes(result)))
        self.assertEqual(wb.sheetnames,['Fragments','Assemblies','Overhangs'])
        expected=fragment_rows(result);sheet=wb['Fragments']
        self.assertEqual(sheet.cell(2,4).value,expected[0]['Target sequence'])
        self.assertEqual(sheet.cell(2,5).value,expected[0]['Full sequence'])
        self.assertEqual(wb['Assemblies'].cell(2,1).value,"'=1+1")
        result['generated'][0]['full_sequence']='A'*32768
        with self.assertRaisesRegex(ValueError,'cell limit'):workbook_bytes(result)


class TableAPITests(unittest.TestCase):
    def setUp(self):
        self.folder=tempfile.TemporaryDirectory();self.old=os.environ.get('GOLDEN_GATE_DB')
        path=Path(self.folder.name)/'test.sqlite3';os.environ['GOLDEN_GATE_DB']=str(path)
        path.with_suffix('.sources.json').write_text(json.dumps({'plate_files':[],'genbank_files':[]}))
        self.client=TestClient(app);self.client.__enter__()
        fixture=NamedTableTests();fixture.setUp()
        self.payload=dict(name='Paired library',enzyme='BsaI',entries=fixture.entries,left_overhang='TTTG',right_overhang='ACCT',use_table=True,table_rows=fixture.rows)

    def tearDown(self):
        self.client.__exit__(None,None,None)
        if self.old is None:os.environ.pop('GOLDEN_GATE_DB',None)
        else:os.environ['GOLDEN_GATE_DB']=self.old
        self.folder.cleanup()

    def test_all_download_formats_use_same_named_rows_and_full_target_records(self):
        response=self.client.post('/api/design/preview?generate=true',json=self.payload)
        self.assertEqual(response.status_code,200,response.text)
        result=response.json();self.assertTrue(result['valid'],result['errors'])
        for format in ('csv','xlsx','fasta','target_fasta','assemblies_csv','gb'):
            exported=self.client.post('/api/design/export?format='+format,json=self.payload)
            self.assertEqual(exported.status_code,200,exported.text[:200] if format!='xlsx' else '')
            if format=='csv':
                rows=list(csv.DictReader(io.StringIO(exported.content.decode('utf-8-sig'))))
                self.assertEqual([row['Full sequence'] for row in rows],[row['full_sequence'] for row in result['generated']])
                self.assertEqual([row['Target sequence'] for row in rows],[row['target'] for row in result['generated']])
            elif format=='gb':
                self.assertEqual([row['sequence'] for row in parse_sequence_file(exported.text)],[row['sequence'] for row in result['assemblies']])
            elif format in ('fasta','target_fasta'):
                values=[''.join(record.splitlines()[1:]) for record in exported.text.split('>')[1:]]
                key='target_sequence' if format=='target_fasta' else 'full_sequence'
                self.assertEqual(values,[row[key] for row in result['generated']])
            elif format=='assemblies_csv':
                rows=list(csv.DictReader(io.StringIO(exported.content.decode('utf-8-sig'))))
                self.assertEqual([row['Name'] for row in rows],['Alpha','Beta'])

    def test_missing_inputs_and_internal_sites_fail_quick_checks_and_exports(self):
        for value in ('ACGN','GGTC TC'):
            payload=deepcopy(self.payload);payload['table_rows'][0]['dynamic_inputs']['insert:100:200']=value
            result=self.client.post('/api/design/preview',json=payload).json()
            self.assertTrue(result['errors']);self.assertFalse(result['table_rows'][0]['valid'])
            self.assertEqual(self.client.post('/api/design/export',json=payload).status_code,422)
        payload=deepcopy(self.payload);payload['table_rows'][0]['dynamic_inputs'].pop('insert:100:200')
        self.assertEqual(self.client.post('/api/design/export',json=payload).status_code,422)

    def test_table_rows_persist_with_adapter_settings_in_private_workspace(self):
        payload=self.payload|dict(left_adapter_enzyme='BsmBI',right_adapter_enzyme='BsaI',include_re_sites=False)
        state={'design':{'activeId':'one','designs':[payload|dict(id='one')]}}
        self.assertEqual(self.client.put('/api/workspace',json={'data':state}).status_code,200)
        self.assertEqual(self.client.get('/api/workspace').json()['data'],state)


if __name__=='__main__':unittest.main()
