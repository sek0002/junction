"""Nested cassette sequence, physical-input and persistence regression checks."""
import io
import json
import os
from pathlib import Path
import tempfile
import unittest

from auth_client import AuthenticatedClient as TestClient
from openpyxl import load_workbook
from app.main import app
from app.exports import parse_sequence_file
from app.sequence_engine import digest_fragment, DUAL_ENZYME


def donor(enzyme,left,target,right):
    f,r=('GGTCTC','GAGACC') if enzyme=='BsaI' else ('CGTCTC','GAGACG')
    return f+'A'+left+target+right+'C'+r


def cassette(left='CACC',right='CTCG',reverse='GAGACG',forward='CGTCTC'):
    return left+'C'+reverse+'TTAA'+forward+'A'+right


class NestedAssemblyTests(unittest.TestCase):
    def setUp(self):
        self.folder=tempfile.TemporaryDirectory()
        self.old=os.environ.get('GOLDEN_GATE_DB')
        path=Path(self.folder.name)/'test.sqlite3'
        os.environ['GOLDEN_GATE_DB']=str(path)
        path.with_suffix('.sources.json').write_text(json.dumps({'plate_files':[],'genbank_files':[]}))
        self.client=TestClient(app)
        self.client.__enter__()
        self.parent=donor('BsaI','TTTG','ACACAC'+cassette()+'TATATA','ATGG')
        self.child=donor('BsmBI','CACC','GATGAA','CTCG')
        self.tail=donor('BsaI','ATGG','AATACG','TTTG')
        self.ids=self.upload([('Parent',self.parent),('Child',self.child),('Tail',self.tail)])
        self.slot=digest_fragment(self.parent,'BsaI')['nested_cassettes'][0]['id']

    def tearDown(self):
        self.client.__exit__(None,None,None)
        if self.old is None:os.environ.pop('GOLDEN_GATE_DB',None)
        else:os.environ['GOLDEN_GATE_DB']=self.old
        self.folder.cleanup()

    def upload(self,rows,plate='nested'):
        text='Name,Plate ID,Well,Sequence,Yield ng\n'+'\n'.join(f'{name},{plate},A{i},{sequence},30' for i,(name,sequence) in enumerate(rows,1))
        r=self.client.post('/api/uploads/plate?no_backbone=true',files={'file':(plate+'.csv',text,'text/csv')})
        self.assertEqual(r.status_code,200,r.text)
        return {p['name']:p['id'] for p in self.client.get('/api/fragments',params={'plate':plate,'limit':1000}).json()}

    def request(self,mode='together'):
        return {'name':'Nested assembly','enzyme':DUAL_ENZYME if mode=='together' else 'BsaI',
            'fragment_ids':[self.ids['Parent'],self.ids['Tail']],
            'nested_assemblies':[{'parent_fragment_id':self.ids['Parent'],'cassette_id':self.slot,
                'fragment_ids':[self.ids['Child']],'mode':mode}],
            'settings':{'planner_stock_mode':'saved'}}

    def preview(self,request):
        r=self.client.post('/api/assemblies/preview',json=request)
        self.assertEqual(r.status_code,200,r.text)
        return r.json()

    def test_reverse_complement_search_build_exports_and_queue_use_the_same_sequence(self):
        from app.sequence_engine import reverse_complement
        reverse=self.upload([('Reverse child',reverse_complement(self.child))],'reverse_child')['Reverse child']
        matches=self.client.get('/api/fragments',params={'enzyme':'BsmBI','left':'CACC','right':'CTCG','q':'Reverse child'}).json()
        self.assertEqual(len(matches),1);self.assertTrue(matches[0]['reverse_complement'])
        self.assertEqual(matches[0]['digest']['left_overhang'],'CACC')
        raw=self.client.get(f'/api/fragments/{reverse}',params={'enzyme':'BsmBI'}).json()
        self.assertEqual(raw['sequence'],reverse_complement(self.child));self.assertFalse(raw['reverse_complement'])
        request=self.request();request['nested_assemblies'][0]['fragment_ids']=[reverse]
        result=self.preview(request);self.assertTrue(result['valid'],result['errors'])
        self.assertTrue(result['nested_assemblies'][0]['fragment_checks'][0]['reverse_complement'])
        child=next(f for f in result['features'] if f.get('source')=='nested_fragment')
        self.assertEqual(child['strand'],-1)
        exported=self.client.post('/api/assemblies/genbank',json=request)
        self.assertEqual(exported.status_code,200,exported.text)
        record=parse_sequence_file(exported.text)[0]
        self.assertEqual(record['sequence'],result['sequence'])
        self.assertEqual(next(f for f in record['features'] if f['label']=='Reverse child')['strand'],-1)
        added=self.client.post('/api/build-list',json=request);self.assertEqual(added.status_code,200,added.text)
        queued=self.client.get('/api/build-list').json()['items'][0]['plan']
        self.assertTrue(queued['nested_assemblies'][0]['fragment_checks'][0]['reverse_complement'])

    def test_reversed_nested_parent_mirrors_child_feature_coordinates(self):
        from app.sequence_engine import reverse_complement
        backbone='ACGT'*20+'CCATCGAGACCTTAAGGTCTCACAAA'+'ACGT'*20
        uploaded=self.client.post('/api/uploads/plasmid?topology=circular',files={'file':('bb.fa','>RC_parent_backbone\n'+backbone,'text/plain')})
        self.assertEqual(uploaded.status_code,200,uploaded.text)
        request=self.request();request.update(backbone_id=uploaded.json()['ids'][0],fragment_ids=[self.ids['Parent']])
        result=self.preview(request);self.assertTrue(result['valid'],result['errors'])
        parent=next(c for c in result['fragment_checks'] if c['id']==self.ids['Parent'] and not c['is_backbone'])
        self.assertTrue(parent['reverse_complement'])
        child=next(f for f in result['features'] if f.get('source')=='nested_fragment')
        self.assertEqual(child['strand'],-1)
        self.assertEqual(result['sequence'][child['start']:child['end']],reverse_complement('CACCGATGAA'))

    def test_detect_inactive_and_active_internal_cassette(self):
        for enzyme in ('BsaI',DUAL_ENZYME):
            found=digest_fragment(self.parent,enzyme)['nested_cassettes']
            self.assertEqual(len(found),1)
            self.assertEqual((found[0]['left_overhang'],found[0]['right_overhang']),('CACC','CTCG'))
            self.assertEqual(found[0]['enzyme'],'BsmBI')
        bad=donor('BsaI','TTTG','CACC'+'CGTCTCACTCG'+'TTAA','ATGG')
        self.assertEqual(digest_fragment(bad,'BsaI')['nested_cassettes'],[])
        mixed=donor('BsaI','TTTG','AAA'+cassette(reverse='GAGACC')+'AAA','ATGG')
        self.assertEqual(digest_fragment(mixed,'BsaI')['nested_cassettes'][0]['enzyme'],DUAL_ENZYME)

    def test_together_sequence_mass_and_nested_features(self):
        result=self.preview(self.request())
        expected='TTTGACACACCACCGATGAACTCGTATATAATGGAATACG'
        self.assertTrue(result['valid'],result['errors'])
        self.assertTrue(result['circular'])
        self.assertEqual(result['sequence'],expected)
        self.assertFalse(any(s['enzyme'] in ('BsaI','BsmBI') for s in result['restriction_sites']))
        row=next(p for p in result['reaction']['rows'] if p['name']=='Parent')
        self.assertEqual(row['length_bp'],len(self.parent))
        self.assertEqual(len(result['reaction']['rows']),3)
        self.assertEqual(result['nested_reactions'],[])
        workbook=self.client.post('/api/assemblies/worksheet',json=self.request())
        book=load_workbook(io.BytesIO(workbook.content))
        selection=next(row for row in book['Fragment selection'].iter_rows(min_row=4,values_only=True) if row[0]=='Parent')
        self.assertIn('Together / BsmBI',selection[7])
        self.assertNotIn('ERROR:',selection[7])
        feature=next(f for f in result['features'] if f.get('source')=='nested_fragment')
        self.assertEqual(feature['label'],'Child')
        self.assertEqual(expected[feature['start']:feature['end']],'CACCGATGAA')
        self.assertTrue(all(p['valid'] for p in result['fragment_checks']))

    def test_separate_stock_is_not_inherited_or_assumed(self):
        request=self.request('separate')
        request['settings']['planner_stock_mode']='custom'
        result=self.preview(request)
        self.assertTrue(result['valid'],result['errors'])
        rows=result['reaction']['rows']
        self.assertEqual([r['name'] for r in rows],['Parent (prepared)','Tail'])
        self.assertIsNone(rows[0]['concentration_ng_ul'])
        self.assertIsNone(rows[0]['volume_ul'])
        self.assertEqual(rows[0]['length_bp'],len(self.parent)-12)
        self.assertEqual(len(result['nested_reactions']),1)
        self.assertEqual([r['name'] for r in result['nested_reactions'][0]['reaction']['rows']],['Parent','Child'])
        request['nested_assemblies'][0]['prepared_stock_ng_ul']=1
        result=self.preview(request)
        self.assertAlmostEqual(result['reaction']['rows'][0]['volume_ul'],.05*(len(self.parent)-12)*.66)
        self.assertIn('Prepared parent',result['reaction']['rows'][0]['concentration_basis'])

    def test_together_requires_active_inner_enzyme(self):
        request=self.request();request['enzyme']='BsaI'
        result=self.preview(request)
        self.assertFalse(result['valid'])
        self.assertTrue(any('Together requires' in e for e in result['errors']))
        self.assertFalse(next(p for p in result['fragment_checks'] if p['id']==self.ids['Parent'])['valid'])

    def test_bad_nested_junction_blocks_genbank(self):
        ids=self.upload([('Wrong',donor('BsmBI','CACT','AAA','CTCG'))],'bad')
        request=self.request();request['nested_assemblies'][0]['fragment_ids']=[ids['Wrong']]
        result=self.preview(request)
        self.assertFalse(result['valid'])
        self.assertTrue(any('5′ junction requires CACC' in e for e in result['errors']))
        self.assertFalse(result['nested_assemblies'][0]['applied'])
        self.assertEqual(self.client.post('/api/assemblies/genbank',json=request).status_code,422)

    def test_unfilled_or_unpaired_active_sites_block(self):
        parent=donor('BsaI','TTTG','AAA'+cassette()+'ACGTCTCAAAA','ATGG')
        ids=self.upload([('Extra',parent)],'extra')
        for mode in ('together','separate'):
            request=self.request(mode);request['fragment_ids'][0]=ids['Extra']
            request['nested_assemblies'][0].update(parent_fragment_id=ids['Extra'],cassette_id=digest_fragment(parent,'BsaI')['nested_cassettes'][0]['id'])
            result=self.preview(request)
            self.assertFalse(result['valid'])
            if mode=='separate':self.assertTrue(any('outside the selected cassette' in e for e in result['errors']))

    def test_collision_between_main_and_nested_is_mode_specific(self):
        parent=donor('BsaI','TTTG','AAA'+cassette(left='TTTG')+'AAA','ATGG')
        child=donor('BsmBI','TTTG','AGAGAA','CTCG')
        ids=self.upload([('Repeat parent',parent),('Repeat child',child)],'collision')
        for mode in ('together','separate'):
            request=self.request(mode);request['fragment_ids'][0]=ids['Repeat parent']
            request['nested_assemblies'][0].update(parent_fragment_id=ids['Repeat parent'],cassette_id=digest_fragment(parent,'BsaI')['nested_cassettes'][0]['id'],fragment_ids=[ids['Repeat child']])
            result=self.preview(request)
            self.assertEqual(result['valid'],mode=='separate',result['errors'])
            if mode=='together':self.assertTrue(any('collision' in e for e in result['errors']))

    def test_multiple_children_keep_shared_overhang_once(self):
        ids=self.upload([('First',donor('BsmBI','CACC','AA','GATG')),('Second',donor('BsmBI','GATG','CC','CTCG'))],'chain')
        request=self.request();request['nested_assemblies'][0]['fragment_ids']=[ids['First'],ids['Second']]
        result=self.preview(request)
        self.assertTrue(result['valid'],result['errors'])
        self.assertIn('CACCAAGATGCCCTCG',result['sequence'])
        self.assertEqual(len([f for f in result['features'] if f.get('source')=='nested_fragment']),2)

    def test_multiple_cassettes_require_all_openings_in_stage(self):
        parent=donor('BsaI','TTTG','AAA'+cassette()+'AAA'+cassette('AGTG','TGCC')+'AAA','ATGG')
        ids=self.upload([('Two cassettes',parent),('Other child',donor('BsmBI','AGTG','AATT','TGCC'))],'multiple')
        slots=digest_fragment(parent,'BsaI')['nested_cassettes']
        request=self.request('separate');request['fragment_ids'][0]=ids['Two cassettes']
        request['nested_assemblies'][0].update(parent_fragment_id=ids['Two cassettes'],cassette_id=slots[0]['id'])
        self.assertFalse(self.preview(request)['valid'])
        request['nested_assemblies'].append({'parent_fragment_id':ids['Two cassettes'],'cassette_id':slots[1]['id'],'fragment_ids':[ids['Other child']],'mode':'separate'})
        result=self.preview(request)
        self.assertTrue(result['valid'],result['errors'])
        self.assertEqual(len(result['nested_reactions']),1)
        self.assertEqual(len(result['nested_reactions'][0]['reaction']['rows']),3)

    def test_queue_and_exports_preserve_nested_recipe_and_features(self):
        request=self.request('separate')
        created=self.client.post('/api/build-list',json=request)
        self.assertEqual(created.status_code,200,created.text)
        item=created.json()
        self.assertEqual(item['request']['nested_assemblies'],[request['nested_assemblies'][0] | {'prepared_stock_ng_ul':None}])
        self.assertEqual(item['request']['nested_assemblies'][0]['mode'],'separate')
        self.assertEqual(len(item['plan']['nested_reactions']),1)
        self.assertEqual(item['status'],'Check volumes')
        for url in ('/api/assemblies/worksheet','/api/build-list/worksheet'):
            response=self.client.post(url,json=request) if url.startswith('/api/assemblies') else self.client.get(url)
            self.assertEqual(response.status_code,200,response.text[:100])
            book=load_workbook(io.BytesIO(response.content))
            self.assertIn('Nested assemblies',book.sheetnames)
            self.assertIn('Nested reactions',book.sheetnames)
            self.assertIn('Nested resuspension',book.sheetnames)
            values=list(book['Nested reactions'].values)
            self.assertTrue(any(row[3]=='Child' for row in values[1:]))
            if url.startswith('/api/assemblies'):
                prepared=next(row for row in book['Assembly pipetting'].iter_rows(min_row=4,values_only=True) if row[2]=='Parent (prepared)')
                self.assertIsNone(prepared[9])
                self.assertIsNone(prepared[11])
        response=self.client.post('/api/assemblies/genbank',json=request)
        self.assertEqual(response.status_code,200,response.text)
        record=parse_sequence_file(response.text)[0]
        self.assertTrue(any(f['label']=='Child' for f in record['features']))
        saved=self.client.post('/api/assemblies',json=request)
        self.assertEqual(saved.status_code,200,saved.text)

    def test_invalid_cassette_parent_and_duplicate_selections_rejected(self):
        request=self.request();request['nested_assemblies'][0]['cassette_id']='missing'
        self.assertEqual(self.client.post('/api/assemblies/preview',json=request).status_code,422)
        request=self.request();request['nested_assemblies']*=2
        self.assertEqual(self.client.post('/api/assemblies/preview',json=request).status_code,422)
        request=self.request();request['nested_assemblies'][0]['fragment_ids']=[self.ids['Parent']]
        self.assertEqual(self.client.post('/api/assemblies/preview',json=request).status_code,422)
        request=self.request();request['nested_assemblies'][0]['parent_fragment_id']=99999
        self.assertEqual(self.client.post('/api/assemblies/preview',json=request).status_code,422)

    def test_separate_stage_amount_covers_main_prepared_input(self):
        request=self.request('separate')
        request['settings'].update(insert_ratio=3,backbone_pmol=.075)
        result=self.preview(request)
        required=result['reaction']['rows'][0]['target_fmol']
        stage=result['nested_reactions'][0]['reaction']
        self.assertAlmostEqual(stage['rows'][0]['target_fmol'],required)
        self.assertAlmostEqual(stage['rows'][1]['target_fmol'],required*3)
        self.assertTrue(any('recovery' in message for message in stage['warnings']))

    def test_new_parent_child_junction_site_is_detected_in_stage(self):
        opening='GTCTA'+'GAGACG'+'TTAA'+'CGTCTCACTCG'
        parent=donor('BsaI','TTTG','ACACAC'+opening+'TATATA','ATGG')
        child=donor('BsmBI','GTCT','CAAA','CTCG')
        ids=self.upload([('Junction parent',parent),('Junction child',child)],'newsite')
        request=self.request('separate');request['fragment_ids'][0]=ids['Junction parent']
        request['nested_assemblies'][0].update(parent_fragment_id=ids['Junction parent'],cassette_id=digest_fragment(parent,'BsaI')['nested_cassettes'][0]['id'],fragment_ids=[ids['Junction child']])
        result=self.preview(request)
        self.assertFalse(result['valid'])
        self.assertTrue(any('new junction' in e for e in result['errors']))

    def test_same_inner_outer_enzyme_is_possible_together_but_not_staged(self):
        parent=donor('BsmBI','TTTG','ACACAC'+cassette()+'TATATA','ATGG')
        tail=donor('BsmBI','ATGG','AATACG','TTTG')
        ids=self.upload([('BsmBI parent',parent),('BsmBI tail',tail)],'sameenzyme')
        for mode in ('together','separate'):
            request=self.request(mode);request['enzyme']='BsmBI'
            request['fragment_ids']=[ids['BsmBI parent'],ids['BsmBI tail']]
            request['nested_assemblies'][0].update(parent_fragment_id=ids['BsmBI parent'],cassette_id=digest_fragment(parent,'BsmBI')['nested_cassettes'][0]['id'])
            result=self.preview(request)
            self.assertEqual(result['valid'],mode=='together',result['errors'])
