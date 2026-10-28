import io
import unittest

from openpyxl import load_workbook

from app.exports import calculate_volumes, genbank_bytes, make_workbook, parse_genbank, parse_sequence_file, reaction_summary, resuspension_plan


class ExportTests(unittest.TestCase):
    def test_palindromic_overhang_cells_are_yellow_and_other_values_are_unchanged(self):
        part={'id':1,'name':'TCGA','plate_id':'Plate','well':'A1','sequence':'GGTCTCAATAT'+'ACGTTACG'*5+'ACCTCGAGACC','yield_ng':300}
        sheet=load_workbook(io.BytesIO(make_workbook([part])) )['Fragment selection']
        self.assertEqual(sheet['E4'].value,'ATAT')
        self.assertEqual(sheet['E4'].fill.fgColor.rgb,'00FFF1A3')
        self.assertIn('self-ligate',sheet['E4'].comment.text)
        self.assertEqual(sheet['F4'].value,'ACCT')
        self.assertNotEqual(sheet['F4'].fill.fgColor.rgb,'00FFF1A3')
        self.assertEqual(sheet['A4'].value,'TCGA')
        self.assertNotEqual(sheet['A4'].fill.fgColor.rgb,'00FFF1A3')

    def test_compound_features_round_trip(self):
        sequence = "ACGT" * 50
        features = [{"type": "CDS", "start": 0, "end": 200, "label": "Wrapped feature", "raw_location": "complement(join(181..200,1..21))", "qualifiers": {"gene": ["test_gene"], "translation": ["MPEPTIDE"], "note": ['contains "quotes"']}}]
        exported = genbank_bytes("Example circular construct", sequence, features)
        record = parse_genbank(exported.decode())[0]
        self.assertEqual(record["sequence"], sequence)
        self.assertTrue(record["circular"])
        cds = next(feature for feature in record["features"] if feature["type"] == "CDS")
        self.assertEqual(cds["raw_location"], features[0]["raw_location"])
        self.assertEqual(cds["segments"], [{"start": 180, "end": 200, "strand": -1}, {"start": 0, "end": 21, "strand": -1}])
        self.assertEqual(cds["qualifiers"]["note"], ['contains "quotes"'])
        self.assertEqual(len(parse_genbank((exported + exported).decode())), 2)
        reround = parse_genbank(genbank_bytes(record["name"], record["sequence"], record["features"]).decode())[0]
        self.assertEqual(len(reround["features"]), 2)

    def test_fasta_and_invalid_dna(self):
        records = parse_sequence_file(">one\nACGT\n>two\nGATTACA\n", "example.fa")
        self.assertEqual([record["sequence"] for record in records], ["ACGT", "GATTACA"])
        with self.assertRaises(ValueError):
            parse_sequence_file("ACGT123", "bad.txt")

    def test_actual_physical_length_and_unknown_stock(self):
        parts = [
            {"name": "backbone", "role": "backbone", "physical_length": 5000, "stock_ng_ul": 25},
            {"name": "insert", "sequence": "A" * 1000, "construct_sequence": "A" * 1200, "yield_ng": 600},
        ]
        settings = {"backbone_ng": 50, "insert_ratio": 2, "target_stock_ng_ul": 20, "master_mix_volume_ul": 2}
        rows = calculate_volumes(parts, settings)
        self.assertAlmostEqual(rows[0]["target_fmol"], 15.15151515)
        self.assertEqual(rows[0]["volume_ul"], 2)
        self.assertAlmostEqual(rows[1]["target_ng"], 24)
        self.assertEqual(rows[1]["resuspension_ul"], 30)
        self.assertIsNone(rows[1]["volume_ul"])
        self.assertIsNone(reaction_summary(rows, settings)["water_volume_ul"])
        assumed = calculate_volumes(parts, dict(settings, assume_target_stock=True))
        self.assertAlmostEqual(assumed[1]["volume_ul"], 1.2)
        self.assertTrue(assumed[1]["concentration_basis"].startswith("ASSUMED"))
        self.assertAlmostEqual(reaction_summary(assumed, settings)["water_volume_ul"], 14.8)

    def test_print_workbook_formulas_and_missing_values(self):
        parts = [
            {"id": "b", "name": "Backbone", "role": "backbone", "plate_id": "plate1", "well": "A1", "sequence": "A" * 5000, "stock_ng_ul": 25},
            {"id": "i", "name": "=unsafe", "plate_id": "plate2", "well": "B2", "sequence": "A" * 1000, "yield_ng": 600},
        ]
        data = make_workbook(parts, {"name": "test", "parts": parts}, {"master_mix_volume_ul": 2})
        workbook = load_workbook(io.BytesIO(data))
        inventory = workbook["Resuspension inventory"]
        self.assertEqual(inventory["B3"].value, "'=unsafe")
        self.assertIsNone(inventory["J3"].value)
        self.assertIn("Actual stock unknown", inventory["M3"].value)
        self.assertIn("F3/G3", inventory["H3"].value)
        self.assertIn("IF", workbook["Assembly pipetting"]["K5"].value)
        self.assertEqual(inventory.page_setup.orientation, "landscape")
        self.assertEqual(inventory.print_title_rows, "$1:$1")
        self.assertTrue(inventory.print_area)

    def test_one_ul_inventory_plan_uses_reference_and_physical_length(self):
        fragments = [
            {"id": "one", "name": "One", "plate_id": "P1", "well": "A1", "sequence": "A" * 1000, "physical_length": 1200, "yield_ng": 600},
            {"id": "two", "name": "Two", "plate_id": "P1", "well": "A2", "physical_length": 2000, "yield_ng": 1000},
        ]
        settings = {"resuspension_mode": "one_ul", "backbone_length_bp": 5000, "backbone_ng": 50}
        rows = resuspension_plan(fragments, settings)
        for row, expected in zip(rows, [12, 20]):
            self.assertAlmostEqual(row["target_ng_ul"], expected)
            self.assertAlmostEqual(row["resuspension_ul"], 50)
        self.assertAlmostEqual(rows[0]["target_fmol"], rows[1]["target_fmol"])
        self.assertEqual(rows[0]["physical_length"], 1200)
        self.assertEqual(rows[0]["input_volume_ul"], 1)
        self.assertNotIn("stock_ng_ul", fragments[0])
        doubled = resuspension_plan(fragments, dict(settings, insert_ratio=2, input_volume_ul=2))
        self.assertAlmostEqual(doubled[0]["target_ng_ul"], 12)
        self.assertAlmostEqual(doubled[0]["target_fmol"], rows[0]["target_fmol"] * 2)

    def test_one_ul_inventory_requires_reference_and_does_not_assume_first_backbone(self):
        fragment = {"name": "F1", "physical_length": 1000, "yield_ng": 0}
        unresolved = resuspension_plan([fragment], {"resuspension_mode": "one_ul"})[0]
        self.assertIsNone(unresolved["target_ng_ul"])
        self.assertIsNone(unresolved["resuspension_ul"])
        self.assertTrue(any("reference length missing" in item for item in unresolved["warnings"]))
        settings = {"resuspension_mode": "one_ul", "backbone_length_bp": 5000, "backbone_ng": 50, "insert_ratio": 3}
        row = resuspension_plan([fragment], settings)[0]
        self.assertEqual(row["target_ng_ul"], 30)
        self.assertEqual(row["resuspension_ul"], 0)
        backbone_row = resuspension_plan([dict(fragment, is_backbone=True)], settings)[0]
        self.assertEqual(backbone_row["target_ng_ul"], 10)
        fixed = resuspension_plan([dict(fragment, yield_ng=600)])[0]
        self.assertEqual(fixed["target_ng_ul"], 20)
        self.assertEqual(fixed["resuspension_ul"], 30)

    def test_one_ul_reaction_plans_one_ul_but_preserves_recorded_stock(self):
        parts = [
            {"id": "b", "name": "Backbone", "role": "backbone", "physical_length": 5000, "yield_ng": 1000},
            {"id": "i", "name": "Insert", "physical_length": 1200, "yield_ng": 600},
        ]
        settings = {"resuspension_mode": "one_ul", "backbone_ng": 50, "assume_target_stock": True}
        rows = calculate_volumes(parts, settings)
        self.assertEqual([row["volume_ul"] for row in rows], [1, 1])
        for row, expected in zip(rows, [50, 12]):
            self.assertAlmostEqual(row["target_ng_ul"], expected)
        self.assertTrue(all("molar-matched" in row["concentration_basis"] for row in rows))
        measured = calculate_volumes([parts[0], dict(parts[1], stock_ng_ul=24, concentration_basis="measured")], settings)
        self.assertAlmostEqual(measured[1]["volume_ul"], .5)
        self.assertEqual(measured[1]["concentration_basis"], "measured")
        self.assertTrue(any("instead of the planned 1" in warning for warning in measured[1]["warnings"]))
        self.assertAlmostEqual(measured[1]["resuspension_ul"], 50)
        unknown = calculate_volumes(parts, dict(settings, assume_target_stock=False))
        self.assertTrue(all(row["volume_ul"] is None for row in unknown))

    def test_one_ul_workbook_dynamic_targets_and_printable_snapshot(self):
        parts = [
            {"id": "b", "name": "Backbone", "role": "backbone", "plate_id": "P1", "well": "A1", "physical_length": 5000, "yield_ng": 1000},
            {"id": "i", "name": "Insert", "plate_id": "P1", "well": "A2", "physical_length": 1200, "yield_ng": 600, "stock_ng_ul": 24},
        ]
        settings = {"resuspension_mode": "one_ul", "assume_target_stock": True, "insert_ratio": 2}
        workbook = load_workbook(io.BytesIO(make_workbook(parts, {"parts": parts}, settings)))
        self.assertEqual(workbook["Settings"]["B8"].value, 5000)
        self.assertEqual(workbook["Settings"]["B11"].value, "one_ul")
        self.assertEqual(workbook["Settings"]["B12"].value, 1)
        inventory = workbook["Resuspension inventory"]
        self.assertIn('Settings!B11="one_ul"', inventory["G3"].value)
        self.assertIn("Settings!B8", inventory["G3"].value)
        self.assertIn("Settings!B12", inventory["G3"].value)
        self.assertIn("Settings!B5", inventory["G3"].value)
        self.assertNotIn("Settings!B5", inventory["G2"].value)
        snapshot = workbook["Resuspension plan"]
        self.assertEqual(snapshot["F3"].value, 50)
        self.assertEqual(snapshot["F4"].value, 24)
        self.assertEqual(snapshot["G4"].value, 25)
        self.assertEqual(snapshot["H4"].value, 1)
        self.assertIn("unmeasured", snapshot["J4"].value)
        pipetting = workbook["Assembly pipetting"]
        self.assertIn("I5/Settings!B12", pipetting["J5"].value)
        self.assertIn("molar-matched", pipetting["M5"].value)

    def test_constrained_plan_uses_fifty_fmol_and_safe_rounding(self):
        fragments = [{"name": "F2", "physical_length": 344, "yield_ng": 824, "plate_id": "P", "well": "B1"}]
        settings = {"resuspension_mode": "constrained", "backbone_length_bp": 1808, "backbone_pmol": .05, "insert_ratio": 1}
        row = resuspension_plan(fragments, settings)[0]
        self.assertTrue(row["feasible"])
        self.assertEqual(row["target_fmol"], 50)
        self.assertAlmostEqual(row["target_ng"], 11.352)
        self.assertAlmostEqual(row["minimum_resuspension_ul"], 72.5863284003)
        self.assertEqual(row["resuspension_ul"], 72.59)
        self.assertGreaterEqual(row["input_volume_ul"], 1)
        self.assertLessEqual(row["input_volume_ul"], 5)
        self.assertAlmostEqual(row["target_ng_ul"], 824 / 72.59)

    def test_constrained_infeasible_mass_limit_and_unknowns_are_not_plans(self):
        settings = {"resuspension_mode": "constrained", "backbone_length_bp": 1808}
        cases = [
            ({"physical_length": 300, "yield_ng": 9}, "insufficient_mass"),
            ({"physical_length": 300, "yield_ng": 990}, "resuspension_limit"),
            ({"physical_length": 300}, "missing_data"),
        ]
        for fragment, reason in cases:
            row = resuspension_plan([fragment], settings)[0]
            self.assertEqual(row["feasibility"], reason)
            self.assertFalse(row["feasible"])
            self.assertIsNone(row["target_ng_ul"])
            self.assertIsNone(row["resuspension_ul"])
            self.assertIsNone(row["input_volume_ul"])
            self.assertTrue(row["warnings"])
        invalid = resuspension_plan([{"physical_length": 300, "yield_ng": 100}], dict(settings, backbone_length_bp=1499))[0]
        self.assertTrue(invalid["feasible"])
        self.assertEqual(invalid["target_fmol"],50)

    def test_constrained_strict_under_hundred_and_decimal_precision(self):
        settings = {"resuspension_mode": "constrained", "backbone_length_bp": 1808}
        allowed = resuspension_plan([{"physical_length": 300, "yield_ng": 985.05}], settings)[0]
        self.assertTrue(allowed["feasible"])
        self.assertEqual(allowed["resuspension_ul"], 99.5)
        self.assertLess(allowed["resuspension_ul"], 100)
        unrepresentable = resuspension_plan([{"physical_length": 300, "yield_ng": 49.5495}], dict(settings, min_input_volume_ul=1, max_input_volume_ul=1))[0]
        self.assertEqual(unrepresentable["feasibility"], "resuspension_precision")
        self.assertIsNone(unrepresentable["resuspension_ul"])

    def test_constrained_actual_stock_overrides_and_flags_outside_limits(self):
        parts = [
            {"name": "Backbone", "physical_length": 1808, "yield_ng": 1743, "role": "backbone"},
            {"name": "Insert", "physical_length": 344, "yield_ng": 824, "stock_ng_ul": 30},
        ]
        settings = {"resuspension_mode": "constrained", "backbone_pmol": .05, "assume_target_stock": True}
        rows = calculate_volumes(parts, settings)
        self.assertEqual(rows[0]["target_fmol"], 50)
        self.assertIsNone(rows[0]["resuspension_ul"])
        self.assertEqual(rows[0]["volume_ul"],1)
        self.assertEqual(rows[0]['feasibility'],'default_volume')
        self.assertIsNone(rows[0]['concentration_ng_ul'])
        self.assertNotIn("Backbone stock missing",rows[0]["warnings"])
        self.assertEqual(rows[1]["concentration_ng_ul"], 30)
        self.assertFalse(rows[1]["feasible"])
        self.assertEqual(rows[1]["feasibility"], "input_volume_limit")
        self.assertTrue(any("allowed input" in warning for warning in rows[1]["warnings"]))
        self.assertTrue(any("cannot meet" in warning for warning in reaction_summary(rows, settings)["warnings"]))

    def test_no_backbone_plans_equal_moles_for_short_inserts_and_no_backbone_stock(self):
        parts = [
            {"id": 1, "name": "Short first", "physical_length": 300, "yield_ng": 100, "plate_id": "P", "well": "A1"},
            {"id": 2, "name": "Second", "physical_length": 600, "yield_ng": 200, "plate_id": "P", "well": "A2"},
        ]
        settings = {"resuspension_mode": "constrained", "no_backbone": True,
            "backbone_stock_ng_ul": 100, "assume_target_stock": True}
        rows = calculate_volumes(parts, settings)
        for row in rows:
            self.assertEqual(row['role'], 'insert')
            self.assertEqual(row['target_fmol'], 50)
            self.assertTrue(row['feasible'], row['warnings'])
            self.assertGreaterEqual(row['volume_ul'], 1)
            self.assertLessEqual(row['volume_ul'], 5)
            self.assertNotEqual(row['concentration_ng_ul'], 100)
            self.assertAlmostEqual(row['target_ng'], 50 * row['length_bp'] * 660 / 1e6)
        workbook = load_workbook(io.BytesIO(make_workbook(parts, {'parts': parts}, settings)))
        self.assertEqual(workbook['Settings']['B18'].value, 1)
        self.assertIsNone(workbook['Settings']['B8'].value)
        self.assertEqual(workbook['Resuspension plan']['I3'].value, 50)
        self.assertEqual(workbook['Resuspension plan']['I4'].value, 50)
        self.assertEqual(workbook['Assembly pipetting']['B4'].value, 'Insert')
        self.assertEqual(workbook['Assembly pipetting']['B5'].value, 'Insert')
        self.assertIn('Settings!B18<>1', workbook['Resuspension inventory']['G2'].value)
        self.assertIn('Settings!B18<>1', workbook['Assembly pipetting']['H4'].value)
        self.assertFalse(any('first selected part treated as backbone' in str(cell.value)
            for row in workbook['Assembly pipetting'] for cell in row))
        workbook.close()

    def test_constrained_workbook_settings_and_formulas(self):
        parts = [{"id": "b", "name": "F1", "role": "backbone", "physical_length": 1808, "yield_ng": 1743}, {"id": "i", "name": "F2", "physical_length": 344, "yield_ng": 824}]
        workbook = load_workbook(io.BytesIO(make_workbook(parts, {"parts": parts}, {"resuspension_mode": "constrained", "assume_target_stock": True})))
        settings = workbook["Settings"]
        self.assertEqual(settings["B11"].value, "constrained")
        self.assertEqual(settings["B13"].value, .05)
        self.assertEqual(settings["B15"].value, 99.99)
        self.assertEqual(settings["B16"].value, 1)
        self.assertEqual(settings["B17"].value, 5)
        inventory = workbook["Resuspension inventory"]
        self.assertIn("CEILING", inventory["G3"].value)
        self.assertIn("FLOOR", inventory["G3"].value)
        self.assertIn("Settings!B15", inventory["G3"].value)
        self.assertEqual(workbook["Resuspension plan"]["G4"].value, 72.59)
        self.assertGreaterEqual(workbook["Resuspension plan"]["H4"].value, 1)
        self.assertIn("Settings!B13*1000", workbook["Assembly pipetting"]["H5"].value)

    def test_manual_backbone_workbook_has_no_estimated_stock_or_input_limits(self):
        parts=[{'id':'b','name':'Backbone','role':'backbone','physical_length':12000},
            {'id':'i','name':'Insert','physical_length':500,'yield_ng':100}]
        settings={'resuspension_mode':'constrained','backbone_pmol':.05,
            'backbone_stock_ng_ul':10,'assume_target_stock':True}
        row=calculate_volumes(parts,settings)[0]
        self.assertAlmostEqual(row['volume_ul'],39.6)
        self.assertTrue(row['feasible'])
        self.assertEqual(row['warnings'],[])
        wb=load_workbook(io.BytesIO(make_workbook(parts,{'parts':parts},settings)))
        self.assertNotIn('B10',wb['Assembly pipetting']['J4'].value)
        self.assertEqual(wb['Assembly pipetting']['N4'].value,'=""')
        self.assertIsNone(wb['Resuspension inventory']['H2'].value)
        self.assertIsNone(wb['Resuspension plan']['G3'].value)

    def test_unknown_backbone_uses_one_ul_and_keeps_stock_unknown(self):
        parts=[{'id':'b','name':'Backbone','role':'backbone','physical_length':5000},
            {'id':'i','name':'Insert','physical_length':1000,'stock_ng_ul':10,'yield_ng':500,'plate_id':'P1','well':'A1'}]
        settings={'resuspension_mode':'constrained','backbone_pmol':.05,'insert_ratio':1,'reaction_volume_ul':20,'master_mix_volume_ul':4,'min_input_volume_ul':2,'max_input_volume_ul':5}
        rows=calculate_volumes(parts,settings)
        self.assertEqual(rows[0]['volume_ul'],1)
        self.assertIsNone(rows[0]['concentration_ng_ul']);self.assertIsNone(rows[0]['molarity_nm'])
        self.assertTrue(rows[0]['feasible']);self.assertEqual(rows[0]['warnings'],[])
        self.assertEqual(rows[0]['concentration_basis'],'Default backbone volume: 1 µL')
        summary=reaction_summary(rows,settings)
        self.assertAlmostEqual(summary['water_volume_ul'],11.7)
        plan=resuspension_plan([parts[0]],settings|{'backbone_length_bp':5000})[0]
        self.assertEqual(plan['input_volume_ul'],1);self.assertTrue(plan['feasible'])
        self.assertIsNone(plan['target_ng_ul']);self.assertIsNone(plan['resuspension_ul'])
        known=calculate_volumes(parts,settings|{'backbone_stock_ng_ul':50})
        self.assertAlmostEqual(known[0]['volume_ul'],3.3)
        self.assertEqual(known[0]['concentration_ng_ul'],50)
        stored=calculate_volumes([parts[0]|{'stock_ng_ul':100},parts[1]],settings)
        self.assertAlmostEqual(stored[0]['volume_ul'],1.65)
        missing=calculate_volumes([parts[0],parts[1]|{'stock_ng_ul':None,'_planner_stock_fixed':True}],settings)
        self.assertEqual(missing[0]['volume_ul'],1);self.assertIsNone(missing[1]['volume_ul'])
        self.assertIsNone(reaction_summary(missing,settings)['water_volume_ul'])
        wb=load_workbook(io.BytesIO(make_workbook(parts,{'parts':parts},settings)))
        self.assertEqual(wb['Resuspension plan']['H3'].value,1)
        self.assertIn('),1,IF(',wb['Assembly pipetting']['K4'].value)
        self.assertIsNone(wb['Settings']['B4'].value)
        pipetting=wb['Assembly pipetting']
        labelled={row[0].value:row[0].row for row in pipetting if isinstance(row[0].value,str)}
        water_formula=pipetting.cell(labelled['Water µL'],11).value
        self.assertIn(f"Settings!B6-K{labelled['Total DNA volume µL']}-K{labelled['Master mix / other reagents µL']}",water_formula)
        wb.close()

    def test_invalid_assembly_prominent_and_dynamic_estimate_toggle(self):
        parts = [
            {"id": "b", "name": "Backbone", "role": "backbone", "sequence": "A" * 5000},
            {"id": "i", "name": "Insert", "sequence": "A" * 1000, "stock_ng_ul": 20},
        ]
        workbook = load_workbook(io.BytesIO(make_workbook(parts, {
            "name": "pIVT example", "parts": parts,
            "errors": ["F4 to F10 overhang mismatch"],
            "warnings": ["Internal BsmBI cassette retained"],
        }, {"assume_target_stock": True, "master_mix_volume_ul": 2})))
        self.assertEqual(workbook["Settings"]["B10"].value, 1)
        checks = workbook["Assembly checks"]
        self.assertIn("BLOCKED", checks["B2"].value)
        self.assertEqual(checks["B5"].value, "F4 to F10 overhang mismatch")
        self.assertEqual(checks["B6"].value, "Internal BsmBI cassette retained")
        pipetting = workbook["Assembly pipetting"]
        self.assertIn("BLOCKED", pipetting["A2"].value)
        self.assertIn("Settings!B10=1", pipetting["J4"].value)
        self.assertIn("'Resuspension inventory'!J3", pipetting["J5"].value)
        self.assertIn("ASSUMED", pipetting["M5"].value)
        total = next(row[10].value for row in pipetting.iter_rows() if row[0].value == "Total DNA volume µL")
        self.assertIn('COUNT(K4:K5)<2,""', total)

    def test_selection_sheet_infers_outer_enzyme_and_fusion_labels(self):
        parts = [
            {"name": "BsaI donor", "plate_id": "plateA", "well": "A1", "sequence": "GGTCTCATTTC" + "A" * 20 + "ACCTCGAGACC"},
            {"name": "BsmBI donor", "plate_id": "plateB", "well": "B2", "sequence": "CGTCTCATAAG" + "C" * 30 + "CTGCCGAGACG"},
        ]
        selection = load_workbook(io.BytesIO(make_workbook(parts)))["Fragment selection"]
        self.assertEqual([selection.cell(4, column).value for column in (4, 5, 6, 7)], ["BsaI", "TTTC", "ACCT", 20])
        self.assertEqual([selection.cell(5, column).value for column in (4, 5, 6, 7)], ["BsmBI", "TAAG", "CTGC", 30])
        self.assertEqual(selection["B4"].value, "plateA")
        self.assertEqual(selection["C5"].value, "B2")
        self.assertEqual(selection.page_setup.orientation, "landscape")
        self.assertEqual(selection.print_title_rows, "$1:$3")


if __name__ == "__main__":
    unittest.main()
