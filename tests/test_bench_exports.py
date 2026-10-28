"""Bench-guide exports must agree with the planner, across all formats."""
import csv
import io
import json
import os
from pathlib import Path
import tempfile
import unittest

from auth_client import AuthenticatedClient as TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.bench_exports import assembly_sections, build_sections, plate_key, guide_bytes
from app.exports import calculate_volumes, reaction_summary
from app.main import app, ReactionSettings


def fixture():
    parts = [dict(id=1, name='Backbone', role='backbone', physical_length=5000, stock_ng_ul=100, plate_id='pSHPs0405B231201IQ', well='A1'),
             dict(id=2, name='Insert', role='insert', physical_length=1000, stock_ng_ul=10, plate_id='pSHPs0622B228801IQ', well='B12', yield_ng=500, resuspension_ul=50)]
    settings = dict(resuspension_mode='constrained', backbone_pmol=.05, backbone_stock_ng_ul=100, insert_ratio=1, master_mix_volume_ul=4, reaction_volume_ul=20, planner_stock_mode='saved')
    rows = calculate_volumes(parts, settings)
    summary = reaction_summary(rows, settings)
    assembly = dict(name='Assembly A', parts=parts, errors=[], reaction=dict(rows=rows, water_ul=summary['water_volume_ul'], master_mix_ul=summary['master_mix_volume_ul'], reaction_volume_ul=20, warnings=summary['warnings']))
    return assembly, settings


def csv_table(content,title):
    rows=list(csv.reader(io.StringIO(content.decode('utf-8-sig'))));start=rows.index([title])+1
    header=rows[start];data=[]
    for row in rows[start+1:]:
        if not row:break
        data.append(dict(zip(header,row)))
    return data


class BenchExportTests(unittest.TestCase):
    def test_assembly_export_is_a4_portrait_in_pdf_and_all_excel_sheets(self):
        from openpyxl import Workbook
        assembly,settings=fixture();sections=assembly_sections(assembly,settings)
        legacy=Workbook();legacy.active.title='Detailed plan';legacy.active.page_setup.orientation='landscape'
        content=io.BytesIO();legacy.save(content)
        book=load_workbook(io.BytesIO(guide_bytes('xlsx',sections,lambda:content.getvalue())))
        for sheet in book:
            self.assertEqual(sheet.page_setup.orientation,'portrait',sheet.title)
            self.assertEqual(str(sheet.page_setup.paperSize),str(sheet.PAPERSIZE_A4),sheet.title)
            self.assertEqual(sheet.page_setup.fitToWidth,1)
            self.assertEqual(sheet.page_setup.fitToHeight,0)
        for page in PdfReader(io.BytesIO(guide_bytes('pdf',sections))).pages:
            self.assertAlmostEqual(float(page.mediabox.width),595.2756,places=3)
            self.assertAlmostEqual(float(page.mediabox.height),841.8898,places=3)

    def test_portrait_pagination_preserves_fifty_assemblies_and_long_metadata(self):
        from copy import deepcopy
        assembly,settings=fixture();sections=[]
        for index in range(50):
            item=deepcopy(assembly);item['name']=f'Assembly {index+1:02d}'
            sections.extend(assembly_sections(item,settings))
        pdf=PdfReader(io.BytesIO(guide_bytes('pdf',sections)))
        text='\n'.join(page.extract_text() for page in pdf.pages)
        for index in range(50):self.assertIn(f'Assembly {index+1:02d}',text)
        self.assertIn('Fragment metadata',text);self.assertIn('Plate key',text)
        self.assertTrue(all(page.mediabox.height>page.mediabox.width for page in pdf.pages))

    def test_multiple_assemblies_have_one_row_each_and_distinct_stock_metadata(self):
        from copy import deepcopy
        assembly,settings=fixture();second=deepcopy(assembly);second['name']='Assembly B'
        second['reaction']['rows'][0].update(concentration_ng_ul=50,volume_ul=3.3)
        sections=assembly_sections(assembly,settings)+assembly_sections(second,settings)
        book=load_workbook(io.BytesIO(guide_bytes('xlsx',sections)))
        self.assertEqual(book['Assemblies'].max_row,5)
        self.assertEqual(book['Fragment metadata'].max_row,6)
        rows=csv_table(guide_bytes('csv',sections),'Assemblies')
        self.assertEqual(len(rows),2)
        self.assertEqual(rows[0]['Inputs (Plate-Well-µL)'],'(0405-A1-1.65)\n(0622-B12-3.30)')
        self.assertEqual(rows[1]['Inputs (Plate-Well-µL)'],'(0405-A1-3.30)\n(0622-B12-3.30)')
        self.assertFalse(any(header.startswith('M0') for header in rows[0]))
        metadata=csv_table(guide_bytes('csv',sections),'Fragment metadata')
        self.assertEqual([float(r['Stock ng/µL']) for r in metadata],[100,10,50])

    def test_stored_stocks_and_water_agree_in_excel_csv_and_pdf(self):
        assembly, settings = fixture()
        sections = assembly_sections(assembly, settings, 'BsaI')
        self.assertAlmostEqual(sections[0]['rows'][0]['use'], 1.65)
        self.assertAlmostEqual(sections[0]['rows'][1]['use'], 3.3)
        self.assertEqual(sections[0]['rows'][1]['resuspend'], 50)
        self.assertAlmostEqual(sections[0]['rows'][-2]['use'], 11.05)
        book = load_workbook(io.BytesIO(guide_bytes('xlsx', sections)))
        self.assertEqual(book.active.title, 'Assemblies')
        self.assertEqual(book.active.max_row,4)
        self.assertEqual(book.active['C4'].value,'(0405-A1-1.65)  (0622-B12-3.30)')
        self.assertAlmostEqual(book.active['D4'].value,4.95)
        self.assertAlmostEqual(book.active['F4'].value,11.05)
        self.assertEqual(book['Fragment metadata']['G5'].value,50)
        self.assertEqual(book['Fragment metadata']['C4'].value,'0405')
        self.assertEqual(book['Fragment metadata']['C4'].font.color.rgb,book['Plate key']['A4'].font.color.rgb)
        rows=csv_table(guide_bytes('csv',sections),'Assemblies')
        self.assertEqual(rows[0]['Inputs (Plate-Well-µL)'],'(0405-A1-1.65)\n(0622-B12-3.30)')
        self.assertTrue({'Stage','DNA input pmol','Insert:backbone'}.isdisjoint(rows[0]))
        self.assertEqual(rows[0]['Other reagents µL'],'4.00')
        metadata=csv_table(guide_bytes('csv',sections),'Fragment metadata')
        self.assertEqual(metadata[1]['Plate'],'0622')
        pdf=PdfReader(io.BytesIO(guide_bytes('pdf',sections)))
        text='\n'.join(page.extract_text() for page in pdf.pages)
        for required in ['Assembly A','B12','0405','3.30','11.05','Other reagents','Plate key','Fragment metadata']:
            self.assertIn(required,' '.join(text.split()))

    def test_plate_text_colors_match_rich_text_metadata_pdf_and_csv_key(self):
        from openpyxl.cell.rich_text import TextBlock
        from app.assembly_matrix import matrix_data
        assembly,settings=fixture();sections=assembly_sections(assembly,settings)
        headers,rows,_,key=matrix_data(sections)
        self.assertTrue({'Stage','DNA input pmol','Insert:backbone'}.isdisjoint(headers))
        self.assertEqual(len(headers),len(rows[0]))
        colors={entry['short']:entry['text_color'] for entry in key.values()}
        self.assertEqual(len(colors),len(set(colors.values())))
        book=load_workbook(io.BytesIO(guide_bytes('xlsx',sections)),rich_text=True)
        inputs=book['Assemblies']['C4'].value
        self.assertEqual(str(inputs),'(0405-A1-1.65)  (0622-B12-3.30)')
        runs={run.text:run.font.color.rgb[-6:] for run in inputs if isinstance(run,TextBlock)}
        self.assertEqual(runs,colors)
        for i,entry in enumerate(key.values(),4):
            self.assertEqual(book['Plate key'].cell(i,1).font.color.rgb[-6:],entry['text_color'])
            self.assertEqual(book['Plate key'].cell(i,2).font.color.rgb[-6:],entry['text_color'])
            self.assertEqual(book['Fragment metadata'].cell(i,3).font.color.rgb[-6:],entry['text_color'])
        csv_key=csv_table(guide_bytes('csv',sections),'Plate key')
        self.assertEqual({e['Plate']:e['Color'].lstrip('#') for e in csv_key},colors)
        pdf=PdfReader(io.BytesIO(guide_bytes('pdf',sections)))
        page=pdf.pages[0]
        text=' '.join(page.extract_text().split())
        self.assertNotIn('DNA input pmol',text);self.assertNotIn('Insert:backbone',text);self.assertNotIn('Stage',text)
        rgb_commands=[tuple(float(x) for x in operands) for operands,operator in page.get_contents().operations if operator==b'rg']
        for color in colors.values():
            expected=tuple(int(color[i:i+2],16)/255 for i in (0,2,4))
            self.assertTrue(any(all(abs(a-b)<.00001 for a,b in zip(command,expected)) for command in rgb_commands))

    def test_many_plates_and_colliding_aliases_have_unique_readable_text_colors(self):
        from copy import deepcopy
        from app.assembly_matrix import matrix_data
        assembly,settings=fixture();sections=assembly_sections(assembly,settings)
        sample=sections[0]['rows'][0]
        sections[0]['rows']=[dict(deepcopy(sample),plate=f'Plate{i:02d}0405') for i in range(30)]
        *_,key=matrix_data(sections)
        colors=[entry['text_color'] for entry in key.values()]
        self.assertEqual(len(set(colors)),30)
        self.assertEqual(len({e['short'] for e in key.values()}),30)
        self.assertEqual(key,matrix_data(sections)[-1])
        for color in colors:
            def linear(channel):
                c=int(channel,16)/255
                return c/12.92 if c<=.04045 else ((c+.055)/1.055)**2.4
            luminance=sum(weight*linear(color[i:i+2]) for weight,i in zip((.2126,.7152,.0722),(0,2,4)))
            self.assertGreater(1.05/(luminance+.05),4.5)

    def test_twelve_large_assemblies_fit_one_portrait_page_without_descriptions(self):
        from copy import deepcopy
        assembly,settings=fixture();sections=[]
        for index in range(12):
            section=assembly_sections(assembly,settings)[0]
            section['name']=f'Build {index+1:02d}'
            sample=section['rows'][1]
            inputs=[dict(deepcopy(sample),id=100+j,name=f'Insert {j+1}',well=f'B{j+1}') for j in range(12)]
            section['rows']=inputs+[row for row in section['rows'] if row['kind'] in ('reagent','total')]
            sections.append(section)
        pdf=PdfReader(io.BytesIO(guide_bytes('pdf',sections)))
        first_page=pdf.pages[0].extract_text()
        for index in range(12):self.assertIn(f'Build {index+1:02d}',first_page)
        self.assertEqual(len(pdf.pages),2)
        for well in range(1,13):self.assertEqual(first_page.count(f'(0622-B{well}-3.30)'),12)
        self.assertNotIn('Fragment metadata',first_page)
        sizes=[]
        pdf.pages[0].extract_text(visitor_text=lambda text,cm,tm,font,size:sizes.append(size) if '3.30)' in text or text.strip()=='0622' else None)
        self.assertTrue(sizes);self.assertTrue(all(size>=8 for size in sizes))
        book=load_workbook(io.BytesIO(guide_bytes('xlsx',sections)))
        self.assertEqual(book['Assemblies']['C4'].value.splitlines()[0],'(0622-B1-3.30)  (0622-B2-3.30)')
        self.assertLess(book['Assemblies'].row_dimensions[4].height,110)
        self.assertEqual(book['Assemblies']['C4'].font.sz,11)
        for title in ('Assemblies','Fragment metadata','Plate key'):
            self.assertIsNone(book[title]['A2'].value)
            self.assertTrue(book[title].row_dimensions[2].hidden)
        text=' '.join(page.extract_text() for page in pdf.pages)
        csv_text=guide_bytes('csv',sections).decode('utf-8-sig')
        for description in ('Transfers in','Inputs are','Stock snapshot','Plate labels use'):
            self.assertNotIn(description,text);self.assertNotIn(description,csv_text)

    def test_metadata_prints_compactly_while_retaining_stock_audit_fields(self):
        from copy import deepcopy
        assembly,settings=fixture();sections=[]
        assembly['reaction']['rows'][1]['recorded_resuspension_ul']=50.01234567890123
        for index in range(50):
            item=deepcopy(assembly);item['name']=f'Assembly {index+1:02d}'
            sections.extend(assembly_sections(item,settings))
        book=load_workbook(io.BytesIO(guide_bytes('xlsx',sections)))
        metadata=book['Fragment metadata']
        self.assertEqual(metadata.max_row,5)
        self.assertIn('Assembly 50 / Final assembly',metadata['J4'].value)
        for column in ('I','J'):self.assertTrue(metadata.column_dimensions[column].hidden)
        self.assertIn('$H$5',str(metadata.print_area))
        self.assertLess(metadata.row_dimensions[4].height,35)
        self.assertLess(metadata.row_dimensions[5].height,35)
        self.assertEqual(metadata['G5'].value,50.01234567890123)
        self.assertEqual(metadata['F4'].value,100)
        csv_rows=csv_table(guide_bytes('csv',sections),'Fragment metadata')
        self.assertIn('Assembly 50 / Final assembly',csv_rows[0]['Assemblies'])
        pdf=PdfReader(io.BytesIO(guide_bytes('pdf',sections)))
        page=next(index for index,page in enumerate(pdf.pages) if 'Fragment metadata' in page.extract_text())
        self.assertEqual(len(pdf.pages)-page,1)
        self.assertNotIn('Stock basis',pdf.pages[page].extract_text())
        self.assertIn('Backbone',pdf.pages[page].extract_text())
        self.assertIn('Plate key',pdf.pages[page].extract_text())

    def test_missing_stocks_and_overfilled_reactions_are_visible(self):
        assembly, settings = fixture()
        assembly['reaction']['rows'][1].update(volume_ul=None, concentration_ng_ul=None, warnings=['Stock missing'])
        assembly['reaction'].update(water_ul=None, warnings=['Incomplete DNA volumes'])
        sections = assembly_sections(assembly, settings)
        rows=csv_table(guide_bytes('csv',sections),'Assemblies')
        self.assertEqual(rows[0]['Inputs (Plate-Well-µL)'],'(0405-A1-1.65)\n(0622-B12-?)')
        self.assertIn('Stock missing',rows[0]['Checks'])
        self.assertEqual(rows[0]['Water µL'],'')
        assembly['reaction'].update(water_ul=-2, warnings=['Reaction overfilled'])
        text = '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(guide_bytes('pdf', assembly_sections(assembly, settings)))).pages)
        self.assertIn('-2', text)
        self.assertIn('Reaction overfilled', text)

    def test_plate_aliases_are_unique_and_csv_strings_are_literal(self):
        assembly, settings = fixture()
        assembly['reaction']['rows'][0]['plate_id'] = 'pSHPs0622B999901IQ'
        assembly['reaction']['rows'][0]['name'] = '=HYPERLINK("bad")'
        sections = assembly_sections(assembly, settings)
        key = plate_key(sections)
        aliases = [entry['short'] for entry in key.values()]
        self.assertEqual(len(aliases), len(set(aliases)))
        self.assertEqual(set(aliases), {'0622-1', '0622-2'})
        book = load_workbook(io.BytesIO(guide_bytes('xlsx', sections)))
        self.assertTrue(book['Fragment metadata']['A4'].value.startswith("'="))
        rows=csv_table(guide_bytes('csv',sections),'Fragment metadata')
        self.assertTrue(rows[0]['Fragment'].startswith("'="))

    def test_separate_nested_reactions_precede_final_and_unavailable_builds_remain(self):
        assembly, settings = fixture()
        stage = dict(name='Nested first', settings=settings, enzyme='BsmBI', reaction=assembly['reaction'], errors=[])
        assembly['nested_reactions'] = [stage]
        sections = build_sections([dict(request=dict(name='Assembly A', enzyme='BsaI', settings=settings), plan=assembly), dict(request=dict(name='Gone', enzyme='BsaI', settings=settings), plan=None, error='Fragment archived')])
        self.assertEqual([s['stage'] for s in sections], ['Nested first', 'Final assembly', 'Unavailable'])
        self.assertIn('Fragment archived', guide_bytes('csv', sections).decode())
        text = ' '.join(' '.join(p.extract_text() for p in PdfReader(io.BytesIO(guide_bytes('pdf', sections))).pages).split())
        self.assertLess(text.index('Nested first'), text.rindex('Assembly A'))
        self.assertIn('Fragment archived', text)


class WorksheetFormatAPITests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.old = os.environ.get('GOLDEN_GATE_DB')
        path = Path(self.folder.name)/'test.sqlite3'
        os.environ['GOLDEN_GATE_DB'] = str(path)
        path.with_suffix('.sources.json').write_text(json.dumps({'plate_files': [], 'genbank_files': []}))
        self.client = TestClient(app)
        self.client.__enter__()
        content = 'Name,Plate ID,Well,Sequence,Yield ng,Stock ng/uL\nLeft,pSHPs0405B231201IQ,A1,GGTCTCATTTGAACCGGATGGC GAGACC,30,5\nRight,pSHPs0622B228801IQ,B12,GGTCTCAATGGCCAAAATTTGCGAGACC,30,5\n'
        r = self.client.post('/api/uploads/plate?no_backbone=true', files={'file': ('plate.csv', content, 'text/csv')})
        self.assertEqual(r.status_code, 200, r.text)
        ids = [row['id'] for row in self.client.get('/api/fragments').json()]
        self.request = dict(name='Two fragment assembly', fragment_ids=ids, settings={'planner_stock_mode': 'saved'})

    def tearDown(self):
        self.client.__exit__(None, None, None)
        if self.old is None:
            os.environ.pop('GOLDEN_GATE_DB', None)
        else:
            os.environ['GOLDEN_GATE_DB'] = self.old
        self.folder.cleanup()

    def test_plasmid_stock_groups_only_update_matching_backbones(self):
        pids=[]
        for i in (1,2):
            sequence='TTTGCGAGACC'+'AAAA'+'GGTCTCATTTG'+'ACGT'*(300+i*50)
            response=self.client.post('/api/uploads/plasmid?topology=circular&backbone_template=true',files={'file':('bb.fa',f'>BB{i}\n{sequence}','text/plain')})
            self.assertEqual(response.status_code,200,response.text)
            pids.append(next(p['id'] for p in self.client.get('/api/plasmids').json() if p['name']==f'BB{i}'))
        for pid in (pids[0],pids[0],pids[1]):
            response=self.client.post('/api/build-list',json=self.request|{'backbone_id':pid,'settings':{'planner_stock_mode':'saved','backbone_stock_ng_ul':100}})
            self.assertEqual(response.status_code,200,response.text)
        response=self.client.put('/api/build-list/backbone-stocks/apply',json={'kind':'plasmid','id':pids[0],'stock_ng_ul':40})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(len(response.json()['items']),2)
        items=self.client.get('/api/build-list').json()['items']
        self.assertEqual([item['request']['settings']['backbone_stock_ng_ul'] for item in items],[40,40,100])
        self.assertTrue(all(item['plan']['reaction']['rows'][0]['concentration_ng_ul']==40 for item in items[:2]))

    def test_export_routes_and_defaults_preserve_records(self):
        self.assertEqual(ReactionSettings().master_mix_volume_ul, 4)
        r = self.client.post('/api/assemblies/preview', json=self.request)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()['reaction']['master_mix_ul'], 4)
        self.assertEqual(r.json()['effective_settings']['insert_ratio'], 1)
        before = self.client.get('/api/inventory/records').json()
        self.assertEqual(self.client.post('/api/build-list', json=self.request).status_code, 200)
        summary=self.client.get('/api/build-list').json()['export']
        self.assertEqual(len(summary['rows']),1)
        self.assertIn('Inputs (Plate-Well-µL)',summary['headers'])
        self.assertNotIn('Input',summary['metadata_headers'])
        csv_rows=csv_table(self.client.get('/api/build-list/worksheet?format=csv').content,'Assemblies')
        self.assertEqual(csv_rows[0]['Inputs (Plate-Well-µL)'],summary['rows'][0][summary['headers'].index('Inputs (Plate-Well-µL)')])
        for format, mime in [('xlsx', 'application/vnd.openxmlformats'), ('csv', 'text/csv'), ('pdf', 'application/pdf')]:
            requests = [self.client.post('/api/assemblies/worksheet?format='+format, json=self.request),
                        self.client.get('/api/build-list/worksheet?format='+format),
                        self.client.post('/api/inventory/worksheet?format='+format, json={'no_backbone': True, 'planner_stock_mode': 'saved'}),
                        self.client.get('/api/inventory/worksheet', params={'format': format, 'no_backbone': True, 'planner_stock_mode': 'saved'})]
            for response in requests:
                self.assertEqual(response.status_code, 200, response.text[:200] if response.status_code != 200 else '')
                self.assertIn(mime, response.headers['content-type'])
                self.assertIn('.'+format, response.headers['content-disposition'])
                if format == 'xlsx':
                    book = load_workbook(io.BytesIO(response.content))
                    self.assertIn(book.active.title,('Assemblies','Bench guide'))
                elif format == 'pdf':
                    self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertEqual(before, self.client.get('/api/inventory/records').json())
        self.assertEqual(self.client.get('/api/build-list/worksheet?format=exe').status_code, 422)
